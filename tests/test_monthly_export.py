from copy import deepcopy
import hashlib
import sqlite3

from openpyxl import load_workbook
import pytest

from fba_replenishment.core import connect
from fba_replenishment.monthly_export import export_monthly_report
from fba_replenishment.monthly_plan import calculate_monthly_plan
from fba_replenishment.review import REVIEW_SCHEMA
from fba_replenishment.sku_collection import INBOUND_EXCEPTION_SCHEMA


def read_rows(ws):
    values = list(ws.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def api_row(sku='A-size 5', **changes):
    result = dict(
        marketplace='US', sku=sku, asin='ASIN-A', family=None,
        status='SCENARIO_ONLY', parameters_confirmed=False, parameter_status='SCENARIO_ONLY',
        fulfillable=5, reserved=2, unfulfillable=1,
        inbound_working=0, inbound_shipped=20, inbound_receiving=0,
        inventory_inbound=20, eligible_inbound=20, unknown_eta_inbound=0,
        inventory_complete=True, inbound_complete=True, data_complete=True,
        planning_data_complete=True, recommended_qty=85, daily_sales=1,
        coverage_days=5, stockout_risk='STOCKOUT_BEFORE_REPLENISHMENT',
        inbound_rule='confirmed_eta', lead_days=35, target_days=75, safety_days=0,
        sales_timezone='America/Los_Angeles', asof='2026-10-13',
        window_start='2026-07-16', stock_observed_at='2026-10-15T01:20:30Z',
        anomalies=[], anomaly_codes=[], stop_reason='', moq=1, case_pack=1,
    )
    for days in (7, 30, 60, 90):
        result.update({f'sales_{days}d': days, f'daily_sales_{days}d': 1,
                       f'sales_complete_{days}d': True,
                       f'observed_sales_{days}d': days,
                       f'observed_sales_days_{days}d': days})
    result.update(changes)
    return result


def monthly_row(sku='A-size 5', units=60, **api_changes):
    return calculate_monthly_plan([api_row(sku, **api_changes)], [dict(
        sku=sku, monthly_units=units, monthly_status='VALID',
        source_rows=[2], source_sheet='销量', source_sha256='a' * 64,
    )])[0]


def export_rows(tmp_path, rows, *, con=None, settings=None, **kwargs):
    path = tmp_path / 'monthly.xlsx'
    summary = export_monthly_report(rows, path, con, settings or {},
                                   input_audit=kwargs.pop('input_audit', {}), **kwargs)
    return load_workbook(path), summary


def test_monthly_demand_is_not_replaced_by_api_30_day_demand(tmp_path):
    item = monthly_row()
    assert item['recommended_qty'] == 195
    wb, summary = export_rows(tmp_path, [item])
    main = read_rows(wb['补货清单'])[0]
    assert main['用户月销量'] == 60
    assert main['月销日均件数'] == 2
    assert main['建议补货件数'] == 195
    assert main['API近30天销量'] == 30
    compare = read_rows(wb['月销与7_30_90对比'])[0]
    assert compare['API近30天日均'] == 1
    assert compare['原API30天试算补货件数'] == 85
    assert compare['月销量基准补货件数'] == 195
    assert summary['recommended_units'] == 195


def test_required_chinese_sheets_and_explicit_scenario_assumptions(tmp_path):
    wb, summary = export_rows(tmp_path, [monthly_row()])
    assert {'补货清单', '建议补货数量', '缺货风险', '月销与7_30_90对比', 'FBA库存明细', '在途货件明细',
            '异常与待确认', '已停补SKU', '月销表外风险SKU', '参数与计算说明',
            '数据来源与审计', '汇总统计'} <= set(wb.sheetnames)
    parameters = {r['项目']: r['内容'] for r in read_rows(wb['参数与计算说明'])}
    assert parameters['补货覆盖总天数'] == 110
    assert parameters['参数确认状态'] == summary['parameter_status'] == 'SCENARIO_ONLY'
    assert '不得作为正式采购依据' in parameters['重要提示']
    assert '30天' in parameters['月销量期间假设']
    assert 'ATVPDKIKX0DER' in parameters['美国站范围']
    assert '2026-07-16' in parameters['API销量日期']
    assert '2026-10-13' in parameters['API销量日期']
    assert wb['参数与计算说明']['B2'].font.bold


def test_month_assumption_cannot_become_confirmed_from_settings_or_flags(tmp_path):
    item = monthly_row()
    item.update(parameters_confirmed=True, parameter_status='CONFIRMED')
    wb, summary = export_rows(tmp_path, [item], settings={'parameters_confirmed': True})
    row = read_rows(wb['补货清单'])[0]
    assert row['参数确认状态'] == summary['parameter_status'] == 'SCENARIO_ONLY'
    assert '不得作为正式采购依据' in row['建议用途']


def test_effective_period_and_parameter_confirmation_are_both_required(tmp_path):
    item = calculate_monthly_plan([api_row()], [dict(
        sku='A-size 5', monthly_units=60, monthly_status='VALID', source_rows=[2])],
        monthly_period_status='USER_CONFIRMED_PERIOD', parameters_confirmed=True)[0]
    wb, summary = export_rows(tmp_path, [item])
    assert summary['parameter_status'] == 'CONFIRMED'
    assert read_rows(wb['补货清单'])[0]['参数确认状态'] == 'CONFIRMED'


def test_unmatched_us_sku_remains_unknown_and_retains_input_source(tmp_path):
    item = calculate_monthly_plan([], [dict(sku='unknown', monthly_units=7,
        monthly_status='VALID', source_rows=[3, 4])])[0]
    wb, summary = export_rows(tmp_path, [item])
    row = read_rows(wb['补货清单'])[0]
    assert row['用户月销量'] == 7
    assert row['美国站精确匹配状态'] == 'NOT_FOUND_US'
    assert row['建议补货件数'] is row['FBA可售'] is row['API近30天销量'] is None
    assert row['上传表来源行'] == '3; 4'
    assert summary['matched_sku_count'] == 0
    assert summary['unmatched_sku_count'] == summary['blocked_sku_count'] == 1


@pytest.mark.parametrize('inventory_complete', [True, False])
def test_direct_stopped_sku_is_always_zero(tmp_path, inventory_complete):
    item = monthly_row(status='STOPPED', recommended_qty=0,
                       inventory_complete=inventory_complete, stop_reason='Manual stop')
    item['recommended_qty'] = 999
    wb, summary = export_rows(tmp_path, [item, monthly_row('same-family')])
    items = {r['SKU']: r for r in read_rows(wb['补货清单'])}
    assert items['A-size 5']['建议补货件数'] == 0
    assert items['same-family']['建议补货件数'] == 195
    assert summary['stopped_sku_count'] == 1
    assert read_rows(wb['已停补SKU'])[0]['SKU'] == 'A-size 5'


def test_original_safety_gate_and_pending_order_remain_blocked(tmp_path):
    item = monthly_row(planning_data_complete=False, status='DATA_REVIEW_REQUIRED',
                       anomalies=[dict(code='ORDERS_PENDING_STATUS', blocking=True,
                                       reason='Pending order has unconfirmed quantities', source='SP-API')])
    item['recommended_qty'] = 100
    wb, summary = export_rows(tmp_path, [item])
    assert read_rows(wb['补货清单'])[0]['建议补货件数'] is None
    assert summary['blocked_sku_count'] == 1
    assert read_rows(wb['异常与待确认'])[0]['异常代码'] == 'ORDERS_PENDING_STATUS'


def test_unknown_api_comparison_values_do_not_become_zero(tmp_path):
    item = monthly_row(sales_30d=0, sales_complete_30d=False,
                       observed_sales_days_30d=29,
                       status='DATA_REVIEW_REQUIRED', planning_data_complete=False)
    wb, _ = export_rows(tmp_path, [item])
    compare = read_rows(wb['月销与7_30_90对比'])[0]
    assert compare['API近30天销量'] is compare['API近30天日均'] is None
    assert compare['上传月销与API30天比值'] is None
    assert compare['月销量基准补货件数'] is None


def test_export_defense_blocks_inconsistent_complete_flags(tmp_path):
    item = monthly_row()
    item.update(sales_complete_90d=False, planning_data_complete=True,
                recommended_qty=999, status='SCENARIO_ONLY')
    wb, summary = export_rows(tmp_path, [item])
    assert read_rows(wb['补货清单'])[0]['建议补货件数'] is None
    assert summary['blocked_sku_count'] == 1


def test_sales_reference_disagreement_is_nonblocking_and_not_averaged(tmp_path):
    wb, _ = export_rows(tmp_path, [monthly_row()])
    row = read_rows(wb['月销与7_30_90对比'])[0]
    assert row['API7天折合30天件数'] == 30
    assert row['API90天折合30天件数'] == 30
    assert row['上传月销折合30天件数'] == 60
    assert row['上传月销与API30天比值'] == 2
    assert 'MONTHLY_DIFFERS_FROM_API30' in row['销量差异人工审核提示']
    assert row['月销量基准补货件数'] == 195
    issues = read_rows(wb['异常与待确认'])
    assert any(r['异常代码'] == 'MONTHLY_DIFFERS_FROM_API30' and r['是否阻断补货'] is False for r in issues)


def test_api_zero_ratio_is_undefined_not_infinite(tmp_path):
    wb, _ = export_rows(tmp_path, [monthly_row(sales_30d=0)])
    compare = read_rows(wb['月销与7_30_90对比'])[0]
    assert compare['上传月销与API30天比值'] is None
    assert 'API30_ZERO_RATIO_UNDEFINED' in compare['销量差异人工审核提示']


def test_conflicting_inbound_preserves_raw_quantities_and_unknown_balance(tmp_path):
    con = connect(tmp_path / 'source.db')
    con.executescript(REVIEW_SCHEMA + INBOUND_EXCEPTION_SCHEMA)
    con.execute('INSERT INTO inbound VALUES(?,?,?,?,?,?)', ('US', 'A-size 5', 'FBA-X', 0, '', 'RECEIVING'))
    con.execute('INSERT INTO sku_inbound_exception VALUES(?,?,?,?,?,?,?,?,?)',
                ('US', 'A-size 5', 'FBA-X', 'FNSKU-X', 'RECEIVING', 10, 25, None, 'INBOUND_QUANTITY_CONFLICT'))
    con.commit()
    wb, summary = export_rows(tmp_path, [monthly_row()], con=con)
    inbound = read_rows(wb['在途货件明细'])
    assert len(inbound) == 1
    assert inbound[0]['原始发运件数'] == 10
    assert inbound[0]['原始接收件数'] == 25
    assert inbound[0]['剩余在途件数'] is None
    assert read_rows(wb['补货清单'])[0]['建议补货件数'] is None
    assert read_rows(wb['补货清单'])[0]['可抵扣在途'] is None
    assert summary['recommended_units'] == 0
    assert summary['blocked_sku_count'] == 1
    con.close()


def test_database_access_is_batched_read_only_and_scoped(tmp_path):
    path = tmp_path / 'source.db'
    con = connect(path)
    con.executescript(REVIEW_SCHEMA)
    for market, sku, shipment in [('US', 'A-size 5', 'FBA-US'),
                                  ('CA', 'A-size 5', 'FBA-CA'),
                                  ('US', 'outside', 'FBA-OUT')]:
        con.execute('INSERT INTO inbound VALUES(?,?,?,?,?,?)', (market, sku, shipment, 20, '', 'SHIPPED'))
    con.commit()
    con.close()
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    con = sqlite3.connect(path.as_uri() + '?mode=ro', uri=True)
    statements = []
    con.set_trace_callback(statements.append)
    inputs = [monthly_row()]
    original = deepcopy(inputs)
    wb, _ = export_rows(tmp_path, inputs, con=con)
    assert inputs == original
    assert {r['货件编号'] for r in read_rows(wb['在途货件明细'])} == {'FBA-US'}
    assert len([s for s in statements if 'FROM inbound ' in s]) == 1
    assert all(not s.lstrip().upper().startswith(('UPDATE', 'INSERT', 'DELETE', 'CREATE')) for s in statements)
    con.close()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_missing_monthly_values_for_omitted_cohort_are_not_zero(tmp_path):
    wb, summary = export_rows(tmp_path, [monthly_row()], omitted_rows=[api_row('outside')])
    row = read_rows(wb['月销表外风险SKU'])[0]
    assert row['月销量状态'] == 'MONTHLY_SALES_NOT_PROVIDED'
    assert row['月销量基准补货件数'] is None
    assert row['API近30天销量'] == 30
    assert summary['omitted_risk_sku_count'] == 1
    assert summary['cohort_count'] == 1


def test_omitted_direct_stopped_sku_keeps_explicit_zero_without_monthly_demand(tmp_path):
    stopped = api_row('omitted-stop', status='STOPPED', recommended_qty=999,
                      stop_reason='Manual stop remains effective')
    wb, summary = export_rows(tmp_path, [monthly_row()],
                              omitted_rows=[stopped, api_row('omitted-normal')])
    rows = {r['SKU']: r for r in read_rows(wb['月销表外风险SKU'])}
    assert rows['omitted-stop']['补货状态'] == 'STOPPED'
    assert rows['omitted-stop']['月销量基准补货件数'] == 0
    assert '不解除停补' in rows['omitted-stop']['处理说明']
    assert 'Manual stop remains effective' in rows['omitted-stop']['处理说明']
    assert rows['omitted-normal']['月销量基准补货件数'] is None
    assert summary['omitted_risk_sku_count'] == 2


def test_sku_priority_has_stockout_before_covered_and_no_sales(tmp_path):
    urgent = monthly_row('urgent', units=60, fulfillable=0)
    covered = monthly_row('covered', units=1, fulfillable=1000)
    zero = monthly_row('zero', units=0)
    wb, _ = export_rows(tmp_path, [zero, covered, urgent])
    assert [r['SKU'] for r in read_rows(wb['补货清单'])] == ['urgent', 'covered', 'zero']


def test_beijing_inventory_time_and_source_hash_are_retained(tmp_path):
    wb, _ = export_rows(tmp_path, [monthly_row()], input_audit={
        'source_file': 'sku销量.xlsx', 'source_sha256': 'a' * 64,
        'input_data_rows': 1489, 'record_count': 1489,
        'issues': ['MONTHLY_QUANTITY_HEADER_NUMERIC'], 'quantity_header': 1,
    }, settings={'inventory_observed_at': '2026-10-15T01:20:30Z'})
    assert read_rows(wb['补货清单'])[0]['库存采集时间_北京时间'] == '2026-10-15T09:20:30+08:00'
    audit = {r['审计项目']: r['审计值'] for r in read_rows(wb['数据来源与审计'])}
    assert audit['source_sha256'] == 'a' * 64
    assert audit['input_data_rows'] == 1489
    assert audit['quantity_header'] == 1
    assert audit['inventory_observed_at_北京时间'] == '2026-10-15T09:20:30+08:00'


def test_all_external_strings_are_literal_and_unallowed_payload_is_omitted(tmp_path):
    formula = '=HYPERLINK("https://example.invalid","SKU")'
    item = monthly_row(formula)
    item.update(access_token='SECRET_SENTINEL', customer_name='PII_SENTINEL',
                order_id='ORDER_SENTINEL', raw_response='RAW_SENTINEL')
    wb, _ = export_rows(tmp_path, [item], settings={'access_token': 'SECRET_SENTINEL'},
        input_audit={'source_sha256': 'a' * 64, 'authorization': 'SECRET_SENTINEL'},
        collection_audit={'summary': {'orders_collected': 11, 'buyer_name': 'PII_SENTINEL'},
                          'audit': [{'kind': 'orders', 'orders': 11, 'order_id': 'ORDER_SENTINEL',
                                     'raw_response': 'RAW_SENTINEL'}]})
    found_formula = False
    for ws in wb:
        for row in ws:
            for cell in row:
                assert cell.value not in {'SECRET_SENTINEL', 'PII_SENTINEL', 'ORDER_SENTINEL', 'RAW_SENTINEL'}
                if cell.value == formula:
                    found_formula = True
                    assert cell.data_type == 's'
    assert found_formula


def test_atomic_serialization_failure_preserves_previous_workbook(tmp_path, monkeypatch):
    output = tmp_path / 'monthly.xlsx'
    output.write_bytes(b'previous-valid-workbook')
    from openpyxl import Workbook

    def failed_save(self, path):
        path.write_bytes(b'incomplete')
        raise OSError('controlled serialization failure')

    monkeypatch.setattr(Workbook, 'save', failed_save)
    with pytest.raises(OSError):
        export_monthly_report([monthly_row()], output, None, {}, input_audit={})
    assert output.read_bytes() == b'previous-valid-workbook'
    assert sorted(p.name for p in tmp_path.iterdir()) == ['monthly.xlsx']


def test_summary_matches_exact_uploaded_cohort_and_not_omitted_rows(tmp_path):
    unknown = calculate_monthly_plan([], [dict(sku='unknown', monthly_units=7,
        monthly_status='VALID', source_rows=[8])])[0]
    rows = [monthly_row(), monthly_row('stopped', status='STOPPED'),
            monthly_row('covered', units=1, fulfillable=1000), unknown]
    wb, summary = export_rows(tmp_path, rows, omitted_rows=[api_row('outside')])
    assert summary['cohort_count'] == 4
    assert summary['matched_sku_count'] == 3
    assert summary['unmatched_sku_count'] == 1
    assert summary['monthly_units_total'] == 128
    assert summary['recommended_sku_count'] == 1
    assert summary['recommended_units'] == 195
    assert summary['blocked_sku_count'] == summary['stopped_sku_count'] == 1
    assert summary['inventory_positive_sku_count'] == 3
    assert summary['omitted_risk_sku_count'] == 1
    assert summary['parameter_status'] == 'SCENARIO_ONLY'
    assert len(read_rows(wb['补货清单'])) == 4
    queued = read_rows(wb['建议补货数量'])
    assert len(queued) == summary['recommended_sku_count']
    assert sum(r['建议补货件数'] for r in queued) == summary['recommended_units']
    assert {r['SKU'] for r in queued} == {'A-size 5'}


@pytest.mark.parametrize('mode', ['duplicate', 'canada', 'wrong-setting', 'overlap'])
def test_invalid_identity_or_overlapping_cohorts_are_rejected(tmp_path, mode):
    rows, settings, omitted = [monthly_row()], {}, []
    if mode == 'duplicate':
        rows.append(monthly_row())
    elif mode == 'canada':
        rows[0]['marketplace'] = 'CA'
    elif mode == 'wrong-setting':
        settings = {'marketplace': 'CA'}
    else:
        omitted = [api_row()]
    with pytest.raises(ValueError):
        export_monthly_report(rows, tmp_path / 'report.xlsx', None, settings,
                              input_audit={}, omitted_rows=omitted)
    assert not (tmp_path / 'report.xlsx').exists()
