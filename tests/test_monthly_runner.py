"""End-to-end offline monthly reporting retains source data and safety gates."""
from datetime import datetime, timedelta, timezone
import hashlib
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from fba_replenishment.core import connect
from fba_replenishment.discovery import ensure_schema
from fba_replenishment.review import REVIEW_SCHEMA
from fba_replenishment.monthly_runner import run_monthly_report


@pytest.fixture
def monthly_source(tmp_path):
    db = tmp_path/'source.db'
    con = connect(db)
    con.executescript(REVIEW_SCHEMA)
    ensure_schema(con)
    stamp = datetime.now(timezone.utc).isoformat()
    end = datetime.now(timezone.utc).date() - timedelta(days=1)
    con.execute("INSERT INTO marketplace_binding VALUES('US','ATVPDKIKX0DER','UTC')")
    for sku in ('selected', 'omitted', 'stopped'):
        con.execute('INSERT INTO sku VALUES(?,?,?,?,1)', ('US',sku,None,'B000000001'))
        con.execute('INSERT INTO stock VALUES(?,?,?,?,?,?)', ('US',sku,5,100,200,stamp))
        con.execute('INSERT INTO inventory_details VALUES(?,?,?,?,?,?)', ('US',sku,0,0,0,stamp))
        con.executemany('INSERT INTO sales VALUES(?,?,?,?)',
            [('US',sku,(end-timedelta(days=i)).isoformat(),1) for i in range(90)])
    con.execute("INSERT INTO stop_flag VALUES('US','stopped','human stop','2026-01-01',1)")
    for kind in ('stock','sales','inbound'):
        con.execute('INSERT INTO sync_log(at,marketplace,kind,status,row_count,detail) VALUES(?,?,?,?,?,?)',
                    (stamp,'US',kind,'SUCCESS',3,'complete snapshots'))
    con.commit()
    con.close()
    source = tmp_path/'monthly.xlsx'
    wb = Workbook()
    ws = wb.active
    for values in [('sku',1),('selected',60),('stopped',60),('missing',30)]:
        ws.append(values)
    wb.save(source)
    return db,source,end


def _hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def test_monthly_run_uses_uploaded_demand_and_preserves_both_sources(monthly_source,tmp_path,monkeypatch):
    db,source,end = monthly_source
    before = [_hash(p) for p in (db,source)]
    monkeypatch.setattr('fba_replenishment.cli.SPAPIClient',lambda *a,**k:pytest.fail('Offline monthly report must not create a live client'))
    out = tmp_path/'report.xlsx'
    result = run_monthly_report(source,db,out,asof=end.isoformat(),
        collection_audit={'summary':{'inventory_observed_at':'2026-10-15T03:22:17+00:00',
                                     'inbound_observed_at':'2026-10-15T03:23:11+00:00'}})
    assert result['cohort_count']==3 and result['unmatched_sku_count']==1
    assert result['recommended_units']==215  # 60/30*110 - 5; reserved is not deducted again.
    assert result['stopped_sku_count']==1 and result['blocked_sku_count']==1
    assert result['omitted_risk_sku_count']==1
    assert [_hash(p) for p in (db,source)]==before
    wb = load_workbook(out,read_only=True,data_only=True)
    ws = wb['补货清单']
    iterator = ws.iter_rows(values_only=True)
    header = next(iterator)
    records = {r[header.index('SKU')]:dict(zip(header,r)) for r in iterator}
    assert records['selected']['用户月销量']==60
    assert records['selected']['API近30天销量']==30
    assert records['missing']['建议补货件数'] is None
    assert records['stopped']['建议补货件数']==0
    audit_ws=wb['数据来源与审计']
    audit_rows=list(audit_ws.iter_rows(values_only=True))
    assert ('inventory_observed_at_北京时间','2026-10-15T11:22:17+08:00') in audit_rows
    assert ('inbound_observed_at_北京时间','2026-10-15T11:23:11+08:00') in audit_rows
    wb.close()


def test_cli_monthly_does_not_connect_or_request_amazon(monthly_source,tmp_path,monkeypatch,capsys):
    from fba_replenishment.cli import main
    db,source,end=monthly_source
    monkeypatch.setattr('fba_replenishment.cli.connect',lambda *a,**k:pytest.fail('Offline command must not migrate or write the database'))
    monkeypatch.setattr('fba_replenishment.cli.SPAPIClient',lambda *a,**k:pytest.fail('Offline command must not authenticate'))
    monkeypatch.setattr('sys.argv',['fba-plan','--db',str(db),'monthly-plan',str(source),
        '--source-db',str(db),'--asof',end.isoformat(),'--output',str(tmp_path/'cli.xlsx')])
    main()
    assert 'SCENARIO_ONLY' in capsys.readouterr().out


@pytest.mark.parametrize('target', ['source','db'])
def test_output_never_overwrites_a_source(monthly_source,target):
    db,source,end=monthly_source
    output=source if target=='source' else db
    before=_hash(output)
    with pytest.raises(ValueError,match='overwrite'):
        run_monthly_report(source,db,output,asof=end.isoformat())
    assert _hash(output)==before


def test_transport_confirmation_without_monthly_scope_is_still_trial(monthly_source,tmp_path):
    db,source,end=monthly_source
    result=run_monthly_report(source,db,tmp_path/'unconfirmed.xlsx',asof=end.isoformat(),
        monthly_period_status='USER_CONFIRMED_PERIOD',settings={'parameters_confirmed':True})
    assert result['parameter_status']=='SCENARIO_ONLY'


def test_text_confirmation_is_not_accepted_as_a_boolean(monthly_source,tmp_path):
    db,source,end=monthly_source
    with pytest.raises(ValueError,match='booleans'):
        run_monthly_report(source,db,tmp_path/'invalid.xlsx',asof=end.isoformat(),
                           settings={'parameters_confirmed':'False'})
