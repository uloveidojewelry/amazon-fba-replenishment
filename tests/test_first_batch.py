from hashlib import sha256
import sqlite3

from openpyxl import load_workbook
import pytest

from fba_replenishment.batch_review import prepare_batch_review
from fba_replenishment.core import connect
from fba_replenishment.discovery import discover
from fba_replenishment.first_batch import extract_first_batch, readonly_database, simulate_first_batch_file
from fba_replenishment.first_batch_export import export_first_batch


@pytest.fixture
def prepared(tmp_path,request):
    db=tmp_path/'first.db';con=connect(db)
    def inventory(sku):
        return {'sellerSku':sku,'asin':'B000000001','fnSku':'FN-'+sku,'condition':'NewItem',
                'inventoryDetails':{'fulfillableQuantity':1,'reservedQuantity':{'totalReservedQuantity':0},
                'unfulfillableQuantity':{'totalUnfulfillableQuantity':0},'inboundWorkingQuantity':0,
                'inboundShippedQuantity':0,'inboundReceivingQuantity':0}}
    families=getattr(request,'param',{'a':'F','b':'F'})
    discover(con,[inventory('a'),inventory('b')],'US','ATVPDKIKX0DER',
             source_candidates={s:{'candidate_family':families[s],'basis':'USER_WORKBOOK_EXACT_SKU'} for s in ('a','b')})
    result=prepare_batch_review(con,tmp_path/'source')
    con.close()
    yield db,result['review_path'],tmp_path


def extract(path):
    return extract_first_batch(path,expected_skus=2,expected_groups=1)


def test_first_source_never_copies_previous_approval(prepared):
    _,path,_=prepared
    w=load_workbook(path);s=w['批量审核决定'];h={c.value:c.column for c in s[1]}
    s.cell(2,h['decision'],'APPROVED');s.cell(2,h['reviewer'],'Prior reviewer');s.cell(2,h['approved_family'],'F');w.save(path)
    p=extract(path)
    assert p['groups'][0]['decision']=='NEEDS_REVIEW'
    assert p['groups'][0]['approved_family'] is None and p['groups'][0]['reviewer'] is None
    assert p['summary']['approval_count']==0


def test_exact_source_modified_asin_is_rejected(prepared):
    _,path,_=prepared
    w=load_workbook(path);s=w['精确匹配优先'];h={c.value:c.column for c in s[1]}
    s.cell(2,h['asin'],'B000000099');w.save(path)
    with pytest.raises(ValueError,match='Exact queue differs'):
        extract(path)


def test_wrong_expected_cohort_count_is_rejected(prepared):
    _,path,_=prepared
    with pytest.raises(ValueError,match='count'):
        extract_first_batch(path)


def test_unknown_stock_and_unknown_sales_are_not_zero(prepared):
    _,path,_=prepared
    p=extract(path)
    assert p['groups'][0]['fulfillable'] is None
    assert p['groups'][0]['sample_units_observed'] is None
    assert all(r['recent_units_observed'] is None for r in p['items'])
    assert p['summary']['sales_unknown_sku_count']==2


def test_readonly_connection_prohibits_mapping_writes(prepared):
    db,_,_=prepared
    with readonly_database(db) as con:
        with pytest.raises(sqlite3.OperationalError):
            con.execute("UPDATE sku SET family='Forged'")


def test_simulation_preserves_database_bytes_and_no_import_permission(prepared):
    db,path,dest=prepared
    p=extract(path);first=export_first_batch(p,dest/'first')['review_path']
    before=sha256(db.read_bytes()).hexdigest()
    report=simulate_first_batch_file(first,db,path)
    assert report['structure_status']=='READY'
    assert report['decision_status']=='PENDING'
    assert report['import_permitted'] is False
    assert before==sha256(db.read_bytes()).hexdigest()


def test_stop_row_identity_tampering_blocks_simulation(prepared):
    db,path,dest=prepared
    p=extract(path);first=export_first_batch(p,dest/'first')['review_path']
    w=load_workbook(first);s=w['停补标记'];h={c.value:c.column for c in s[1]}
    s.cell(2,h['sku'],'another-sku');w.save(first)
    report=simulate_first_batch_file(first,db,path)
    assert 'STOP_VISIBLE_MEMBER_IDENTITY_CHANGED' in report['global_blockers']
    assert report['import_permitted'] is False


@pytest.mark.parametrize('prepared',[{'a':'F','b':'G'}],indirect=True)
@pytest.mark.parametrize('remove_all',[False,True])
def test_removing_complete_groups_from_all_sheets_is_blocked(prepared,remove_all):
    db,path,dest=prepared
    p=extract_first_batch(path,expected_skus=2,expected_groups=2)
    first=export_first_batch(p,dest/'first')['review_path']
    removed=p['groups'][0]['group_id']
    w=load_workbook(first)
    for name in ('审核决定','SKU完整明细','停补标记'):
        sheet=w[name];headers={c.value:c.column for c in sheet[1]}
        for index in range(sheet.max_row,1,-1):
            if remove_all or sheet.cell(index,headers['group_id']).value==removed:
                sheet.delete_rows(index)
    w.save(first)
    report=simulate_first_batch_file(first,db,path)
    assert 'FIRST_COHORT_SCOPE_CHANGED' in report['global_blockers']
    assert report['structure_status']=='BLOCKED'
    assert report['import_permitted'] is False


def test_replacing_workbook_with_another_legal_snapshot_is_blocked(prepared):
    db,path,dest=prepared
    con=connect(db)
    con.execute("UPDATE discovery_sku SET last_seen_at='2030-01-01T00:00:00+00:00'")
    con.commit()
    newer=prepare_batch_review(con,dest/'newer')['review_path'];con.close()
    p=extract(newer)
    first=export_first_batch(p,dest/'first')['review_path']
    report=simulate_first_batch_file(first,db,path)
    assert 'FIRST_COHORT_SCOPE_CHANGED' in report['global_blockers']
    assert report['structure_status']=='BLOCKED'
    assert report['import_permitted'] is False
