from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from fba_replenishment.core import calculate, connect
from fba_replenishment.review import build_review_records, run_review
from fba_replenishment.spapi import SPAPIError
import pytest


def inventory(sku='J123-size 5', asin='B000000001'):
    return {'sellerSku': sku, 'asin': asin, 'fnSku': 'FN1', 'condition': 'NewItem',
            'inventoryDetails': {'fulfillableQuantity': 10, 'reservedQuantity': {'totalReservedQuantity': 2},
             'unfulfillableQuantity': {'totalUnfulfillableQuantity': 1}, 'inboundWorkingQuantity': 0,
             'inboundShippedQuantity': 3, 'inboundReceivingQuantity': 0}}


def test_candidates_never_approve_or_split_first_hyphen():
    rows = build_review_records([inventory('random-name-size 5')])
    assert rows[0]['candidate_family'] is None
    assert rows[0]['review_status'] == 'UNAPPROVED'
    rows = build_review_records([inventory()], source_candidates={
        'J123-size 5': {'candidate_family': 'User-defined', 'basis': 'USER_WORKBOOK_EXACT_SKU'}})
    assert rows[0]['candidate_family'] == 'User-defined'
    assert rows[0]['recommended_qty'] == 0


def test_catalog_403_produces_review_files_and_all_zero(tmp_path):
    class Client:
        marketplace = 'ATVPDKIKX0DER'
        def inventory_summaries(self): return [inventory()]
        def catalog_relationships(self, selected):
            raise SPAPIError('SP-API HTTP 403', 'SPAPI_HTTP_403')
    con = connect(tmp_path / 'review.db')
    asof = (datetime.now(timezone.utc).date()-timedelta(days=2)).isoformat()
    result = run_review(con, Client(), tmp_path, sample_summary={'asof': asof})
    assert result['summary']['calculable_sku_count'] == 0
    assert result['summary']['catalog_failed_count'] == 1
    assert con.execute('SELECT family FROM sku').fetchone()[0] is None
    assert calculate(con)[0]['recommended_qty'] == 0
    for path in [result['mapping_path'], result['stop_path'], result['blocked_plan_path']]:
        wb = load_workbook(path, read_only=True)
        assert len(wb.sheetnames) >= 4
        wb.close()
    con.close()


def test_partial_external_inventory_never_replaces_current_snapshot(tmp_path):
    class Client:
        marketplace = 'ATVPDKIKX0DER'
    con = connect(tmp_path / 'review.db')
    with pytest.raises(SPAPIError, match='completed'):
        run_review(con, Client(), tmp_path, inventory=[inventory()], inventory_audit={'complete': False, 'rows': 1})
    assert con.execute('SELECT COUNT(*) FROM stock').fetchone()[0] == 0
    con.close()


def test_daily_review_preserves_manual_approval_and_source_success(tmp_path):
    from fba_replenishment.discovery import approve_mapping
    class Client:
        marketplace = 'ATVPDKIKX0DER'
        def inventory_summaries(self): return [inventory()]
    con = connect(tmp_path / 'review.db')
    run_review(con, Client(), tmp_path, catalog_sample_size=0)
    approve_mapping(con, 'US', 'J123-size 5', 'J123', 'Manager')
    with con:
        con.execute("INSERT INTO sync_log(at,marketplace,kind,status,row_count,detail) VALUES('now','US','sales','SUCCESS',30,'Validated snapshot')")
    after = run_review(con, Client(), tmp_path, catalog_sample_size=0)
    assert con.execute('SELECT family FROM sku').fetchone()[0] == 'J123'
    assert con.execute('SELECT mapping_status FROM discovery_sku').fetchone()[0] == 'APPROVED'
    assert con.execute("SELECT COUNT(*) FROM data_quality_block WHERE code='MAPPING_UNAPPROVED'").fetchone()[0] == 0
    assert con.execute("SELECT status FROM sync_log WHERE kind='sales' ORDER BY id DESC LIMIT 1").fetchone()[0] == 'SUCCESS'
    assert after['summary']['mapping_pending_sku_count'] == 0
    assert after['summary']['mapping_unknown_sku_count'] == 0
    con.close()
