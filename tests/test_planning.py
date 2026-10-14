import csv
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from fba_replenishment.core import calculate, connect, import_csv, resume, stop
from fba_replenishment.export import export


@pytest.fixture
def con(tmp_path):
    c = connect(tmp_path / 'plan.db')
    today = datetime.now(timezone.utc).date()
    with c:
        for market in ['US', 'CA']:
            for sku, family in [('a', 'family'), ('b', 'family'), ('c', 'other')]:
                c.execute('INSERT INTO sku(marketplace,sku,family) VALUES(?,?,?)', (market, sku, family))
                c.execute('INSERT INTO stock VALUES(?,?,?,?,?,?)', (market, sku, 0, 100, 100, datetime.now(timezone.utc).isoformat()))
                c.executemany('INSERT INTO sales VALUES(?,?,?,?)', [(market, sku, (today - timedelta(days=i)).isoformat(), 1) for i in range(1, 31)])
    yield c
    c.close()


def bysku(con, market='US', **kwargs):
    return {r['sku']: r for r in calculate(con, marketplace=market, **kwargs)}


def test_multiple_stops_resume_marketplace_and_new_child(con):
    stop(con, 'US', 'a', 'stop a')
    stop(con, 'US', 'b', 'stop b')
    assert all(bysku(con)[s]['recommended_qty'] == 0 for s in ['a', 'b'])
    assert bysku(con, 'CA')['b']['recommended_qty'] > 0
    assert bysku(con)['c']['recommended_qty'] > 0
    with pytest.raises(ValueError):
        resume(con, 'US', 'a', ' ')
    resume(con, 'US', 'a', 'manager')
    assert bysku(con)['a']['status'] == 'STOPPED'
    con.execute("INSERT INTO sku(marketplace,sku,family) VALUES('US','new','family')")
    con.commit()
    assert bysku(con)['new']['status'] == 'STOPPED'
    resume(con, 'US', 'b', 'manager')
    assert bysku(con)['a']['recommended_qty'] > 0
    assert bysku(con)['new']['status'] == 'MISSING_STOCK'


def test_stop_family_survives_remapping_and_reopening(con, tmp_path):
    stop(con, 'US', 'a', 'discontinued')
    file = tmp_path / 'mapping.csv'
    file.write_text('sku,family\na,other\n')
    import_csv(con, 'sku', file, 'US')
    reopened = connect(tmp_path / 'plan.db')
    assert bysku(reopened)['b']['status'] == 'STOPPED'
    assert bysku(reopened)['c']['status'] == 'STOPPED'
    reopened.close()
    resume(con, 'US', 'a', 'manager')
    assert bysku(con)['b']['status'] == 'ACTIVE'


def test_incomplete_unknown_stock_and_stale(con):
    con.execute("DELETE FROM sales WHERE marketplace='US' AND sku='a' AND day=(SELECT MIN(day) FROM sales)")
    con.execute("DELETE FROM stock WHERE marketplace='US' AND sku='b'")
    con.execute("UPDATE stock SET observed_at='2020-01-01T00:00:00+00:00' WHERE marketplace='US' AND sku='c'")
    con.commit()
    result = bysku(con)
    assert result['a']['status'] == 'INCOMPLETE_SALES'
    assert result['b']['status'] == 'MISSING_STOCK'
    assert result['c']['status'] == 'STALE_STOCK'
    assert all(r['recommended_qty'] == 0 for r in result.values())


def test_moq_casepack_and_inventory_not_double_subtracted(con):
    result = bysku(con, lead_days=0, target_days=1, moq=11, case_pack=6)
    assert result['a']['recommended_qty'] == 12
    # Reserved and unfulfillable are already separate from fulfillable.
    assert bysku(con, lead_days=0, target_days=10)['a']['recommended_qty'] == 10


def test_inbound_terminal_unknown_eta_and_received_exclusion(con):
    con.executemany('INSERT INTO inbound VALUES(?,?,?,?,?,?)', [
        ('US', 'a', 'closed', 1000, '2000-01-01', 'CLOSED'),
        ('US', 'a', 'cancel', 1000, '', 'CANCELLED'),
        ('US', 'b', 'active', 12, '', 'SHIPPED'),
        ('US', 'c', 'on-time', 1000, '2000-01-01', 'IN_TRANSIT'),
    ])
    con.commit()
    result = bysku(con)
    assert result['a']['recommended_qty'] == 110
    assert result['b']['status'] == 'INBOUND_REVIEW_REQUIRED'
    assert result['c']['recommended_qty'] == 0


def test_csv_atomic_upsert_and_missing_stock_rejected(con, tmp_path):
    p = tmp_path / 'sales.csv'
    day = datetime.now(timezone.utc).date() - timedelta(days=1)
    p.write_text(f'sku,day,units\na,{day},10\n')
    import_csv(con, 'sales', p, 'US')
    import_csv(con, 'sales', p, 'US')
    assert con.execute("SELECT COUNT(*) FROM sales WHERE marketplace='US' AND sku='a'").fetchone()[0] == 30
    assert bysku(con)['a']['sales_30d'] == 39
    p.write_text('sku,fulfillable\na,10\nb,\n')
    with pytest.raises(ValueError):
        import_csv(con, 'stock', p, 'US')
    assert con.execute("SELECT fulfillable FROM stock WHERE marketplace='US' AND sku='a'").fetchone()[0] == 0


def test_excel_contains_inputs_alerts_parameters_and_literal_strings(con, tmp_path):
    stop(con, 'US', 'a', '=HYPERLINK("malicious","click")')
    # Trigger SKU with formula-like text must remain a literal in workbook.
    con.execute("INSERT INTO sku(marketplace,sku,family) VALUES('US','=1+1','family')")
    con.commit()
    path = tmp_path / 'plan.xlsx'
    export(calculate(con), path, con, 'US', {'max_stock_age_days': 2})
    wb = load_workbook(path)
    assert set(wb.sheetnames) == {'Replenishment', 'Sales', 'Inventory', 'Inbound', 'Orders', 'Alerts', 'Sync Log', 'Audit', 'Parameters'}
    assert wb['Sales'].max_row == 91
    assert wb['Inventory'].max_row == 4
    assert wb['Alerts'].max_row == 4
    for row in wb['Replenishment'].iter_rows(min_row=2):
        if row[1].value == '=1+1':
            assert row[1].data_type == 's'
            assert row[17].value == 0
    assert wb['Audit'].cell(2, 6).data_type == 's'


def test_disabled_child_stops_whole_family_and_requires_approved_resume(con):
    con.execute("UPDATE sku SET enabled=0 WHERE marketplace='US' AND sku='a'")
    con.commit()
    result = bysku(con)
    assert result['a']['status'] == result['b']['status'] == 'STOPPED'
    assert result['a']['recommended_qty'] == result['b']['recommended_qty'] == 0
    assert bysku(con, 'CA')['b']['status'] == 'ACTIVE'
    with pytest.raises(ValueError, match='approver'):
        resume(con, 'US', 'a', '')
    resume(con, 'US', 'a', 'manager')
    assert bysku(con)['b']['status'] == 'ACTIVE'
