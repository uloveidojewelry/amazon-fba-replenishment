from openpyxl import load_workbook
import pytest

from fba_replenishment.batch_review import prepare_batch_review, import_batch_review
from fba_replenishment.batch_review import import_confirmed_stops
from fba_replenishment.core import connect
from fba_replenishment.discovery import discover

US='ATVPDKIKX0DER'


def inventory(sku):
    return {'sellerSku':sku,'asin':'B000000001','fnSku':'FN-'+sku,'condition':'NewItem',
            'inventoryDetails':{'fulfillableQuantity':1,'reservedQuantity':{'totalReservedQuantity':0},
            'unfulfillableQuantity':{'totalUnfulfillableQuantity':0},'inboundWorkingQuantity':0,
            'inboundShippedQuantity':0,'inboundReceivingQuantity':0}}


@pytest.fixture
def prepared(tmp_path):
    con=connect(tmp_path/'review.db')
    source={s:{'candidate_family':'F','basis':'USER_WORKBOOK_EXACT_SKU'} for s in ('a','b')}
    discover(con,[inventory('a'),inventory('b')],'US',US,source_candidates=source)
    paths=prepare_batch_review(con,tmp_path)
    yield con,paths,source
    con.close()


def approve_file(path, action='APPROVED'):
    wb=load_workbook(path)
    s=wb['批量审核决定'];h={c.value:c.column for c in s[1]}
    for n in range(2,s.max_row+1):
        s.cell(n,h['decision'],action)
        s.cell(n,h['approved_family'],'F')
        s.cell(n,h['reviewer'],'manager')
        s.cell(n,h['reason'],'Explicitly reviewed the exact visible SKU set')
    wb.save(path)


def test_unedited_workbook_never_approves(prepared):
    con,paths,_=prepared
    result=import_batch_review(con,paths['review_path'],'manager',confirm=True)
    assert result['status']=='NO_APPROVED_DECISIONS'
    assert con.execute('SELECT COUNT(*) FROM batch_review_decision').fetchone()[0]==0
    assert con.execute("SELECT COUNT(*) FROM discovery_sku WHERE family IS NOT NULL").fetchone()[0]==0


def test_explicit_family_workbook_approves_only_frozen_members(prepared):
    con,paths,source=prepared
    approve_file(paths['review_path'])
    discover(con,[inventory('a'),inventory('b'),inventory('new-sibling')],'US',US,source_candidates=source)
    result=import_batch_review(con,paths['review_path'],'manager',confirm=True)
    assert result['approved_skus']==2
    assert con.execute('SELECT family FROM discovery_sku WHERE sku=?',('new-sibling',)).fetchone()[0] is None
    assert con.execute('SELECT COUNT(*) FROM batch_review_audit WHERE action=?',('BATCH_MAPPING_APPROVED',)).fetchone()[0]==2


def test_removed_visible_member_rejects_entire_batch(prepared):
    con,paths,_=prepared
    approve_file(paths['review_path'])
    wb=load_workbook(paths['review_path']);wb['SKU明确集合'].delete_rows(2);wb.save(paths['review_path'])
    with pytest.raises(ValueError,match='Visible workbook membership'):
        import_batch_review(con,paths['review_path'],'manager',confirm=True)
    assert con.execute('SELECT COUNT(*) FROM batch_review_decision').fetchone()[0]==0
    assert con.execute("SELECT COUNT(*) FROM discovery_sku WHERE family IS NOT NULL").fetchone()[0]==0


def test_reject_decision_is_persistent_and_never_approves(prepared):
    con,paths,_=prepared
    approve_file(paths['review_path'],'REJECTED')
    result=import_batch_review(con,paths['review_path'],'manager',confirm=True)
    assert result['rejected_groups']==1
    assert con.execute('SELECT decision_type FROM batch_review_decision').fetchone()[0]=='REJECTED'
    assert con.execute("SELECT COUNT(*) FROM discovery_sku WHERE family IS NOT NULL").fetchone()[0]==0


def test_workbook_requires_matching_explicit_reviewer(prepared):
    con,paths,_=prepared
    approve_file(paths['review_path'])
    with pytest.raises(ValueError,match='identity'):
        import_batch_review(con,paths['review_path'],'different person',confirm=True)
    with pytest.raises(ValueError,match='confirmation'):
        import_batch_review(con,paths['review_path'],'manager')


def test_multiple_workbook_decisions_roll_back_if_later_scope_is_stale(tmp_path):
    con=connect(tmp_path/'multiple.db')
    source={s:{'candidate_family':s.upper(),'basis':'USER_WORKBOOK_EXACT_SKU'} for s in ('a','b')}
    discover(con,[inventory('a'),inventory('b')],'US',US,source_candidates=source)
    paths=prepare_batch_review(con,tmp_path)
    wb=load_workbook(paths['review_path']);s=wb['批量审核决定'];h={c.value:c.column for c in s[1]}
    for n in range(2,s.max_row+1):
        s.cell(n,h['decision'],'APPROVED');s.cell(n,h['approved_family'],s.cell(n,h['candidate_family']).value)
        s.cell(n,h['reviewer'],'manager');s.cell(n,h['reason'],'Reviewed')
    wb.save(paths['review_path'])
    with con:con.execute("UPDATE discovery_sku SET mapping_basis='CHANGED' WHERE sku='b'")
    with pytest.raises(ValueError,match='Stale'):
        import_batch_review(con,paths['review_path'],'manager',confirm=True)
    assert con.execute('SELECT COUNT(*) FROM batch_review_decision').fetchone()[0]==0
    assert con.execute("SELECT COUNT(*) FROM discovery_sku WHERE family IS NOT NULL").fetchone()[0]==0
    con.close()


def test_stop_file_cannot_mutate_another_selected_namespace(prepared,tmp_path):
    con,_,_=prepared
    path=tmp_path/'stops.csv'
    path.write_text('marketplace,sku,family,stop,reason,effective_date,review_status,reviewer\nCA,a,F,TRUE,Review,2026-10-14,APPROVED,manager\n')
    with pytest.raises(ValueError,match='different marketplace'):
        import_confirmed_stops(con,path,'manager',marketplace='US')
    assert con.execute('SELECT COUNT(*) FROM stop_flag').fetchone()[0]==0


def test_not_yet_effective_stop_does_not_write_during_readonly_planning(tmp_path):
    import sqlite3
    from datetime import datetime,timedelta,timezone
    from fba_replenishment.batch_decisions import import_stop_list
    from fba_replenishment.core import calculate
    from fba_replenishment.discovery import approve_mapping
    db=tmp_path/'future.db'
    con=connect(db)
    discover(con,[inventory('a')],'US',US)
    approve_mapping(con,'US','a','F','manager')
    future=(datetime.now(timezone.utc).date()+timedelta(days=1)).isoformat()
    path=tmp_path/'future.csv'
    path.write_text('marketplace,sku,family,stop,reason,effective_date,review_status,reviewer\nUS,a,F,TRUE,Review,'+future+',APPROVED,manager\n')
    import_stop_list(con,path,'manager')
    assert con.execute('SELECT COUNT(*) FROM stop_flag').fetchone()[0]==0
    before=con.execute('SELECT COUNT(*) FROM batch_review_audit').fetchone()[0]
    con.close()
    reader=sqlite3.connect('file:'+str(db)+'?mode=ro',uri=True);reader.row_factory=sqlite3.Row
    try:
        assert calculate(reader)[0]['recommended_qty']==0
        assert reader.execute('SELECT COUNT(*) FROM batch_review_audit').fetchone()[0]==before
    finally:reader.close()


def test_mapping_approval_does_not_turn_missing_inbound_evidence_into_zero(tmp_path):
    from datetime import datetime,timedelta,timezone
    from fba_replenishment.core import calculate
    from fba_replenishment.discovery import approve_mapping
    con=connect(tmp_path/'planning.db')
    discover(con,[inventory('a')],'US',US)
    approve_mapping(con,'US','a','F','manager')
    asof=datetime.now(timezone.utc).date()-timedelta(days=2)
    with con:
        con.execute('INSERT INTO stock VALUES(?,?,?,?,?,?)',('US','a',0,0,0,datetime.now(timezone.utc).isoformat()))
        con.executemany('INSERT INTO sales VALUES(?,?,?,?)',[('US','a',(asof-timedelta(days=i)).isoformat(),1) for i in range(30)])
    row=calculate(con,asof=asof.isoformat())[0]
    assert row['status']=='INCOMPLETE_INBOUND' and row['recommended_qty']==0
    con.close()
