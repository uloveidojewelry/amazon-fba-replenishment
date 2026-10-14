from copy import deepcopy
import sqlite3

from openpyxl import load_workbook
import pytest

from fba_replenishment.core import connect
from fba_replenishment.review import REVIEW_SCHEMA
from fba_replenishment.sku_export import export_sku_report


def read_rows(ws):
    values = list(ws.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def plan(sku='sku-a', **changes):
    result = {
        'marketplace': 'US', 'sku': sku, 'asin': 'ASIN-A', 'family': None,
        'status': 'ACTIVE', 'replenishment_status': 'ACTIVE',
        'sales_status': 'SOLD', 'fulfillable': 5, 'reserved': 2, 'unfulfillable': 1,
        'inbound_working': 0, 'inbound_shipped': 20, 'inbound_receiving': 0,
        'inventory_inbound': 20, 'eligible_inbound': 20, 'unknown_eta_inbound': 0,
        'inventory_complete': True, 'inbound_complete': True,
        'recommended_qty': 50, 'daily_sales': 1, 'coverage_days': 5,
        'stockout_risk': 'STOCKOUT_BEFORE_ARRIVAL', 'parameter_status': 'SCENARIO_ONLY',
        'parameters_confirmed': False, 'lead_days': 35, 'target_days': 75,
        'safety_days': 7, 'moq': 1, 'case_pack': 1,
        'inbound_rule': 'CONFIRMED_ETA_ONLY', 'sales_timezone': 'America/Los_Angeles',
        'asof': '2026-10-13', 'window_start': '2026-07-16',
        'stock_observed_at': '2026-10-14T01:00:00+00:00', 'stock_age_days': 1,
        'anomaly_codes': [], 'anomalies': [],
    }
    for days in (7, 30, 60, 90):
        result.update({f'sales_{days}d': days, f'observed_sales_{days}d': days,
                       f'observed_sales_days_{days}d': days,
                       f'sales_complete_{days}d': True, f'daily_sales_{days}d': 1})
    result.update(changes)
    return result


def export_rows(tmp_path, rows, con=None, settings=None, collection_audit=None):
    path = tmp_path / 'report.xlsx'
    summary = export_sku_report(rows, path, con, 'US', settings or {}, collection_audit)
    return load_workbook(path), summary


def test_independent_report_has_required_sheets_and_does_not_require_family(tmp_path):
    wb, summary = export_rows(tmp_path, [plan()])
    assert {'SKU补货总表', '7_30_60_90天销量', 'FBA库存明细', '在途库存明细',
            '建议补货数量', '缺货风险', '数据异常', '已停补SKU', '计算参数',
            '采集审计', '报表统计', '入仓库存核对'} <= set(wb.sheetnames)
    row = read_rows(wb['SKU补货总表'])[0]
    assert row['family'] is None
    assert row['recommended_qty'] == 50
    assert row['daily_sales'] == 1
    assert summary['sku_count'] == summary['sold_sku_count'] == 1
    assert summary['family_mapping_required'] is summary['family_stop_applied'] is False
    assert summary['data_collection_completeness'] == 1


def test_windows_are_complete_calendar_totals_and_30_day_basis(tmp_path):
    item = plan(sales_7d=21, sales_30d=60, sales_60d=60, sales_90d=360)
    wb, summary = export_rows(tmp_path, [item])
    row = read_rows(wb['7_30_60_90天销量'])[0]
    assert [row[f'daily_sales_{days}d'] for days in (7, 30, 60, 90)] == [3, 2, 1, 4]
    assert read_rows(wb['SKU补货总表'])[0]['daily_sales'] == 2
    assert summary['recommendation_basis_days'] == 30


def test_unknown_sales_are_blank_even_if_source_claims_zero_or_sync_success(tmp_path):
    item = plan(sales_30d=0, sales_90d=0, sales_complete_30d=False,
                sales_complete_90d=False, observed_sales_days_30d=29,
                observed_sales_days_90d=89, sales_status='NO_SALES')
    wb, summary = export_rows(tmp_path, [item], collection_audit={
        'summary': {'collection_status': 'SUCCESS'}, 'audit': [{'status': 'SUCCESS'}]})
    row = read_rows(wb['SKU补货总表'])[0]
    assert row['sales_30d'] is row['sales_90d'] is row['daily_sales'] is None
    assert row['recommended_qty'] is None
    assert row['sales_status'] == 'UNKNOWN'
    assert summary['no_sales_sku_count'] == 0
    assert summary['unknown_data_sku_count'] == 1
    assert summary['data_collection_completeness'] == 0


def test_certified_no_sales_is_distinct_from_unknown_and_sold(tmp_path):
    zero = plan('zero', recommended_qty=0, **{f'sales_{d}d': 0 for d in (7, 30, 60, 90)})
    unknown = plan('unknown', sales_complete_90d=False)
    wb, summary = export_rows(tmp_path, [plan('sold'), zero, unknown])
    assert summary['sku_count'] == 3
    assert summary['sold_sku_count'] == summary['no_sales_sku_count'] == summary['sales_unknown_sku_count'] == 1
    by_sku = {r['sku']: r for r in read_rows(wb['SKU补货总表'])}
    assert by_sku['zero']['sales_90d'] == 0
    assert by_sku['unknown']['sales_90d'] is None


def test_direct_stop_does_not_propagate_to_same_family(tmp_path):
    rows = [plan('stopped', family='J001', status='STOPPED', recommended_qty=999),
            plan('same-family', family='J001'), plan('other-family', family='J002')]
    original = deepcopy(rows)
    wb, summary = export_rows(tmp_path, rows)
    items = {r['sku']: r for r in read_rows(wb['SKU补货总表'])}
    assert items['stopped']['recommended_qty'] == 0
    assert items['same-family']['recommended_qty'] == items['other-family']['recommended_qty'] == 50
    assert rows == original
    assert summary['stopped_sku_count'] == 1
    assert [r['sku'] for r in read_rows(wb['已停补SKU'])] == ['stopped']


def test_stopped_unknown_data_keeps_explicit_zero(tmp_path):
    wb, _ = export_rows(tmp_path, [plan(status='STOPPED', recommended_qty=None,
                                      reserved=None, sales_complete_30d=False)])
    assert read_rows(wb['SKU补货总表'])[0]['recommended_qty'] == 0


def test_unconfirmed_parameters_are_prominent_scenario_not_purchase(tmp_path):
    wb, summary = export_rows(tmp_path, [plan()], settings={'target_days': 75, 'access_token': 'SECRET'})
    row = read_rows(wb['建议补货数量'])[0]
    assert row['parameter_status'] == 'SCENARIO_ONLY'
    assert '不得作为正式采购依据' in row['recommendation_label']
    parameters = {r['parameter']: r['value'] for r in read_rows(wb['计算参数'])}
    assert parameters['target_days'] == 75
    assert parameters['family_stop_applied'] is False
    assert summary['scenario_recommended_units'] == 50
    assert summary['confirmed_recommended_units'] == 0


def test_confirmation_requires_both_flag_and_status(tmp_path):
    wb, summary = export_rows(tmp_path, [plan(parameter_status='CONFIRMED', parameters_confirmed=False)])
    assert read_rows(wb['SKU补货总表'])[0]['parameter_status'] == 'SCENARIO_ONLY'
    wb, summary = export_rows(tmp_path, [plan(parameter_status='CONFIRMED', parameters_confirmed=True)])
    assert summary['parameter_status'] == 'CONFIRMED'
    assert summary['confirmed_recommended_units'] == 50


@pytest.mark.parametrize('bad_value', [None, -1, True, float('nan')])
def test_unknown_or_invalid_inventory_is_not_zero(tmp_path, bad_value):
    wb, summary = export_rows(tmp_path, [plan(reserved=bad_value)])
    row = read_rows(wb['FBA库存明细'])[0]
    assert row['reserved'] is None
    assert row['inventory_complete'] is False
    assert read_rows(wb['SKU补货总表'])[0]['recommended_qty'] is None
    assert summary['inventory_completeness'] == 0
    assert summary['anomaly_sku_count'] == 1


def test_inventory_summary_and_shipment_detail_do_not_double_deduct(tmp_path):
    con = connect(tmp_path / 'data.db')
    con.executescript(REVIEW_SCHEMA)
    con.execute('INSERT INTO inbound VALUES(?,?,?,?,?,?)', ('US', 'sku-a', 'FBA-1', 20, '2026-10-20', 'SHIPPED'))
    con.commit()
    wb, summary = export_rows(tmp_path, [plan()], con)
    reconcile = read_rows(wb['入仓库存核对'])[0]
    assert reconcile['inventory_inbound'] == reconcile['eligible_inbound'] == 20
    assert read_rows(wb['在途库存明细'])[0]['units'] == 20
    assert read_rows(wb['建议补货数量'])[0]['recommended_qty'] == 50
    assert '重复' in wb['FBA库存明细']['F1'].comment.text
    con.close()


def test_inventory_and_inbound_are_batched_read_only_and_other_site_filtered(tmp_path):
    db = tmp_path / 'data.db'
    con = connect(db)
    con.executescript(REVIEW_SCHEMA)
    for sku in ('a', 'b'):
        con.execute('INSERT INTO inventory_details VALUES(?,?,?,?,?,?)', ('US', sku, 0, 20, 0, '2026-10-14T01:00:00Z'))
        con.execute('INSERT INTO inbound VALUES(?,?,?,?,?,?)', ('US', sku, 'FBA-' + sku, 20, '', 'SHIPPED'))
    con.execute('INSERT INTO inbound VALUES(?,?,?,?,?,?)', ('CA', 'a', 'FBA-CA', 99, '', 'SHIPPED'))
    con.execute('INSERT INTO inbound VALUES(?,?,?,?,?,?)', ('US', 'outside-cohort', 'FBA-OTHER', 99, '', 'SHIPPED'))
    con.commit()
    con.close()
    original = db.read_bytes()
    readonly = sqlite3.connect(db.as_uri() + '?mode=ro', uri=True)
    statements = []
    readonly.set_trace_callback(statements.append)
    rows = [plan(sku) for sku in ('a', 'b')]
    for row in rows:
        for field in ('inbound_working', 'inbound_shipped', 'inbound_receiving'):
            del row[field]
    wb, summary = export_rows(tmp_path, rows, readonly)
    assert summary['inventory_complete_sku_count'] == 2
    assert {r['sku'] for r in read_rows(wb['在途库存明细'])} == {'a', 'b'}
    assert all(r['eta'] is None for r in read_rows(wb['在途库存明细']))
    assert len([s for s in statements if 'FROM inbound ' in s]) == 1
    assert len([s for s in statements if 'FROM inventory_details ' in s]) == 1
    assert db.read_bytes() == original
    readonly.close()


def test_unknown_inventory_does_not_get_filled_from_a_stale_database(tmp_path):
    con = connect(tmp_path / 'data.db')
    con.executescript(REVIEW_SCHEMA)
    con.execute('INSERT INTO inventory_details VALUES(?,?,?,?,?,?)', ('US', 'sku-a', 0, 20, 0, '2026-10-01T00:00:00Z'))
    con.commit()
    wb, summary = export_rows(tmp_path, [plan(inbound_shipped=None)], con)
    assert read_rows(wb['FBA库存明细'])[0]['inbound_shipped'] is None
    assert summary['inventory_complete_sku_count'] == 0
    con.close()


def test_absent_inbound_fields_do_not_get_certified_from_a_different_snapshot(tmp_path):
    con = connect(tmp_path / 'data.db')
    con.executescript(REVIEW_SCHEMA)
    con.execute('INSERT INTO inventory_details VALUES(?,?,?,?,?,?)', ('US', 'sku-a', 0, 20, 0, '2026-10-01T01:00:00Z'))
    con.commit()
    item = plan()
    for field in ('inbound_working', 'inbound_shipped', 'inbound_receiving'):
        del item[field]
    wb, summary = export_rows(tmp_path, [item], con)
    assert read_rows(wb['FBA库存明细'])[0]['inbound_shipped'] is None
    assert summary['inventory_complete_sku_count'] == 0
    con.close()


def test_unavailable_inbound_collection_blocks_and_cannot_certify_completeness(tmp_path):
    wb, summary = export_rows(tmp_path, [plan(inbound_complete=False)], collection_audit={
        'summary': {'collection_status': 'SUCCESS'}})
    assert summary['data_collection_completeness'] == 0
    assert summary['inbound_completeness'] == 0
    assert read_rows(wb['SKU补货总表'])[0]['recommended_qty'] is None
    assert 'INBOUND_COLLECTION_INCOMPLETE' in {r['code'] for r in read_rows(wb['数据异常'])}


def test_export_rejects_duplicate_and_other_marketplace_before_saving(tmp_path):
    for rows in ([plan(), plan()], [plan(marketplace='CA')]):
        with pytest.raises(ValueError):
            export_sku_report(rows, tmp_path / 'report.xlsx', None, 'US', {})
        assert not (tmp_path / 'report.xlsx').exists()


def test_external_strings_literal_sensitive_fields_and_api_bodies_are_omitted(tmp_path):
    formula = '=HYPERLINK("https://example.invalid","SKU")'
    item = plan(formula, family=formula, access_token='TOKEN_SENTINEL',
                buyer_name='CUSTOMER_SENTINEL', order_id='ORDER_SENTINEL',
                raw_response='RAW_SENTINEL', anomalies=[{
                    'code': 'SAFE_CODE', 'reason': 'Known inventory conflict', 'blocking': True,
                    'source': 'Inventory validation', 'buyer_name': 'CUSTOMER_SENTINEL',
                    'authorization': 'TOKEN_SENTINEL', 'raw_payload': 'RAW_SENTINEL'}])
    wb, summary = export_rows(tmp_path, [item], settings={'refresh_token': 'TOKEN_SENTINEL'}, collection_audit={
        'summary': {'orders_collected': 22, 'access_token': 'TOKEN_SENTINEL'},
        'audit': [{'kind': 'orders', 'orders': 22, 'complete': True,
                   'begin': '2026-10-01T00:00:00Z', 'issues': {'SAFE_CODE': 1},
                   'order_id': 'ORDER_SENTINEL', 'raw_response': 'RAW_SENTINEL',
                   'detail': 'TOKEN_SENTINEL', 'client_secret': 'TOKEN_SENTINEL'}]})
    assert summary['orders_collected'] == 22
    for ws in wb:
        for row in ws:
            for cell in row:
                assert cell.value not in ('TOKEN_SENTINEL', 'CUSTOMER_SENTINEL', 'ORDER_SENTINEL', 'RAW_SENTINEL')
                if cell.value == formula:
                    assert cell.data_type == 's'
    assert read_rows(wb['采集审计'])[0]['orders'] == 22


def test_completeness_is_sku_intersection_not_just_sales_or_snapshot_success(tmp_path):
    wb, summary = export_rows(tmp_path, [plan('complete'),
        plan('missing-sales', sales_complete_90d=False),
        plan('missing-stock', reserved=None), plan('missing-inbound', inbound_complete=False)],
        collection_audit={'summary': {'collection_status': 'SUCCESS'}})
    assert summary['sku_count'] == 4
    assert summary['sales_completeness'] == summary['inventory_completeness'] == summary['inbound_completeness'] == .75
    assert summary['data_collection_completeness'] == .25
    assert summary['unknown_data_sku_count'] == 3
    assert summary['anomaly_sku_count'] == 3


def test_empty_cohort_completeness_is_unknown_not_one(tmp_path):
    wb, summary = export_rows(tmp_path, [])
    assert summary['sku_count'] == 0
    assert summary['data_collection_completeness'] is None
    assert summary['recommended_units'] == 0


def test_complete_quantities_do_not_mask_identity_conflict(tmp_path):
    item = plan(data_complete=False, status='DATA_REVIEW_REQUIRED', recommended_qty=None,
                anomalies=[{'code': 'SKU_ASIN_IDENTITY_CONFLICT', 'reason': 'SKU ASIN conflict',
                            'blocking': True, 'source': 'Stored identity validation'}])
    wb, summary = export_rows(tmp_path, [item])
    assert summary['sales_completeness'] == summary['inventory_completeness'] == summary['inbound_completeness'] == 1
    assert summary['data_collection_completeness'] == 0
    assert summary['unknown_data_sku_count'] == summary['anomaly_sku_count'] == 1
    assert read_rows(wb['SKU补货总表'])[0]['recommended_qty'] is None


def test_data_sources_are_visible_and_no_recent_demand_is_not_a_stockout(tmp_path):
    item = plan(sales_source='Orders API / AFN', inventory_source='FBA Inventory API',
                inbound_source='Fulfillment Inbound v0', stockout_risk='NO_RECENT_DEMAND')
    wb, _ = export_rows(tmp_path, [item])
    assert read_rows(wb['SKU补货总表'])[0]['sales_source'] == 'Orders API / AFN'
    assert read_rows(wb['7_30_60_90天销量'])[0]['sales_source'] == 'Orders API / AFN'
    assert read_rows(wb['FBA库存明细'])[0]['inventory_source'] == 'FBA Inventory API'
    assert not read_rows(wb['缺货风险'])


def exception_table(con):
    con.execute('''CREATE TABLE sku_inbound_exception (
        marketplace TEXT NOT NULL, sku TEXT NOT NULL, shipment_id TEXT NOT NULL,
        fnsku TEXT NOT NULL DEFAULT '',
        status TEXT, quantity_shipped INTEGER, quantity_received INTEGER,
        remaining INTEGER, code TEXT, PRIMARY KEY(marketplace,sku,shipment_id,fnsku))''')


def test_received_exceeds_shipped_preserves_evidence_and_unknown_remaining(tmp_path):
    db = tmp_path / 'data.db'
    con = connect(db)
    exception_table(con)
    # A previously clamped balance cannot mask contradictory original evidence.
    con.execute('INSERT INTO inbound VALUES(?,?,?,?,?,?)', ('US', 'conflict', 'FBA-C', 0, '', 'RECEIVING'))
    con.execute('INSERT INTO inbound VALUES(?,?,?,?,?,?)', ('US', 'normal', 'FBA-N', 20, '2026-10-20', 'SHIPPED'))
    con.execute('''INSERT INTO sku_inbound_exception
        (marketplace,sku,shipment_id,status,quantity_shipped,quantity_received,remaining,code)
        VALUES(?,?,?,?,?,?,?,?)''',
                ('US', 'conflict', 'FBA-C', 'RECEIVING', 4, 5, None, 'INBOUND_QUANTITY_CONFLICT'))
    con.commit()
    con.close()
    original = db.read_bytes()
    readonly = sqlite3.connect(db.as_uri() + '?mode=ro', uri=True)
    wb, summary = export_rows(tmp_path, [plan('normal'), plan('conflict')], readonly)
    detail = {r['sku']: r for r in read_rows(wb['在途库存明细'])}
    assert len(detail) == 2
    assert detail['conflict']['quantity_shipped'] == 4
    assert detail['conflict']['quantity_received'] == 5
    assert detail['conflict']['remaining'] is detail['conflict']['units'] is None
    assert detail['conflict']['code'] == 'INBOUND_QUANTITY_CONFLICT'
    assert detail['normal']['remaining'] == detail['normal']['units'] == 20
    plans = {r['sku']: r for r in read_rows(wb['SKU补货总表'])}
    assert plans['conflict']['recommended_qty'] is None
    assert plans['conflict']['eligible_inbound'] is None
    assert plans['conflict']['inbound_complete'] is False
    assert plans['normal']['recommended_qty'] == 50
    assert summary['data_collection_completeness'] == .5
    assert summary['unknown_data_sku_count'] == summary['anomaly_sku_count'] == 1
    assert db.read_bytes() == original
    readonly.close()


def test_exception_details_keep_missing_values_blank_and_filter_marketplace_and_cohort(tmp_path):
    con = connect(tmp_path / 'data.db')
    exception_table(con)
    con.executemany('''INSERT INTO sku_inbound_exception
        (marketplace,sku,shipment_id,status,quantity_shipped,quantity_received,remaining,code)
        VALUES(?,?,?,?,?,?,?,?)''', [
        ('US', 'sku-a', 'FBA-A', 'RECEIVING', None, 5, None, 'INBOUND_QUANTITY_CONFLICT'),
        ('CA', 'sku-a', 'FBA-CA', 'RECEIVING', 4, 5, None, 'INBOUND_QUANTITY_CONFLICT'),
        ('US', 'outside', 'FBA-X', 'RECEIVING', 4, 5, None, 'INBOUND_QUANTITY_CONFLICT')])
    con.commit()
    wb, summary = export_rows(tmp_path, [plan()], con)
    detail = read_rows(wb['在途库存明细'])
    assert len(detail) == 1
    assert detail[0]['shipment_id'] == 'FBA-A'
    assert detail[0]['quantity_shipped'] is detail[0]['remaining'] is None
    assert detail[0]['quantity_received'] == 5
    assert summary['sku_count'] == summary['unknown_data_sku_count'] == 1
    con.close()


def test_direct_stop_remains_zero_with_inbound_quantity_conflict(tmp_path):
    con = connect(tmp_path / 'data.db')
    exception_table(con)
    con.execute('''INSERT INTO sku_inbound_exception
        (marketplace,sku,shipment_id,status,quantity_shipped,quantity_received,remaining,code)
        VALUES(?,?,?,?,?,?,?,?)''',
                ('US', 'sku-a', 'FBA-C', 'RECEIVING', 4, 5, None, 'INBOUND_QUANTITY_CONFLICT'))
    con.commit()
    wb, summary = export_rows(tmp_path, [plan(status='STOPPED', stop_reason='Manual direct stop')], con)
    stopped = read_rows(wb['已停补SKU'])[0]
    assert stopped['recommended_qty'] == 0
    assert stopped['stop_reason'] == 'Manual direct stop'
    assert read_rows(wb['在途库存明细'])[0]['remaining'] is None
    assert summary['unknown_data_sku_count'] == summary['stopped_sku_count'] == 1
    con.close()


def test_unknown_data_count_is_union_of_missing_data_missing_result_and_unknown_eta(tmp_path):
    rows = [plan('scenario-valid'),
            plan('missing-data', data_complete=False, sales_complete_90d=False),
            plan('eta-unknown', unknown_eta_inbound=20, inbound_rule='exclude'),
            plan('result-blocked', recommended_qty=None, status='DATA_REVIEW_REQUIRED'),
            plan('known-zero', recommended_qty=0),
            plan('stopped', status='STOPPED', recommended_qty=0)]
    wb, summary = export_rows(tmp_path, rows)
    assert summary['unknown_data_sku_count'] == 3
    assert summary['eta_unknown_sku_count'] == 1
    assert summary['sales_unknown_sku_count'] == 1
    assert summary['complete_sku_count'] == 5
    assert summary['data_collection_completeness'] == 5 / 6
    items = {r['sku']: r for r in read_rows(wb['SKU补货总表'])}
    assert items['scenario-valid']['recommended_qty'] == items['eta-unknown']['recommended_qty'] == 50
    assert items['eta-unknown']['unknown_eta_inbound'] == 20
    assert items['known-zero']['recommended_qty'] == items['stopped']['recommended_qty'] == 0
    assert items['result-blocked']['recommended_qty'] is None


def test_multiple_fnsku_exceptions_keep_all_evidence_but_count_one_unknown_sku(tmp_path):
    con = connect(tmp_path / 'data.db')
    exception_table(con)
    con.executemany('INSERT INTO sku_inbound_exception VALUES(?,?,?,?,?,?,?,?,?)', [
        ('US', 'sku-a', 'FBA-A', 'FN-1', 'RECEIVING', 4, 5, None, 'INBOUND_QUANTITY_CONFLICT'),
        ('US', 'sku-a', 'FBA-A', 'FN-2', 'RECEIVING', 6, 8, None, 'INBOUND_QUANTITY_CONFLICT')])
    con.commit()
    wb, summary = export_rows(tmp_path, [plan()], con)
    detail = read_rows(wb['在途库存明细'])
    assert len(detail) == 2
    assert {row['fnsku'] for row in detail} == {'FN-1', 'FN-2'}
    assert all(row['remaining'] is row['units'] is None for row in detail)
    assert summary['inbound_quantity_conflict_sku_count'] == summary['unknown_data_sku_count'] == 1
    assert summary['inbound_exception_line_count'] == 2
    assert len(read_rows(wb['数据异常'])) == 1
    con.close()
