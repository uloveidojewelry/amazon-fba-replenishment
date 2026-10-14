from copy import deepcopy

import pytest
from openpyxl import load_workbook

from fba_replenishment.first_batch_export import (
    ANOMALIES_SHEET, DECISION_COLUMNS, DECISION_SHEET, DETAIL_COLUMNS,
    MEMBERS_SHEET, QUANTITY_FIELDS, STOP_COLUMNS, STOP_SHEET, export_first_batch,
)


def rows(ws):
    values = list(ws.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def payload():
    quantities = dict(zip(QUANTITY_FIELDS, (0, 0, 0, 0, 0, 0)))
    group_base = dict(snapshot_id='source-snapshot', marketplace='US',
                      api_marketplace='ATVPDKIKX0DER', mapping_basis='USER_WORKBOOK_EXACT_SKU',
                      priority='P0', priority_score=10, decision='APPROVED', approved_family='unsafe')
    member_base = dict(snapshot_id='source-snapshot', marketplace='US',
                       api_marketplace='ATVPDKIKX0DER', mapping_basis='USER_WORKBOOK_EXACT_SKU',
                       exact_reference_match=True, family=None,
                       mapping_status='AUTO_MATCHED_PENDING_REVIEW',
                       replenishment_status='BLOCKED_MAPPING', **quantities)
    return {
        'snapshot_id': 'source-snapshot',
        'summary': {'snapshot_id': 'source-snapshot', 'sku_count': 100, 'import_permitted': True},
        'groups': [
            dict(group_base, group_id='no-evidence', member_digest='digest-none',
                 candidate_family='J003', sku_count=1, batch_eligible=True),
            dict(group_base, group_id='stock', member_digest='digest-stock',
                 candidate_family='J001', sku_count=2, batch_eligible=True),
            dict(group_base, group_id='exception', member_digest='digest-exception',
                 candidate_family='J002', sku_count=1, batch_eligible=False),
        ],
        'items': [
            dict(member_base, group_id='no-evidence', sku='J003-size 5', asin='B000000003',
                 candidate_family='J003', recent_units_observed=None),
            dict(member_base, group_id='stock', sku='J001-size 5', asin='B000000001',
                 candidate_family='J001', inbound_shipped=4, fulfillable=2,
                 inventory_observed_at='2026-10-14T00:00:00+00:00',
                 recent_units_observed=None),
            dict(member_base, group_id='stock', sku='J001-size 6', asin='B000000002',
                 candidate_family='J001', fulfillable=None, recent_units_observed=None),
            dict(member_base, group_id='exception', sku='J002-size weird', asin='B000000004',
                 candidate_family='J002', recent_units_observed=3,
                 recent_sales_source='READ_ONLY_SAMPLE', recent_sales_status='POSITIVE_SAMPLE_ONLY',
                 replenishment_status='STOPPED'),
        ],
        'anomalies': [
            {'priority': 'P0', 'severity': 'BLOCKING', 'group_id': 'exception',
             'sku': 'J002-size weird', 'marketplace': 'US', 'api_marketplace': 'ATVPDKIKX0DER',
             'candidate_family': 'J002', 'code': 'REFERENCE_FORMAT_WARNING',
             'message': '必须人工核对', 'source': 'USER_WORKBOOK_EXACT_SKU'},
        ],
        'validation': [
            {'check': 'FIXED_MEMBERS', 'status': 'READY', 'severity': 'INFO',
             'group_id': None, 'sku': None, 'message': '已确认冻结成员集合'},
            {'check': 'USER_CONFIRMATION', 'status': 'PENDING', 'severity': 'BLOCKING',
             'group_id': None, 'sku': None, 'message': '正式导入需要用户明确确认'},
        ],
        'source': {'file_name': 'sku_batch_review.xlsx', 'sha256': 'source-file-hash',
                   'snapshot_id': 'source-snapshot'},
    }


def test_export_prioritizes_sample_sales_stock_and_explicit_members(tmp_path):
    data = payload()
    before = deepcopy(data)
    paths = export_first_batch(data, tmp_path)
    wb = load_workbook(paths['review_path'])
    assert wb.sheetnames == [DECISION_SHEET, MEMBERS_SHEET, ANOMALIES_SHEET, STOP_SHEET,
                            '模拟校验', '审核统计', '来源快照', '中文说明']
    groups = rows(wb[DECISION_SHEET])
    assert [row['group_id'] for row in groups] == ['exception', 'stock', 'no-evidence']
    assert all(row['decision'] == 'NEEDS_REVIEW' for row in groups)
    assert all(row['approved_family'] is row['reviewer'] is row['reason'] is None for row in groups)
    assert all(row['manifest_id'] == row['group_id'] for row in groups)
    details = rows(wb[MEMBERS_SHEET])
    assert [row['sku'] for row in details] == ['J002-size weird', 'J001-size 5', 'J001-size 6', 'J003-size 5']
    assert all(row['recommended_qty'] == 0 for row in details)
    assert details[1]['group_member_count'] == details[2]['group_member_count'] == 2
    assert details[0]['recent_sales_source'] == 'READ_ONLY_SAMPLE'
    assert details[0]['anomaly_codes'] == 'REFERENCE_FORMAT_WARNING'
    assert details[1]['inventory_observed_at'] == '2026-10-14T00:00:00+00:00'
    assert rows(wb['来源快照'])[0]['value'] == 'sku_batch_review.xlsx'
    assert data == before
    wb.close()


def test_unknown_inventory_is_not_summed_as_zero_and_samples_have_coverage(tmp_path):
    data = payload()
    data['items'][1]['recent_units_observed'] = 2
    wb = load_workbook(export_first_batch(data, tmp_path)['review_path'])
    groups = {row['group_id']: row for row in rows(wb[DECISION_SHEET])}
    stock = groups['stock']
    assert stock['fulfillable'] is None
    assert stock['reserved'] == 0
    assert stock['inbound_shipped'] == 4
    assert stock['inventory_complete_sku_count'] == 1
    assert stock['sample_units_observed'] == 2
    assert stock['sales_observed_sku_count'] == 1
    assert stock['recent_sales_positive_count'] == 1
    assert groups['no-evidence']['sample_units_observed'] is None
    details = {row['sku']: row for row in rows(wb[MEMBERS_SHEET])}
    assert details['J001-size 6']['fulfillable'] is None
    assert details['J003-size 5']['listing_status'] == 'UNKNOWN'
    assert details['J003-size 5']['recent_units_observed'] is None
    wb.close()


def test_protects_exception_approval_and_immutable_scope_but_allows_human_decisions(tmp_path):
    wb = load_workbook(export_first_batch(payload(), tmp_path)['review_path'])
    ws = wb[DECISION_SHEET]
    assert ws.protection.sheet is True
    decision_col = DECISION_COLUMNS.index('decision') + 1
    assert ws.cell(2, decision_col).value == 'NEEDS_REVIEW'
    assert ws.cell(2, decision_col).protection.locked is False
    assert '不能 APPROVE' in ws.cell(2, decision_col).comment.text
    assert ws.cell(2, DECISION_COLUMNS.index('approved_family') + 1).protection.locked is True
    assert ws.cell(2, DECISION_COLUMNS.index('approved_family') + 1).value is None
    assert ws.cell(2, DECISION_COLUMNS.index('reviewer') + 1).protection.locked is False
    assert ws.cell(2, DECISION_COLUMNS.index('reason') + 1).protection.locked is False
    assert ws.cell(3, decision_col).protection.locked is False
    assert ws.cell(3, DECISION_COLUMNS.index('member_digest') + 1).protection.locked is True
    validation, exception_validation = list(ws.data_validations.dataValidation)
    assert validation.formula1 == '"APPROVE,REJECT,NEEDS_REVIEW"'
    assert ws.cell(2, decision_col).coordinate not in validation
    assert ws.cell(3, decision_col).coordinate in validation
    assert exception_validation.formula1 == '"REJECT,NEEDS_REVIEW"'
    assert ws.cell(2, decision_col).coordinate in exception_validation
    assert ws.cell(3, decision_col).coordinate not in exception_validation
    assert validation.showErrorMessage is True
    assert validation.errorStyle == 'stop'
    assert wb[MEMBERS_SHEET].protection.sheet is True
    wb.close()


def test_stop_marks_are_unconfirmed_without_prefilling_formal_family(tmp_path):
    data = payload()
    data['items'][0].update(stop=True, reviewer='previous-reviewer', review_status='APPROVE')
    wb = load_workbook(export_first_batch(data, tmp_path)['review_path'])
    stops = rows(wb[STOP_SHEET])
    assert len(stops) == 4
    assert all(row['review_status'] == 'NEEDS_REVIEW' for row in stops)
    assert all(row['stop'] is row['reason'] is row['effective_date'] is row['reviewer'] is None for row in stops)
    assert all(row['family'] is None for row in stops)
    assert stops[0]['existing_replenishment_status'] == 'STOPPED'
    ws = wb[STOP_SHEET]
    stop_col = STOP_COLUMNS.index('stop') + 1
    assert 'FALSE 不会解除既有停补' in ws.cell(1, stop_col).comment.text
    assert ws.cell(2, stop_col).protection.locked is False
    assert ws.cell(2, STOP_COLUMNS.index('sku') + 1).protection.locked is True
    summary = {row['metric']: row['value'] for row in rows(wb['审核统计'])}
    assert summary['sku_count'] == 4
    assert summary['stock_positive_sku_count'] == 1
    assert summary['stock_or_inbound_positive_sku_count'] == 1
    assert summary['inbound_only_sku_count'] == 0
    assert summary['recent_sales_positive_sku_count'] == 1
    assert summary['exception_group_count'] == 1
    assert summary['anomaly_sku_count'] == 1
    assert summary['import_permitted'] is False
    assert summary['require_confirmation'] is True
    assert summary['full_history_sync'] is False
    wb.close()


def test_inventory_count_separates_fba_stock_from_inbound_only_evidence(tmp_path):
    data = payload()
    data['items'][1]['fulfillable'] = 0
    wb = load_workbook(export_first_batch(data, tmp_path)['review_path'])
    summary = {row['metric']: row['value'] for row in rows(wb['审核统计'])}
    assert summary['stock_positive_sku_count'] == 0
    assert summary['stock_or_inbound_positive_sku_count'] == 1
    assert summary['inbound_only_sku_count'] == 1
    group = next(row for row in rows(wb[DECISION_SHEET]) if row['group_id'] == 'stock')
    assert group['stock_positive_count'] == 1
    wb.close()


def test_chinese_instructions_cover_group_approval_stop_inheritance_and_confirmation(tmp_path):
    wb = load_workbook(export_first_batch(payload(), tmp_path)['review_path'])
    instructions = ' '.join(str(row['操作说明']) for row in rows(wb['中文说明']))
    for text in ('APPROVE', 'NEEDS_REVIEW', '未来所有同前缀 SKU',
                 '本批之外的已确认成员', '明确确认', '不等于零', '当前文件不能强行批准',
                 '异常可以拒绝', '重新生成单体审核集合'):
        assert text in instructions
    wb.close()


def test_formula_injection_is_literal_and_nonoperational_fields_are_omitted(tmp_path):
    data = payload()
    literal = '=HYPERLINK("https://invalid.example","SKU")'
    data['items'][0]['sku'] = literal
    data['summary']['access_token'] = 'TOKEN_SENTINEL'
    data['source']['client_secret'] = 'SECRET_SENTINEL'
    for kind in ('groups', 'items', 'anomalies', 'validation'):
        for row in data[kind]:
            row['refresh_token'] = 'SECRET_SENTINEL'
            row['buyer_name'] = 'CUSTOMER_SENTINEL'
    wb = load_workbook(export_first_batch(data, tmp_path)['review_path'])
    matched = False
    for ws in wb:
        for row in ws:
            for cell in row:
                assert cell.value not in ('TOKEN_SENTINEL', 'SECRET_SENTINEL', 'CUSTOMER_SENTINEL')
                if cell.value == literal:
                    matched = True
                    assert cell.data_type == 's'
    assert matched
    wb.close()


@pytest.mark.parametrize('change', [
    'duplicate_member', 'missing_group', 'bad_count', 'duplicate_group',
    'wrong_snapshot', 'wrong_member_digest', 'not_exact', 'foreign_anomaly',
])
def test_ambiguous_or_nonexact_fixed_sets_are_rejected_before_file_creation(tmp_path, change):
    data = payload()
    if change == 'duplicate_member':
        data['items'].append(dict(data['items'][0]))
    elif change == 'missing_group':
        data['items'][0]['group_id'] = 'unregistered'
    elif change == 'bad_count':
        data['groups'][0]['sku_count'] = 500
    elif change == 'duplicate_group':
        data['groups'].append(dict(data['groups'][0]))
    elif change == 'wrong_snapshot':
        data['items'][0]['snapshot_id'] = 'newer-snapshot'
    elif change == 'wrong_member_digest':
        data['items'][0]['member_digest'] = 'different-set'
    elif change == 'not_exact':
        data['items'][0].update(exact_reference_match=False, mapping_basis='FORMAT_ONLY')
    elif change == 'foreign_anomaly':
        data['anomalies'][0]['sku'] = 'outside-set'
    with pytest.raises(ValueError):
        export_first_batch(data, tmp_path)
    assert not list(tmp_path.glob('*.xlsx'))


def test_large_family_uses_one_row_per_sku_instead_of_a_truncated_member_cell(tmp_path):
    data = payload()
    group = dict(data['groups'][0], sku_count=400)
    template = data['items'][0]
    data['groups'] = [group]
    data['items'] = [dict(template, sku='J003-size ' + str(i) + '-' + ('X' * 80)) for i in range(400)]
    data['anomalies'] = []
    assert len(';'.join(row['sku'] for row in data['items'])) > 32767
    wb = load_workbook(export_first_batch(data, tmp_path)['review_path'], read_only=True)
    details = rows(wb[MEMBERS_SHEET])
    assert len(details) == 400
    assert {row['sku'] for row in details} == {row['sku'] for row in data['items']}
    assert rows(wb[DECISION_SHEET])[0]['sku_count'] == 400
    wb.close()
