from copy import deepcopy

import pytest
from openpyxl import load_workbook

from fba_replenishment.batch_export import (
    DECISION_COLUMNS, DECISION_SHEET, MEMBERS_SHEET, STOP_SHEET, export_batch_review,
)


def rows(ws):
    values = list(ws.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def queue():
    return {
        'snapshot_id': 'snapshot-fixed',
        'summary': {'sku_count': 3, 'group_count': 2, 'exact_match_sku_count': 2},
        'groups': [
            {'group_id': 'fixed-group', 'member_digest': 'fixed-members-digest',
             'marketplace': 'US', 'api_marketplace': 'ATVPDKIKX0DER',
             'candidate_family': 'J001', 'mapping_basis': 'USER_WORKBOOK_EXACT_SKU',
             'naming_format': 'MODEL-size N', 'sku_count': 2, 'exact_match_count': 2,
             'priority': 'P0', 'priority_reasons': '精确匹配；有 FBA 库存',
             'batch_eligible': True},
            {'group_id': 'format-only', 'member_digest': 'unknown-members-digest',
             'marketplace': 'US', 'api_marketplace': 'ATVPDKIKX0DER',
             'candidate_family': None, 'mapping_basis': 'FORMAT_ONLY',
             'naming_format': 'UNKNOWN', 'sku_count': 1, 'exact_match_count': 0,
             'priority': 'P2', 'batch_eligible': False},
        ],
        'items': [
            {'group_id': 'format-only', 'sku': 'odd-sku', 'asin': 'B000000003',
             'mapping_status': 'REVIEW_REQUIRED', 'candidate_family': None,
             'fulfillable': None, 'reserved': None, 'unfulfillable': None,
             'recent_units_observed': None, 'recommended_qty': 999},
            {'group_id': 'fixed-group', 'sku': 'J001-size 6', 'asin': 'B000000002',
             'mapping_status': 'AUTO_MATCHED_PENDING_REVIEW', 'candidate_family': 'J001',
             'exact_reference_match': True,
             'replenishment_status': 'STOPPED', 'fba_stock_status': 'HAS_STOCK',
             'fulfillable': 4, 'reserved': 1, 'unfulfillable': 0,
             'recent_sales_status': 'UNKNOWN', 'recent_units_observed': None},
            {'group_id': 'fixed-group', 'sku': 'J001-size 5', 'asin': 'B000000001',
             'mapping_status': 'AUTO_MATCHED_PENDING_REVIEW', 'candidate_family': 'J001',
             'exact_reference_match': True,
             'replenishment_status': 'BLOCKED_MAPPING', 'fulfillable': 0,
             'listing_status': 'UNKNOWN', 'listing_evidence': 'NOT_QUERIED'},
        ],
        'anomalies': [
            {'severity': 'WARNING', 'sku': 'J001-size 5', 'code': 'UNVERIFIED_LISTING'},
            {'severity': 'BLOCKING', 'sku': 'odd-sku', 'code': 'UNKNOWN_FAMILY'},
        ],
        'audit': [{'at': '2026-10-14T00:00:00Z', 'action': 'REVIEW_QUEUE', 'reviewer': 'tester'}],
    }


def test_excel_prioritizes_explicit_series_sets_without_approving(tmp_path):
    data = queue()
    before = deepcopy(data)
    paths = export_batch_review(data, tmp_path)
    wb = load_workbook(paths['review_path'])
    assert wb.sheetnames[:4] == ['系列优先审核', '数据异常', '精确匹配优先', DECISION_SHEET]
    groups = rows(wb['系列优先审核'])
    assert [row['group_id'] for row in groups] == ['fixed-group', 'format-only']
    assert groups[0]['exact_match_count'] == 2
    assert rows(wb['数据异常'])[0]['severity'] == 'BLOCKING'
    decisions = rows(wb[DECISION_SHEET])
    assert all(row['decision'] == 'PENDING' for row in decisions)
    assert all(row['approved_family'] is row['reviewer'] is row['reason'] is None for row in decisions)
    assert [row['manifest_id'] for row in decisions] == ['fixed-group', 'format-only']
    members = rows(wb[MEMBERS_SHEET])
    assert [row['sku'] for row in members] == ['J001-size 5', 'J001-size 6', 'odd-sku']
    assert all(row['snapshot_id'] == 'snapshot-fixed' for row in members)
    assert members[0]['member_digest'] == members[1]['member_digest'] == 'fixed-members-digest'
    assert members[2]['member_digest'] == 'unknown-members-digest'
    assert all(row['recommended_qty'] == 0 for row in members)
    assert data == before


def test_format_only_groups_cannot_use_approval_dropdown(tmp_path):
    paths = export_batch_review(queue(), tmp_path)
    wb = load_workbook(paths['review_path'])
    ws = wb[DECISION_SHEET]
    decision_col = DECISION_COLUMNS.index('decision') + 1
    assert ws.protection.sheet is True
    assert ws.cell(2, decision_col).protection.locked is False
    assert ws.cell(3, decision_col).protection.locked is True
    validations = list(ws.data_validations.dataValidation)
    assert len(validations) == 1
    assert ws.cell(2, decision_col).coordinate in validations[0]
    assert ws.cell(3, decision_col).coordinate not in validations[0]
    assert '未确认候选系列' in ws.cell(3, decision_col).comment.text


def test_exact_match_queue_remains_pending_and_statistics_explain_workload(tmp_path):
    data = queue()
    data['summary'].update(exact_exception_sku_count=0, exact_review_action_count=1,
                           exact_actions_saved=1, listing_unknown_count=3,
                           fullmatch_candidate_sku_count=0, naming_candidate_sku_count=0)
    paths = export_batch_review(data, tmp_path)
    wb = load_workbook(paths['review_path'])
    precise = rows(wb['精确匹配优先'])
    assert {row['sku'] for row in precise} == {'J001-size 5', 'J001-size 6'}
    assert all(row['mapping_status'] == 'AUTO_MATCHED_PENDING_REVIEW' for row in precise)
    assert all(row['family'] is None and row['recommended_qty'] == 0 for row in precise)
    assert rows(wb[DECISION_SHEET])[0]['decision'] == 'PENDING'
    summary = {row['metric']: row['value'] for row in rows(wb['审核统计'])}
    assert summary['exact_actions_saved'] == summary['exact_review_action_count'] == 1
    assert summary['listing_unknown_count'] == 3
    assert any('不代表实际审核耗时' in (row['操作说明'] or '') for row in rows(wb['使用说明']))


def test_unknown_sales_inventory_and_listing_are_not_zero_or_inactive(tmp_path):
    paths = export_batch_review(queue(), tmp_path)
    wb = load_workbook(paths['review_path'])
    members = {row['sku']: row for row in rows(wb[MEMBERS_SHEET])}
    unknown = members['odd-sku']
    assert unknown['recent_units_observed'] is None
    assert unknown['fulfillable'] is unknown['reserved'] is unknown['unfulfillable'] is None
    assert unknown['recent_sales_status'] == unknown['fba_stock_status'] == 'UNKNOWN'
    assert unknown['listing_status'] == unknown['history_status'] == 'UNKNOWN'
    assert members['J001-size 5']['fulfillable'] == 0
    assert members['J001-size 6']['fba_stock_status'] == 'HAS_STOCK'
    assert members['J001-size 6']['listing_status'] == 'UNKNOWN'


def test_stop_import_template_has_no_inferred_stop_or_authoritative_family(tmp_path):
    paths = export_batch_review(queue(), tmp_path)
    for path in (paths['review_path'], paths['stop_template_path']):
        wb = load_workbook(path)
        stops = rows(wb[STOP_SHEET])
        assert len(stops) == 3
        assert all(row['review_status'] == 'UNAPPROVED' for row in stops)
        assert all(row['stop'] is row['reviewer'] is row['effective_date'] is None for row in stops)
        assert all(row['family'] is None for row in stops)
        assert stops[0]['candidate_family'] == 'J001'
        assert 'FALSE 不自动恢复' in wb[STOP_SHEET].cell(1, 6).comment.text
        assert any('不能批准未来同前缀' in (row['操作说明'] or '') for row in rows(wb['使用说明']))


def test_formula_injection_and_unknown_sensitive_fields_are_not_exported(tmp_path):
    data = queue()
    external = '=HYPERLINK("https://invalid.example","SKU")'
    data['items'][0]['sku'] = external
    data['groups'][1]['candidate_family'] = external
    data['audit'][0]['reason'] = external
    data['summary']['access_token'] = 'TOKEN_SENTINEL'
    for key in ('items', 'groups', 'audit', 'anomalies'):
        for row in data[key]:
            row['client_secret'] = 'SECRET_SENTINEL'
            row['buyer_name'] = 'CUSTOMER_SENTINEL'
            row['raw_response'] = {'access_token': 'TOKEN_SENTINEL'}
    paths = export_batch_review(data, tmp_path)
    for path in paths.values():
        wb = load_workbook(path)
        seen_external = False
        for ws in wb:
            for row in ws:
                for cell in row:
                    assert cell.value not in ('TOKEN_SENTINEL', 'SECRET_SENTINEL', 'CUSTOMER_SENTINEL')
                    if cell.value == external:
                        seen_external = True
                        assert cell.data_type == 's'
        assert seen_external


def test_large_family_exports_all_members_without_a_concatenated_cell(tmp_path):
    data = queue()
    skus = ['J001-size ' + str(i) + '-' + ('X' * 80) for i in range(400)]
    assert len(';'.join(skus)) > 32767
    data['groups'] = [dict(data['groups'][0], sku_count=len(skus))]
    data['items'] = [{'group_id': 'fixed-group', 'sku': sku} for sku in skus]
    paths = export_batch_review(data, tmp_path)
    wb = load_workbook(paths['review_path'], read_only=True)
    members = rows(wb[MEMBERS_SHEET])
    assert len(members) == len(skus)
    assert {row['sku'] for row in members} == set(skus)
    assert rows(wb[DECISION_SHEET])[0]['sku_count'] == len(skus)
    wb.close()


@pytest.mark.parametrize('change', ['duplicate', 'missing_group', 'bad_count'])
def test_export_rejects_ambiguous_membership_before_writing(tmp_path, change):
    data = queue()
    if change == 'duplicate':
        data['items'].append(dict(data['items'][0]))
    elif change == 'missing_group':
        data['items'][0]['group_id'] = 'not-present'
    else:
        data['groups'][0]['sku_count'] = 99
    with pytest.raises(ValueError):
        export_batch_review(data, tmp_path)
    assert not list(tmp_path.glob('*.xlsx'))
