from datetime import date, timedelta

from openpyxl import load_workbook

from fba_replenishment.core import connect
from fba_replenishment.export import export, export_mapping_review


def read_rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def plan_row(sku='a', family='family', market='US', **changes):
    row = {'marketplace': market, 'sku': sku, 'family': family, 'asin': '',
           'status': 'ACTIVE', 'stop_reason': '', 'recommended_qty': 100,
           'sales_30d': 30, 'observed_sales_days': 30, 'daily_sales': 1,
           'fulfillable': 5, 'reserved': 2, 'unfulfillable': 0,
           'eligible_inbound': 0, 'unknown_eta_inbound': 0, 'coverage_days': 5,
           'asof': '2026-10-13', 'lead_days': 35, 'target_days': 75,
           'safety_days': 0, 'moq': 1, 'case_pack': 1, 'sales_timezone': 'America/Los_Angeles'}
    row.update(changes)
    return row


def certified_windows(units=1):
    return {key: value for days in (7, 30, 60, 90) for key, value in (
        (f'sales_{days}d', days * units),
        (f'observed_sales_days_{days}d', days),
        (f'sales_complete_{days}d', True),
    )}


def test_real_acceptance_exports_sources_parameters_risk_and_windows(tmp_path):
    path = tmp_path / 'acceptance.xlsx'
    acceptance = {
        'summary': {'collection_status': 'COMPLETE', 'data_source': 'SP-API', 'orders_collected': 77,
                    'inbound_v2024_status': 'NOT_VERIFIED'},
        'per_sku': {'a': dict(certified_windows(2), inventory_complete=True,
                              sales_source='Orders v0; AFN; canceled excluded',
                              inventory_source='FBA Inventory v1', inbound_source='Fulfillment Inbound v0')},
        'audit': [{'stage': 'full', 'kind': 'orders', 'status': 'SUCCESS', 'row_count': 77,
                   'pages': 4, 'source': 'Orders v0', 'detail': 'Validated complete calendar window'}],
    }
    export([plan_row()], path, acceptance=acceptance, settings={'lead_days': 35})
    wb = load_workbook(path)
    assert {'Family Summary', 'Stopped Families', 'Stockout Risk', 'Data Anomalies',
            'Collection Audit', 'Acceptance'} <= set(wb.sheetnames)
    item = read_rows(wb['Replenishment'])[0]
    assert [item[f'sales_{days}d'] for days in (7, 30, 60, 90)] == [14, 60, 120, 180]
    assert item['inventory_source'] == 'FBA Inventory v1'
    assert item['sales_completeness_90d'] == 1
    assert read_rows(wb['Family Summary'])[0]['sales_90d'] == 180
    assert read_rows(wb['Stockout Risk'])[0]['risk'] == 'STOCKOUT_BEFORE_REPLENISHMENT'
    summary = {r['metric']: r['value'] for r in read_rows(wb['Acceptance'])}
    assert summary['orders_collected'] == 77
    assert summary['inventory_coverage'] == summary['sales_completeness'] == 1
    assert summary['calculable_sku_count'] == 1
    assert read_rows(wb['Collection Audit'])[0]['pages'] == 4
    parameters = {r['parameter']: r['value'] for r in read_rows(wb['Parameters'])}
    assert parameters['lead_days'] == 35
    assert parameters['mode'] == 'READ_ONLY'


def test_sample_does_not_certify_old_database_or_unknown_inventory(tmp_path):
    con = connect(tmp_path / 'data.db')
    end = date(2026, 10, 13)
    con.executemany('INSERT INTO sales VALUES(?,?,?,?)', [
        ('US', 'a', (end - timedelta(days=i)).isoformat(), 0) for i in range(90)])
    con.commit()
    path = tmp_path / 'sample.xlsx'
    export([plan_row(reserved=None)], path, con, acceptance={
        'summary': {'collection_status': 'FAILED_QUALITY_GATE', 'data_source': 'SP-API'},
        'per_sku': {'a': {'observed_sales_days_30d': 2, 'sales_30d': 2}},
    })
    wb = load_workbook(path)
    row = read_rows(wb['Replenishment'])[0]
    assert all(row[f'sales_{days}d'] is None for days in (7, 30, 60, 90))
    assert row['observed_sales_30d'] == 2
    assert row['observed_sales_days_90d'] is None
    assert row['sales_completeness_90d'] is None
    assert row['reserved'] is None
    assert row['daily_sales'] is row['coverage_days'] is None
    assert row['recommended_qty'] == 0
    assert row['status'] == 'DATA_REVIEW_REQUIRED'
    summary = {r['metric']: r['value'] for r in read_rows(wb['Acceptance'])}
    assert summary['inventory_coverage'] == 0
    assert summary['sales_completeness'] == 0
    assert summary['blocked_sku_count'] == 1
    con.close()


def test_calendar_gaps_remain_unknown_and_known_zero_is_preserved(tmp_path):
    con = connect(tmp_path / 'data.db')
    end = date(2026, 10, 13)
    # Ninety explicit zero days are a known zero; omitted dates are unknown.
    for sku in ('known-zero', 'gap'):
        con.executemany('INSERT INTO sales VALUES(?,?,?,?)', [
            ('US', sku, (end - timedelta(days=i)).isoformat(), 0)
            for i in range(90) if sku == 'known-zero' or i != 8])
    con.commit()
    path = tmp_path / 'calendar.xlsx'
    export([plan_row('known-zero', sales_30d=0), plan_row('gap', sales_30d=0)], path, con,
           acceptance={'summary': {'collection_status': 'COMPLETE'}})
    wb = load_workbook(path)
    rows = {r['sku']: r for r in read_rows(wb['Replenishment'])}
    assert rows['known-zero']['sales_90d'] == 0
    assert rows['known-zero']['sales_complete_90d'] is True
    assert rows['gap']['sales_7d'] == 0
    assert rows['gap']['sales_30d'] is None
    assert rows['gap']['sales_90d'] is None
    assert rows['gap']['observed_sales_days_90d'] == 89
    assert rows['gap']['recommended_qty'] == 0
    assert read_rows(wb['Family Summary'])[0]['sales_90d'] is None
    con.close()


def test_stop_rule_applies_to_marketplace_family_and_input_is_unmodified(tmp_path):
    rows = [plan_row('a', status='STOPPED', recommended_qty=0), plan_row('b'),
            plan_row('c', market='CA'), plan_row('d', family='other')]
    path = tmp_path / 'stops.xlsx'
    export(rows, path, acceptance={'per_sku': {r['sku']: certified_windows() for r in rows}})
    wb = load_workbook(path)
    detail = {r['sku']: r for r in read_rows(wb['Replenishment'])}
    assert detail['a']['recommended_qty'] == detail['b']['recommended_qty'] == 0
    assert detail['a']['status'] == detail['b']['status'] == 'STOPPED'
    assert detail['c']['recommended_qty'] == detail['d']['recommended_qty'] == 100
    assert rows[1]['recommended_qty'] == 100
    stopped = read_rows(wb['Stopped Families'])
    assert len(stopped) == 1
    assert stopped[0]['marketplace'] == 'US'
    assert stopped[0]['recommended_qty'] == 0
    assert stopped[0]['sku_count'] == 2
    assert 'STOP_RULE_INCONSISTENCY' in {r['code'] for r in read_rows(wb['Data Anomalies'])}


def test_external_strings_are_literal_and_unlisted_sensitive_fields_omitted(tmp_path):
    text = '=HYPERLINK("example.invalid","external")'
    path = tmp_path / 'literal.xlsx'
    export([plan_row(text, family=text)], path, acceptance={
        'summary': {'collection_status': 'SAMPLE_ONLY', 'access_token': 'TOKEN_SENTINEL'},
        'audit': [{'detail': text, 'raw_response': 'CUSTOMER_SENTINEL', 'access_token': 'TOKEN_SENTINEL'}],
        'anomalies': [{'sku': text, 'code': 'UNKNOWN_SKU', 'message': text,
                       'buyer_name': 'CUSTOMER_SENTINEL', 'refresh_token': 'TOKEN_SENTINEL'}],
    })
    wb = load_workbook(path)
    for ws in wb:
        for row in ws:
            for cell in row:
                assert cell.value not in ('TOKEN_SENTINEL', 'CUSTOMER_SENTINEL')
                if cell.value == text:
                    assert cell.data_type == 's'


def test_missing_family_is_blocked_even_when_input_claims_active(tmp_path):
    path = tmp_path / 'unmapped.xlsx'
    export([plan_row(family='')], path, acceptance={'per_sku': {'a': certified_windows()}})
    wb = load_workbook(path)
    row = read_rows(wb['Replenishment'])[0]
    assert row['status'] == 'FAMILY_REVIEW_REQUIRED'
    assert row['recommended_qty'] == 0
    assert not read_rows(wb['Family Summary'])
    assert read_rows(wb['Data Anomalies'])[0]['code'] == 'FAMILY_REVIEW_REQUIRED'


def test_mapping_and_stop_reviews_never_infer_approval_or_unknown_stop_false(tmp_path):
    mapping, stops = tmp_path / 'mapping.xlsx', tmp_path / 'stop.xlsx'
    records = [{'marketplace': 'US', 'sku': 'J001-size 6', 'asin': 'ASIN-A',
                'candidate_family': 'J001', 'mapping_basis': 'SKU naming candidate only',
                'parent_asins': ['PARENT-B', 'PARENT-A', 'PARENT-A'], 'catalog_status': 'SUCCESS',
                'fulfillable': 15, 'reserved': None, 'unfulfillable': 0,
                'review_status': 'APPROVED', 'recommended_qty': 99,
                'stop_replenishment': False, 'buyer_name': 'CUSTOMER_SENTINEL'},
               {'marketplace': 'US', 'sku': '=1+1', 'asin': 'ASIN-B',
                'candidate_family': None, 'mapping_basis': 'UNKNOWN', 'catalog_status': 'FORBIDDEN'}]
    export_mapping_review(records, mapping, stops, summary={
        'summary': {'collection_status': 'SAMPLE_ONLY', 'mapping_pending_sku_count': 2,
                    'access_token': 'TOKEN_SENTINEL'},
        'audit': [{'kind': 'catalog', 'status': 'PARTIAL', 'row_count': 2, 'detail': 'Review required',
                   'raw_response': 'CUSTOMER_SENTINEL'}],
    })
    mapping_wb, stop_wb = load_workbook(mapping), load_workbook(stops)
    assert set(mapping_wb.sheetnames) == {'SKU Mapping Review', 'Read Me', 'Collection Audit', 'Review Summary'}
    assert set(stop_wb.sheetnames) == {'Stop List Review', 'Read Me', 'Collection Audit', 'Review Summary'}
    mappings = read_rows(mapping_wb['SKU Mapping Review'])
    assert mappings[0]['candidate_family'] == 'J001'
    assert mappings[0]['parent_asins'] == 'PARENT-A; PARENT-B'
    assert mappings[0]['reserved'] is None
    assert mappings[1]['candidate_family'] is None
    for row in mappings:
        assert row['review_status'] == 'UNAPPROVED'
        assert row['recommended_qty'] == 0
        assert row['approved_family'] is row['reviewer'] is row['reviewed_at'] is None
    for row in read_rows(stop_wb['Stop List Review']):
        assert row['stop_replenishment'] is None
        assert row['family'] is row['reason'] is row['effective_date'] is None
        assert row['review_status'] == 'UNAPPROVED'
        assert row['recommended_qty'] == 0
    for wb in (mapping_wb, stop_wb):
        for ws in wb:
            for row in ws:
                for cell in row:
                    assert cell.value not in ('TOKEN_SENTINEL', 'CUSTOMER_SENTINEL')
                    if cell.value == '=1+1':
                        assert cell.data_type == 's'


def test_discovery_sheets_classify_pending_stops_and_conflicts_without_approving(tmp_path):
    path = tmp_path / 'discovery.xlsx'
    records = [
        {'marketplace': 'US', 'sku': 'new', 'asin': 'ASIN-NEW', 'candidate_family': 'J001',
         'family': None, 'effective_family': 'J001', 'mapping_basis': 'Candidate SKU prefix',
         'mapping_status': 'AUTO_DISCOVERED_PENDING_REVIEW', 'replenishment_status': 'STOPPED',
         'first_discovered_at': '2026-10-14T01:00:00Z', 'last_seen_at': '2026-10-14T02:00:00Z',
         'is_new': True, 'recommended_qty': 88, 'access_token': 'TOKEN_SENTINEL'},
        {'marketplace': 'US', 'sku': 'approved', 'asin': 'ASIN-OLD', 'family': 'J001',
         'mapping_status': 'APPROVED', 'replenishment_status': 'ACTIVE', 'is_new': False,
         'recommended_qty': 12},
        {'marketplace': 'CA', 'sku': 'canada', 'asin': 'ASIN-CA', 'family': 'J001',
         'review_status': 'APPROVED', 'replenishment_status': 'ACTIVE', 'is_new': False,
         'recommended_qty': 4},
        {'marketplace': 'US', 'sku': '=1+1', 'asin': 'ASIN-CONFLICT', 'candidate_family': 'J002',
         'family': 'J003', 'mapping_status': 'MAPPING_CONFLICT', 'replenishment_status': 'MAPPING_CONFLICT',
         'is_new': True, 'recommended_qty': 999, 'buyer_name': 'CUSTOMER_SENTINEL'},
    ]
    export([], path, discovery=records)
    wb = load_workbook(path)
    assert {'新发现 SKU', '待审核 SKU', 'SKU 系列映射', '已停补系列', '系列映射冲突'} <= set(wb.sheetnames)
    assert {r['SKU'] for r in read_rows(wb['新发现 SKU'])} == {'new', '=1+1'}
    assert {r['SKU'] for r in read_rows(wb['待审核 SKU'])} == {'new', '=1+1'}
    assert {r['SKU'] for r in read_rows(wb['已停补系列'])} == {'new', 'approved'}
    assert {r['SKU'] for r in read_rows(wb['系列映射冲突'])} == {'=1+1'}
    mapped = {r['SKU']: r for r in read_rows(wb['SKU 系列映射'])}
    assert mapped['new']['正式系列'] is None
    assert mapped['new']['候选系列'] == mapped['new']['安全停补系列'] == 'J001'
    assert mapped['new']['审核状态'] == 'AUTO_DISCOVERED_PENDING_REVIEW'
    assert mapped['new']['首次发现时间'] == '2026-10-14T01:00:00Z'
    assert mapped['new']['最近同步时间'] == '2026-10-14T02:00:00Z'
    assert mapped['approved']['补货状态'] == 'STOPPED'
    assert mapped['new']['建议补货数量'] == mapped['approved']['建议补货数量'] == 0
    assert mapped['=1+1']['建议补货数量'] == 0
    assert mapped['canada']['建议补货数量'] == 4
    assert records[1]['recommended_qty'] == 12
    for ws in wb:
        for row in ws:
            for cell in row:
                assert cell.value not in ('TOKEN_SENTINEL', 'CUSTOMER_SENTINEL')
                if cell.value == '=1+1':
                    assert cell.data_type == 's'


def test_discovery_explicit_empty_list_has_review_sheets_and_legacy_has_none(tmp_path):
    path = tmp_path / 'empty.xlsx'
    export([], path, discovery=[])
    wb = load_workbook(path)
    assert wb['新发现 SKU'].max_row == wb['SKU 系列映射'].max_row == 1
    export([], path)
    wb = load_workbook(path)
    assert 'SKU 系列映射' not in wb.sheetnames
