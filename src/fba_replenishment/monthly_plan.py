"""Read-only user-monthly-demand scenarios over existing safe SKU snapshots.

Workbook data never replaces SP-API sales, identities, or configuration. All
existing API data blockers and direct SKU stops remain effective.
"""
from collections import defaultdict
from collections.abc import Mapping
from copy import deepcopy
from fractions import Fraction
import hashlib
import math
from pathlib import Path

from openpyxl import load_workbook

from .sku_plan import _settings


def _units(value):
    if type(value) not in (int, float) or not math.isfinite(value):
        return None
    return int(value) if value >= 0 and int(value) == value else None


def read_monthly_workbook(path, sheet_name=None):
    """Read two input columns, retaining exact case-sensitive SKU identities.

    Formula cells are rejected rather than evaluated or read through Excel's
    cached result. Equal duplicate quantities are deduplicated; conflicting
    duplicates remain unknown and are never added together.
    """
    path = Path(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name] if sheet_name is not None else workbook.active
        rows = sheet.iter_rows()
        header = next(rows, None)
        if not header or not isinstance(header[0].value, str) or header[0].value.strip() not in ('sku', 'SKU'):
            raise ValueError('Monthly workbook first column must have an explicit SKU header')
        issues = ['MONTHLY_PERIOD_UNSPECIFIED']
        sales_header = header[1].value if len(header) > 1 else None
        if sales_header is None or isinstance(sales_header, str) and not sales_header.strip():
            issues.append('MONTHLY_QUANTITY_HEADER_MISSING')
        elif type(sales_header) in (int, float, bool):
            issues.append('MONTHLY_QUANTITY_HEADER_NUMERIC')
        elif getattr(header[1], 'data_type', '') == 'f':
            issues.append('MONTHLY_QUANTITY_HEADER_FORMULA')
        grouped, invalid_identity = defaultdict(list), []
        audit = dict(source_sha256=digest, source_file=path.name, sheet_name=sheet.title,
                     header_row=1, monthly_period_status='ASSUMED_30_DAY_MONTH',
                     monthly_period=None, quantity_header=sales_header,
                     issues=issues, input_data_rows=0, blank_rows=0,
                     sku_whitespace_trimmed_rows=[], duplicate_identical_skus=0,
                     duplicate_conflict_skus=0)
        for number, cells in enumerate(rows, 2):
            sku_cell = cells[0]
            quantity_cell = cells[1] if len(cells) > 1 else None
            sku_value = sku_cell.value
            value = quantity_cell.value if quantity_cell is not None else None
            if sku_value is None and value is None:
                audit['blank_rows'] += 1
                continue
            audit['input_data_rows'] += 1
            row_issues = []
            if sku_cell.data_type == 'f':
                sku = ''
                row_issues.append('MONTHLY_SKU_FORMULA')
            elif not isinstance(sku_value, str) or not sku_value.strip():
                sku = ''
                row_issues.append('MONTHLY_SKU_INVALID')
            else:
                sku = sku_value.strip()
                if sku != sku_value:
                    row_issues.append('SKU_WHITESPACE_TRIMMED')
                    audit['sku_whitespace_trimmed_rows'].append(number)
            if quantity_cell is not None and quantity_cell.data_type == 'f':
                units = None
                row_issues.append('MONTHLY_QUANTITY_FORMULA')
            else:
                units = _units(value)
                if units is None:
                    row_issues.append('MONTHLY_QUANTITY_INVALID')
            status = 'VALID' if sku and units is not None else 'INVALID'
            record = dict(sku=sku, monthly_units=units, monthly_status=status,
                          source_rows=[number], issues=row_issues,
                          source_sheet=sheet.title, source_sha256=digest)
            if sku:
                grouped[sku].append(record)
            else:
                record['monthly_units'] = None
                invalid_identity.append(record)
        records = []
        for sku, entries in grouped.items():
            first = deepcopy(entries[0])
            if len(entries) > 1:
                first['source_rows'] = [n for entry in entries for n in entry['source_rows']]
                first['issues'] = list(dict.fromkeys(code for entry in entries for code in entry['issues']))
                same = all(entry['monthly_status'] == 'VALID' for entry in entries) and len({entry['monthly_units'] for entry in entries}) == 1
                if same:
                    first['issues'].append('DUPLICATE_IDENTICAL')
                    audit['duplicate_identical_skus'] += 1
                else:
                    first.update(monthly_units=None, monthly_status='DUPLICATE_CONFLICT')
                    first['issues'].append('DUPLICATE_CONFLICT')
                    audit['duplicate_conflict_skus'] += 1
            records.append(first)
        records.extend(invalid_identity)
        records.sort(key=lambda record: record['source_rows'][0])
        audit.update(record_count=len(records), valid_record_count=sum(r['monthly_status'] == 'VALID' for r in records),
                     invalid_record_count=sum(r['monthly_status'] == 'INVALID' for r in records))
        return dict(records=records, audit=audit)
    finally:
        workbook.close()


def calculate_monthly_plan(api_rows, records, *, monthly_days=30, lead_days=35,
                           target_days=75, safety_days=0, moq=1, case_pack=1,
                           parameters_confirmed=False,
                           monthly_period_status='ASSUMED_30_DAY_MONTH'):
    """Use user monthly demand without weakening the safe API snapshot gate.

    ``api_rows`` must come from calculate_skus followed by the latest direct SKU
    stop overlay. Unknowns stay None, and the inputs are never modified.
    """
    _settings(lead_days, target_days, safety_days, moq, case_pack, 2,
              'confirmed_eta', parameters_confirmed)
    if (type(monthly_days) not in (int, float) or not math.isfinite(monthly_days)
            or monthly_days <= 0):
        raise ValueError('monthly_days must be finite and positive')
    if monthly_period_status not in ('ASSUMED_30_DAY_MONTH', 'USER_CONFIRMED_PERIOD'):
        raise ValueError('Monthly period status must be assumed or user confirmed')
    confirmed = parameters_confirmed and monthly_period_status == 'USER_CONFIRMED_PERIOD'
    known = {}
    for row in api_rows:
        if not isinstance(row, Mapping):
            raise ValueError('API planning rows must be mappings')
        if row.get('marketplace') != 'US' or row.get('api_marketplace') not in (None, 'ATVPDKIKX0DER'):
            continue
        sku = row.get('sku')
        if not isinstance(sku, str) or not sku:
            raise ValueError('API SKU identity is invalid')
        if sku in known:
            raise ValueError('Duplicate US API SKU identity')
        known[sku] = row
    result, seen = [], set()
    for record in records:
        if not isinstance(record, Mapping):
            raise ValueError('Monthly records must be mappings')
        sku = record.get('sku')
        if not isinstance(sku, str):
            raise ValueError('Monthly SKU must be a string')
        if sku and sku in seen:
            raise ValueError('Monthly records must be deduplicated before calculation')
        seen.add(sku)
        original = known.get(sku)
        if original is not None and original.get('lead_days') != lead_days:
            raise ValueError('Monthly lead_days must match the API inbound eligibility horizon')
        row = deepcopy(original) if original is not None else dict(
            marketplace='US', api_marketplace='ATVPDKIKX0DER', sku=sku,
            asin=None, family='', fulfillable=None, reserved=None, unfulfillable=None,
            inbound_working=None, inbound_shipped=None, inbound_receiving=None,
            inventory_inbound=None, eligible_inbound=None, unknown_eta_inbound=None,
            inventory_complete=False, inbound_complete=False, data_complete=False,
            planning_data_complete=False, sales_status='UNKNOWN',
            status='DATA_REVIEW_REQUIRED', recommended_qty=None,
            anomalies=[], anomaly_codes=[])
        row['anomalies'] = list(row.get('anomalies') or [])
        row['anomaly_codes'] = list(row.get('anomaly_codes') or [])

        def issue(code, reason):
            if code not in row['anomaly_codes']:
                row['anomaly_codes'].append(code)
                row['anomalies'].append(dict(code=code, reason=reason, blocking=True,
                                             source='User monthly workbook validation'))

        units = _units(record.get('monthly_units'))
        valid = record.get('monthly_status') == 'VALID' and units is not None and bool(sku)
        if not valid:
            units = None
            issue('MONTHLY_SALES_INVALID', 'Monthly demand is missing, invalid, or conflicts across source rows')
        if original is None:
            issue('SKU_NOT_FOUND_US', 'Exact Seller SKU is absent from the verified US snapshot')
            for days in (7, 30, 60, 90):
                row.update({f'sales_{days}d': None, f'daily_sales_{days}d': None,
                            f'sales_complete_{days}d': False,
                            f'observed_sales_{days}d': None,
                            f'observed_sales_days_{days}d': 0,
                            f'sales_completeness_{days}d': 0})
        row.update(uploaded_monthly_units=units, monthly_daily_sales=units / monthly_days if valid else None,
                   monthly_days=monthly_days, demand_basis='USER_MONTHLY_WORKBOOK',
                   monthly_period_status=monthly_period_status,
                   monthly_status=record.get('monthly_status', 'INVALID'),
                   source_rows=list(record.get('source_rows') or []),
                   source_sheet=record.get('source_sheet'), source_sha256=record.get('source_sha256'),
                   monthly_issues=list(record.get('issues') or []),
                   api_30d_recommended_qty=original.get('recommended_qty') if original else None,
                   api_coverage_days=original.get('coverage_days') if original else None,
                   lead_days=lead_days, target_days=target_days, safety_days=safety_days,
                   moq=moq, case_pack=case_pack, parameters_confirmed=confirmed,
                   parameter_status='CONFIRMED' if confirmed else 'SCENARIO_ONLY')
        # These comparisons explain discrepancies for human review. They never
        # replace monthly demand, bypass a safety gate, or create hard blockers.
        references = {}
        for days in (7, 30, 90):
            quantity = _units(original.get(f'sales_{days}d')) if original else None
            references[days] = (quantity if original.get(f'sales_complete_{days}d') is True else None) if original else None
        row['api_7d_month_equivalent'] = references[7] * 30 / 7 if references[7] is not None else None
        row['api_90d_month_equivalent'] = references[90] / 3 if references[90] is not None else None
        monthly_equivalent = Fraction(units * 30) / Fraction(str(monthly_days)) if valid else None
        row['monthly_30d_equivalent'] = float(monthly_equivalent) if monthly_equivalent is not None else None
        row['api_30d_ratio'] = float(monthly_equivalent / references[30]) if valid and references[30] is not None and references[30] > 0 else None
        reviews = []
        if valid and references[30] is not None:
            if references[30] == 0:
                reviews.append('API30_ZERO_RATIO_UNDEFINED')
            if ((monthly_equivalent == 0) != (references[30] == 0)
                    or monthly_equivalent > 0 and references[30] > 0
                    and max(monthly_equivalent / references[30], references[30] / monthly_equivalent) >= 2):
                reviews.append('MONTHLY_DIFFERS_FROM_API30')
        if valid and references[7] is not None and references[7] > 0:
            recent_month = Fraction(references[7] * 30, 7)
            if recent_month >= 2 * monthly_equivalent:
                reviews.append('RECENT_7D_HIGHER_THAN_MONTHLY')
        row['reference_review_codes'] = reviews
        row['reference_review_threshold_ratio'] = 2
        row['daily_sales'] = row['monthly_daily_sales']
        row['recommended_qty'] = None
        safe = (original is not None and valid
                and original.get('status') in ('ACTIVE', 'SCENARIO_ONLY')
                and original.get('planning_data_complete') is True
                and original.get('inventory_complete') is True
                and original.get('inbound_complete') is True
                and _units(original.get('fulfillable')) is not None
                and _units(original.get('eligible_inbound')) is not None
                and not any(a.get('blocking') is True for a in row['anomalies']))
        if original is not None and original.get('status') == 'STOPPED':
            row.update(status='STOPPED', recommended_qty=0)
        elif safe:
            daily = Fraction(units) / Fraction(str(monthly_days))
            horizon = sum((Fraction(str(value)) for value in (lead_days, target_days, safety_days)), Fraction())
            raw = max(0, math.ceil(daily * horizon - row['fulfillable'] - row['eligible_inbound']))
            quantity = ((max(raw, moq) + case_pack - 1) // case_pack) * case_pack if raw else 0
            row.update(recommended_qty=quantity, status='ACTIVE' if confirmed else 'SCENARIO_ONLY', stop_reason='')
        else:
            row.update(status='DATA_REVIEW_REQUIRED', stop_reason='; '.join(row['anomaly_codes']) or 'API_SAFETY_GATE_NOT_PASSED')
        daily = row['monthly_daily_sales']
        stock_known = row.get('inventory_complete') is True and _units(row.get('fulfillable')) is not None
        row['coverage_days'] = round(row['fulfillable'] / daily, 2) if stock_known and daily is not None and daily > 0 else None
        if not stock_known or daily is None:
            risk = 'UNKNOWN_DATA'
        elif daily == 0:
            risk = 'NO_RECENT_DEMAND'
        elif row['coverage_days'] < lead_days:
            risk = 'STOCKOUT_BEFORE_REPLENISHMENT'
        elif row['coverage_days'] < lead_days + safety_days:
            risk = 'SAFETY_STOCK_RISK'
        else:
            risk = 'COVERED'
        row.update(stockout_risk=risk, replenishment_status=row['status'],
                   planning_data_complete=bool(safe),
                   procurement_eligible=bool(confirmed and row['status'] == 'ACTIVE'))
        result.append(row)
    return result
