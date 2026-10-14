"""Atomic, explicit human review decisions for frozen SKU collections.

This module performs only local SQLite changes. A manifest contains exact
marketplace/SellerSKU identities and original states; it never represents a
prefix, naming rule, or authorization to approve future children. Stop imports
are monotonic: blank, FALSE and unapproved rows cannot resume replenishment.
"""
import csv
import hashlib
import json
import uuid
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook

from .discovery import (
    _history, _record, _refresh_safety, _site, _stamp, _text,
    ensure_schema as ensure_discovery_schema,
)


SCHEMA = (
    '''CREATE TABLE IF NOT EXISTS batch_review_manifest (
        manifest_id TEXT PRIMARY KEY, marketplace TEXT NOT NULL,
        api_marketplace TEXT NOT NULL, family TEXT,
        member_sha TEXT NOT NULL, member_count INTEGER NOT NULL,
        created_at TEXT NOT NULL, created_by TEXT NOT NULL,
        batch_eligible INTEGER NOT NULL,
        UNIQUE(api_marketplace,family,member_sha))''',
    '''CREATE TABLE IF NOT EXISTS batch_review_member (
        manifest_id TEXT NOT NULL, sku TEXT NOT NULL, expected_state TEXT NOT NULL,
        PRIMARY KEY(manifest_id,sku))''',
    '''CREATE TABLE IF NOT EXISTS batch_review_decision (
        manifest_id TEXT PRIMARY KEY, at TEXT NOT NULL, reviewer TEXT NOT NULL,
        family TEXT, approved_count INTEGER NOT NULL,
        unchanged_count INTEGER NOT NULL, decision_type TEXT NOT NULL DEFAULT 'APPROVED',
        reason TEXT)''',
    '''CREATE TABLE IF NOT EXISTS batch_review_audit (
        id INTEGER PRIMARY KEY AUTOINCREMENT, at TEXT NOT NULL,
        reviewer TEXT NOT NULL, action TEXT NOT NULL, operation_id TEXT NOT NULL,
        marketplace TEXT NOT NULL, sku TEXT NOT NULL,
        original_values TEXT NOT NULL, changed_values TEXT NOT NULL)''',
    '''CREATE TABLE IF NOT EXISTS batch_scheduled_stop (
        id TEXT PRIMARY KEY, import_id TEXT NOT NULL, marketplace TEXT NOT NULL,
        api_marketplace TEXT NOT NULL, sku TEXT NOT NULL, family TEXT,
        reason TEXT NOT NULL, effective_date TEXT NOT NULL,
        reviewer TEXT NOT NULL, reviewed_at TEXT NOT NULL,
        status TEXT NOT NULL, applied_at TEXT, decision_sha TEXT NOT NULL UNIQUE,
        expected_mapping TEXT NOT NULL, effective_timezone TEXT NOT NULL)''',
)
EXPECTED_FIELDS = (
    'api_marketplace', 'marketplace', 'sku', 'asin', 'candidate_family',
    'candidate_families', 'family', 'effective_family', 'mapping_basis',
    'mapping_status', 'mapping_conflict', 'format_issue', 'presence_status',
    'replenishment_status', 'approved_by', 'approved_at', 'approval_evidence',
    'approved_asin', 'inventory_identities',
)


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def ensure_schema(con):
    ensure_discovery_schema(con)
    for statement in SCHEMA:
        con.execute(statement)
    columns = {row['name'] for row in con.execute('PRAGMA table_info(batch_review_decision)')}
    if 'decision_type' not in columns:
        con.execute("ALTER TABLE batch_review_decision ADD COLUMN decision_type TEXT NOT NULL DEFAULT 'APPROVED'")
    if 'reason' not in columns:
        con.execute('ALTER TABLE batch_review_decision ADD COLUMN reason TEXT')
    # Decisions and frozen review scopes are append-only even via direct SQL.
    for table in ('batch_review_manifest', 'batch_review_member',
                  'batch_review_decision', 'batch_review_audit'):
        for operation in ('UPDATE', 'DELETE'):
            con.execute(f'''CREATE TRIGGER IF NOT EXISTS {table}_no_{operation.lower()}
                BEFORE {operation} ON {table} BEGIN
                SELECT RAISE(ABORT,'Human review records are immutable'); END''')
    con.execute('''CREATE TRIGGER IF NOT EXISTS batch_review_member_no_late_insert
        BEFORE INSERT ON batch_review_member
        WHEN NOT EXISTS(SELECT 1 FROM batch_review_manifest WHERE manifest_id=NEW.manifest_id)
        OR (SELECT COUNT(*) FROM batch_review_member WHERE manifest_id=NEW.manifest_id)
            >= (SELECT member_count FROM batch_review_manifest WHERE manifest_id=NEW.manifest_id)
        OR EXISTS(SELECT 1 FROM batch_review_decision WHERE manifest_id=NEW.manifest_id)
        BEGIN SELECT RAISE(ABORT,'Approved review membership is immutable'); END''')


@contextmanager
def _atomic(con):
    """Keep caller transactions intact; roll back every row on any failure."""
    nested = con.in_transaction
    savepoint = 'batch_review_' + uuid.uuid4().hex
    con.execute('SAVEPOINT ' + savepoint if nested else 'BEGIN IMMEDIATE')
    try:
        yield
    except BaseException:
        if nested:
            con.execute('ROLLBACK TO ' + savepoint)
            con.execute('RELEASE ' + savepoint)
        else:
            con.rollback()
        raise
    else:
        con.execute('RELEASE ' + savepoint) if nested else con.commit()


def _expected_state(con, row):
    state = {key: row[key] for key in EXPECTED_FIELDS}
    for key in ('candidate_families', 'inventory_identities', 'approval_evidence'):
        if state[key] is not None:
            state[key] = json.loads(state[key])
    master = con.execute('SELECT family,asin,enabled FROM sku WHERE marketplace=? AND sku=?',
                         (row['marketplace'], row['sku'])).fetchone()
    state['master'] = dict(master) if master else None
    stop = con.execute('SELECT reason,created_at,active FROM stop_flag WHERE marketplace=? AND sku=?',
                       (row['marketplace'], row['sku'])).fetchone()
    state['stop'] = dict(stop) if stop else None
    families = set(state['candidate_families']) | {row['family'], row['effective_family']}
    state['family_stops'] = [dict(item) for item in con.execute(
        'SELECT family,trigger_sku,active FROM family_stop WHERE marketplace=? ORDER BY family,trigger_sku',
        (row['marketplace'],)) if item['family'] in families]
    return state


def _member_payload(marketplace, api_marketplace, family, members, batch_eligible):
    return {'marketplace': marketplace, 'api_marketplace': api_marketplace,
            'family': family, 'batch_eligible': bool(batch_eligible),
            'members': sorted(members, key=lambda item: item['sku'])}


def _sha(payload):
    return hashlib.sha256(_json(payload).encode('utf-8')).hexdigest()


def _sku(value):
    """SellerSKU is an exact identity, including any spaces Amazon returned."""
    if not isinstance(value, str) or not value.strip() or len(value) > 256 or any(
            ord(character) < 32 for character in value):
        raise ValueError('Missing or invalid review SKU')
    return value


def create_review_manifest(con, marketplace, api_marketplace, family, skus,
                           created_by='review-export', observed_at=None, *, batch_eligible=None):
    """Freeze an enumerated review scope without approving any mapping."""
    ensure_schema(con)
    family = _text(family, 'review target family') if family is not None else None
    if batch_eligible is not None and type(batch_eligible) is not bool:
        raise ValueError('Batch eligibility must be an explicit boolean')
    if batch_eligible is True and not family:
        raise ValueError('An unknown family cannot be eligible for batch approval')
    creator, timestamp = _text(created_by, 'manifest creator'), _stamp(observed_at)
    if not isinstance(skus, (list, tuple, set)) or not skus:
        raise ValueError('Review requires an explicit nonempty SKU collection')
    members = [_sku(sku) for sku in skus]
    if len(set(members)) != len(members):
        raise ValueError('Review scope contains duplicate SKUs')
    with _atomic(con):
        marketplace, api_marketplace = _site(con, marketplace, api_marketplace)
        snapshots = []
        for sku in sorted(members):
            row = _record(con, marketplace, sku)
            if row['api_marketplace'] != api_marketplace:
                raise ValueError('Review member belongs to another API marketplace')
            snapshots.append({'sku': sku, 'expected_state': _expected_state(con, row)})
        eligible = bool(family) and (batch_eligible is True or (batch_eligible is None and all(
            member['expected_state']['family'] == family or
            member['expected_state']['candidate_family'] == family for member in snapshots)))
        member_sha = _sha(_member_payload(marketplace, api_marketplace, family, snapshots, eligible))
        existing = con.execute('''SELECT * FROM batch_review_manifest
            WHERE api_marketplace=? AND family IS ? AND member_sha=?''',
                               (api_marketplace, family, member_sha)).fetchone()
        if existing:
            return dict(existing)
        manifest_id = uuid.uuid4().hex
        con.execute('INSERT INTO batch_review_manifest VALUES(?,?,?,?,?,?,?,?,?)',
                    (manifest_id, marketplace, api_marketplace, family, member_sha,
                     len(snapshots), timestamp, creator, int(eligible)))
        for member in snapshots:
            con.execute('INSERT INTO batch_review_member VALUES(?,?,?)',
                        (manifest_id, member['sku'], _json(member['expected_state'])))
    return {'manifest_id': manifest_id, 'marketplace': marketplace,
            'api_marketplace': api_marketplace, 'family': family,
            'member_sha': member_sha, 'member_count': len(snapshots),
            'created_at': timestamp, 'created_by': creator, 'batch_eligible': eligible}


def review_manifest_members(con, manifest_id):
    ensure_schema(con)
    manifest_id = _text(manifest_id, 'review manifest ID')
    if not con.execute('SELECT 1 FROM batch_review_manifest WHERE manifest_id=?',
                       (manifest_id,)).fetchone():
        raise ValueError('Unknown review manifest')
    return [{'sku': row['sku'], 'expected_state': json.loads(row['expected_state'])}
            for row in con.execute('''SELECT sku,expected_state FROM batch_review_member
                WHERE manifest_id=? ORDER BY sku''', (manifest_id,))]


def _review_audit(con, timestamp, reviewer, action, operation_id, marketplace, sku,
                  original, changed):
    con.execute('''INSERT INTO batch_review_audit
        (at,reviewer,action,operation_id,marketplace,sku,original_values,changed_values)
        VALUES(?,?,?,?,?,?,?,?)''',
                (timestamp, reviewer, action, operation_id, marketplace, sku,
                 _json(original), _json(changed)))
    con.execute('INSERT INTO audit(at,action,marketplace,sku,detail) VALUES(?,?,?,?,?)',
                (timestamp, action, marketplace, sku,
                 _json({'reviewer': reviewer, 'operation_id': operation_id,
                        'original': original, 'changed': changed})))


def _validate_identity(row):
    identities = json.loads(row['inventory_identities'])
    if not row['asin'] or not identities or {item.get('asin') for item in identities} != {row['asin']}:
        raise ValueError('Ambiguous or incomplete SKU/ASIN identity cannot be batch approved')


def approve_batch(con, manifest_id, member_sha, family, reviewer, *, confirm=False,
                  allow_remap=False, resolve_conflicts=False, observed_at=None, reason=None,
                  refresh_safety=True):
    """Approve only frozen members after explicit confirmation and stale checks.

    Existing approved mappings to the same family are retained unchanged. A
    different formal family requires allow_remap=True; conflicts separately
    require resolve_conflicts=True and never permit ambiguous SKU/ASIN identity.
    """
    if type(refresh_safety) is not bool:
        raise ValueError('Safety refresh flag must be an explicit boolean')
    if not refresh_safety and not con.in_transaction:
        raise ValueError('Deferred safety refresh requires an existing caller transaction')
    ensure_schema(con)
    if confirm is not True:
        raise ValueError('Explicit human confirmation is required for batch approval')
    if type(allow_remap) is not bool or type(resolve_conflicts) is not bool:
        raise ValueError('Remap/conflict confirmations must be explicit booleans')
    manifest_id = _text(manifest_id, 'review manifest ID')
    family, reviewer = _text(family, 'approved family'), _text(reviewer, 'reviewer')
    reason = _text(reason, 'approval reason') if reason is not None else None
    timestamp = _stamp(observed_at)
    with _atomic(con):
        manifest = con.execute('SELECT * FROM batch_review_manifest WHERE manifest_id=?',
                               (manifest_id,)).fetchone()
        if manifest is None:
            raise ValueError('Unknown review manifest')
        if not manifest['batch_eligible'] or not manifest['family']:
            raise ValueError('Format-only or unknown family groups cannot be batch approved')
        if manifest['family'] != family or manifest['member_sha'] != member_sha:
            raise ValueError('Review family or membership hash does not match the frozen manifest')
        if con.execute('SELECT 1 FROM batch_review_decision WHERE manifest_id=?', (manifest_id,)).fetchone():
            raise ValueError('Review manifest already has an immutable decision')
        _site(con, manifest['marketplace'], manifest['api_marketplace'])
        members = review_manifest_members(con, manifest_id)
        payload = _member_payload(manifest['marketplace'], manifest['api_marketplace'], family,
                                  members, manifest['batch_eligible'])
        if len(members) != manifest['member_count'] or _sha(payload) != member_sha:
            raise ValueError('Review manifest membership was altered')
        rows = []
        # Validate the whole scope before performing the first mapping write.
        for member in members:
            row = _record(con, manifest['marketplace'], member['sku'])
            if _expected_state(con, row) != member['expected_state']:
                raise ValueError('Stale review scope; regenerate the manifest before approving')
            _validate_identity(row)
            master = member['expected_state']['master']
            if not master or master['asin'] != row['asin']:
                raise ValueError('Core SKU/ASIN identity conflicts with discovery; individual review required')
            existing_families = {row['family'], master['family']} - {None, ''}
            if any(value != family for value in existing_families) and not allow_remap:
                raise ValueError('Existing formal mapping requires explicit remap confirmation')
            if (row['mapping_conflict'] or row['mapping_status'] == 'FAMILY_CONFLICT') and not resolve_conflicts:
                raise ValueError('Family conflict requires explicit conflict-resolution confirmation')
            rows.append((row, member['expected_state']))
        approved_count, unchanged_count = 0, 0
        for row, original in rows:
            sku = row['sku']
            if row['mapping_status'] == 'APPROVED' and row['family'] == family:
                unchanged_count += 1
                _review_audit(con, timestamp, reviewer, 'BATCH_MAPPING_NO_CHANGE', manifest_id,
                              manifest['marketplace'], sku, original, {})
                continue
            con.execute('''UPDATE discovery_sku SET family=?,effective_family=?,mapping_status='APPROVED',
                mapping_conflict=0,approved_by=?,approved_at=?,approval_evidence=?,approved_asin=?
                WHERE api_marketplace=? AND sku=?''',
                        (family, family, reviewer, timestamp, row['candidate_families'], row['asin'],
                         manifest['api_marketplace'], sku))
            con.execute('UPDATE sku SET family=? WHERE marketplace=? AND sku=?',
                        (family, manifest['marketplace'], sku))
            if refresh_safety and con.execute('SELECT 1 FROM stop_flag WHERE marketplace=? AND sku=? AND active=1',
                           (manifest['marketplace'], sku)).fetchone():
                con.execute('INSERT OR REPLACE INTO family_stop VALUES(?,?,?,1)',
                            (manifest['marketplace'], family, sku))
            con.execute("DELETE FROM data_quality_block WHERE marketplace=? AND sku=? AND code='MAPPING_UNAPPROVED'",
                        (manifest['marketplace'], sku))
            changed = {'family': family, 'mapping_status': 'APPROVED',
                       'approved_by': reviewer, 'approved_at': timestamp,
                       'mapping_conflict': 0, 'approved_asin': row['asin']}
            if reason is not None:
                changed['approval_reason'] = reason
            _review_audit(con, timestamp, reviewer, 'BATCH_MAPPING_APPROVED', manifest_id,
                          manifest['marketplace'], sku, original, changed)
            _history(con, timestamp, manifest['api_marketplace'], manifest['marketplace'], sku,
                     'BATCH_MAPPING_APPROVED', {'reviewer': reviewer, 'manifest_id': manifest_id,
                                               'original': original, 'changed': changed})
            approved_count += 1
        if refresh_safety:
            _refresh_safety(con, manifest['marketplace'])
        con.execute('''INSERT INTO batch_review_decision
            (manifest_id,at,reviewer,family,approved_count,unchanged_count,decision_type,reason)
            VALUES(?,?,?,?,?,?,'APPROVED',?)''',
                    (manifest_id, timestamp, reviewer, family, approved_count, unchanged_count, reason))
    return {'manifest_id': manifest_id, 'member_sha': member_sha,
            'marketplace': manifest['marketplace'], 'family': family,
            'approved_count': approved_count, 'unchanged_count': unchanged_count,
            'member_count': len(members), 'reviewer': reviewer, 'reviewed_at': timestamp}


def reject_batch(con, manifest_id, member_sha, reviewer, reason, *, confirm=False,
                 observed_at=None, refresh_safety=True):
    """Persist rejection of a frozen scope without remapping approved children."""
    if type(refresh_safety) is not bool:
        raise ValueError('Safety refresh flag must be an explicit boolean')
    if not refresh_safety and not con.in_transaction:
        raise ValueError('Deferred safety refresh requires an existing caller transaction')
    ensure_schema(con)
    if confirm is not True:
        raise ValueError('Explicit human confirmation is required for batch rejection')
    manifest_id = _text(manifest_id, 'review manifest ID')
    reviewer, reason = _text(reviewer, 'reviewer'), _text(reason, 'rejection reason')
    timestamp = _stamp(observed_at)
    with _atomic(con):
        manifest = con.execute('SELECT * FROM batch_review_manifest WHERE manifest_id=?',
                               (manifest_id,)).fetchone()
        if manifest is None:
            raise ValueError('Unknown review manifest')
        if member_sha != manifest['member_sha']:
            raise ValueError('Membership hash does not match the frozen manifest')
        if con.execute('SELECT 1 FROM batch_review_decision WHERE manifest_id=?', (manifest_id,)).fetchone():
            raise ValueError('Review manifest already has an immutable decision')
        _site(con, manifest['marketplace'], manifest['api_marketplace'])
        members = review_manifest_members(con, manifest_id)
        payload = _member_payload(manifest['marketplace'], manifest['api_marketplace'], manifest['family'],
                                  members, manifest['batch_eligible'])
        if len(members) != manifest['member_count'] or _sha(payload) != member_sha:
            raise ValueError('Review manifest membership was altered')
        snapshots = []
        for member in members:
            row = _record(con, manifest['marketplace'], member['sku'])
            if _expected_state(con, row) != member['expected_state']:
                raise ValueError('Stale review scope; regenerate the manifest before rejecting')
            snapshots.append((row, member['expected_state']))
        retained_approved = 0
        for row, original in snapshots:
            changed = {'review_decision': 'REJECTED', 'reason': reason}
            # Keep the formal mapping, approver, and active stops intact.
            if row['mapping_status'] == 'APPROVED' or row['approved_by']:
                retained_approved += 1
            elif not row['mapping_conflict']:
                con.execute("UPDATE discovery_sku SET mapping_status='REVIEW_REQUIRED' WHERE api_marketplace=? AND sku=?",
                            (manifest['api_marketplace'], row['sku']))
                changed['mapping_status'] = 'REVIEW_REQUIRED'
            _review_audit(con, timestamp, reviewer, 'BATCH_MAPPING_REJECTED', manifest_id,
                          manifest['marketplace'], row['sku'], original, changed)
            _history(con, timestamp, manifest['api_marketplace'], manifest['marketplace'], row['sku'],
                     'BATCH_MAPPING_REJECTED', {'reviewer': reviewer, 'original': original, 'changed': changed})
        if refresh_safety:
            _refresh_safety(con, manifest['marketplace'])
        con.execute('''INSERT INTO batch_review_decision
            (manifest_id,at,reviewer,family,approved_count,unchanged_count,decision_type,reason)
            VALUES(?,?,?,?,0,?,'REJECTED',?)''',
                    (manifest_id, timestamp, reviewer, manifest['family'], retained_approved, reason))
    return {'manifest_id': manifest_id, 'member_sha': member_sha, 'decision_type': 'REJECTED',
            'member_count': len(members), 'rejected_count': len(members),
            'retained_approved_count': retained_approved, 'reviewer': reviewer,
            'reviewed_at': timestamp, 'reason': reason}


STOP_COLUMN_ALIASES = {
    'sku': ('sku', 'SKU'), 'family': ('family', '系列', '正式系列'),
    'marketplace': ('marketplace', 'Marketplace', '站点'),
    'api_marketplace': ('api_marketplace', 'Marketplace ID', 'API Marketplace'),
    'stop': ('stop', 'is_stopped', '是否停补'),
    'reason': ('reason', 'stop_reason', '停补原因'),
    'effective_date': ('effective_date', '生效日期'),
    'review_status': ('review_status', '审核状态'),
    'reviewer': ('reviewer', '审核人'),
}


def normalize_stop_row(row):
    """Normalize headers only; unknown decision values never become FALSE."""
    result = {}
    for key, names in STOP_COLUMN_ALIASES.items():
        values = [row[name] for name in names if name in row and row[name] not in (None, '')]
        if len({_json(value.isoformat() if isinstance(value, (date, datetime)) else value)
                for value in values}) > 1:
            raise ValueError('Conflicting stop-list column aliases')
        result[key] = values[0] if values else None
    status = result['review_status']
    if status is None or status == '':
        result['review_status'] = 'UNAPPROVED'
    elif not isinstance(status, str):
        raise ValueError('Invalid stop-list review status')
    else:
        result['review_status'] = status.strip().upper()
    value = result['stop']
    if value is None or (isinstance(value, str) and not value.strip()):
        result['stop'] = 'NO_CHANGE'
    elif value is True or (isinstance(value, str) and value.strip().upper() in ('TRUE', 'YES', '是', 'STOP')):
        result['stop'] = 'TRUE'
    elif value is False or (isinstance(value, str) and value.strip().upper() in ('FALSE', 'NO', '否', 'NO_CHANGE')):
        result['stop'] = 'NO_CHANGE'
    else:
        raise ValueError('Unknown stop decision; use explicit TRUE or NO_CHANGE')
    return result


def _read_stop_rows(path):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                raise ValueError('Missing or duplicate stop-list headers')
            return list(reader)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Stop list must be an Excel .xlsx or CSV file')
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        candidates = [sheet for sheet in workbook if sheet.title in ('停补导入', '停止补货', '停补名单', 'Stop List')]
        sheet = candidates[0] if candidates else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, ())
        if not headers or len(set(headers)) != len(headers):
            raise ValueError('Missing or duplicate stop-list headers')
        result = []
        for values in iterator:
            if not any(value is not None for value in values):
                continue
            if any(isinstance(value, str) and value.startswith('=') for value in values):
                raise ValueError('Stop-list formulas cannot authorize a human decision')
            result.append(dict(zip(headers, values)))
        return result
    finally:
        workbook.close()


def _effective_day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str) or not value.strip():
        raise ValueError('Approved stop requires an effective date')
    try:
        return date.fromisoformat(value.strip()).isoformat()
    except ValueError:
        raise ValueError('Invalid stop effective date; use YYYY-MM-DD') from None


def _stop_original(con, marketplace, sku, family):
    old = con.execute('SELECT * FROM stop_flag WHERE marketplace=? AND sku=?',
                      (marketplace, sku)).fetchone()
    family_old = [] if not family else [dict(row) for row in con.execute(
        'SELECT * FROM family_stop WHERE marketplace=? AND family=? ORDER BY trigger_sku',
        (marketplace, family))]
    return {'stop_flag': dict(old) if old else None, 'family_stops': family_old}


def _stop_mapping(con, row):
    master = con.execute('SELECT family,asin,enabled FROM sku WHERE marketplace=? AND sku=?',
                         (row['marketplace'], row['sku'])).fetchone()
    direct = con.execute('SELECT reason,created_at,active FROM stop_flag WHERE marketplace=? AND sku=?',
                         (row['marketplace'], row['sku'])).fetchone()
    return {key: row[key] for key in ('api_marketplace', 'marketplace', 'sku', 'asin', 'family',
                                      'mapping_status', 'approved_asin', 'approved_by', 'approved_at')} | {
        'master': dict(master) if master else None, 'direct_stop': dict(direct) if direct else None}


def _stop_timezone(con, marketplace):
    binding = con.execute('SELECT sales_timezone FROM marketplace_binding WHERE marketplace=?',
                          (marketplace,)).fetchone()
    name = (binding['sales_timezone'] if binding else None) or 'UTC'
    try:
        ZoneInfo(name)
    except (ZoneInfoNotFoundError, TypeError):
        raise ValueError('Invalid marketplace timezone for stop effective date') from None
    return name


def _apply_stop(con, scheduled, timestamp):
    marketplace, sku, family = scheduled['marketplace'], scheduled['sku'], scheduled['family']
    original = _stop_original(con, marketplace, sku, family)
    current = original['stop_flag']
    # Preserve an existing manual reason and its creation time.
    if not current or not current['active']:
        con.execute('''INSERT INTO stop_flag VALUES(?,?,?,?,1)
            ON CONFLICT(marketplace,sku) DO UPDATE SET reason=excluded.reason,
            created_at=excluded.created_at,active=1''',
                    (marketplace, sku, scheduled['reason'], timestamp))
    if family:
        con.execute('''INSERT INTO family_stop VALUES(?,?,?,1)
            ON CONFLICT(marketplace,family,trigger_sku) DO UPDATE SET active=1''',
                    (marketplace, family, sku))
    con.execute("UPDATE batch_scheduled_stop SET status='ACTIVE',applied_at=? WHERE id=?",
                (timestamp, scheduled['id']))
    changed = {'stop': 'TRUE', 'family': family, 'reason': scheduled['reason'],
               'effective_date': scheduled['effective_date'], 'status': 'ACTIVE'}
    _review_audit(con, timestamp, scheduled['reviewer'], 'STOP_IMPORT_APPLIED',
                  scheduled['import_id'], marketplace, sku, original, changed)
    _history(con, timestamp, scheduled['api_marketplace'], marketplace, sku,
             'STOP_IMPORT_APPLIED', {'reviewer': scheduled['reviewer'], 'original': original,
                                     'changed': changed})


def _pending_stop_conflict(con, scheduled, timestamp, current):
    con.execute("UPDATE batch_scheduled_stop SET status='CONFLICT_PENDING' WHERE id=?", (scheduled['id'],))
    changed = {'status': 'CONFLICT_PENDING', 'reason': 'Approved mapping changed before stop effective date',
               'current_mapping': current}
    _review_audit(con, timestamp, scheduled['reviewer'], 'STOP_IMPORT_SCOPE_CONFLICT',
                  scheduled['import_id'], scheduled['marketplace'], scheduled['sku'],
                  json.loads(scheduled['expected_mapping']), changed)
    # Preserve a conservative local review gate on the reviewed series and
    # trigger. A remap cannot turn a scheduled stop into automatic permission.
    skus = {scheduled['sku']} | {row['sku'] for row in con.execute(
        "SELECT sku FROM discovery_sku WHERE marketplace=? AND family=? AND mapping_status='APPROVED'",
        (scheduled['marketplace'], scheduled['family']))}
    for sku in skus:
        con.execute('INSERT OR REPLACE INTO data_quality_block VALUES(?,?,?,?)',
                    (scheduled['marketplace'], sku, 'SCHEDULED_STOP_SCOPE_CONFLICT',
                     'A previously reviewed stop changed scope; explicit human review required'))


def refresh_pending_stops(con, observed_at=None, manage_transaction=True):
    """Apply due, previously reviewed stops; never replay resumed active stops."""
    ensure_schema(con)
    timestamp = _stamp(observed_at)

    def apply():
        due = [dict(row) for row in con.execute('''SELECT * FROM batch_scheduled_stop
            WHERE status='PENDING' ORDER BY effective_date,id''') if row['effective_date'] <=
            datetime.fromisoformat(timestamp).astimezone(ZoneInfo(row['effective_timezone'])).date().isoformat()]
        applied_count, conflict_count = 0, 0
        for scheduled in due:
            target = con.execute('SELECT * FROM discovery_sku WHERE marketplace=? AND sku=?',
                                 (scheduled['marketplace'], scheduled['sku'])).fetchone()
            current = _stop_mapping(con, target) if target else None
            if current != json.loads(scheduled['expected_mapping']):
                _pending_stop_conflict(con, scheduled, timestamp, current)
                conflict_count += 1
                continue
            _apply_stop(con, scheduled, timestamp)
            applied_count += 1
        for marketplace in {row['marketplace'] for row in due}:
            _refresh_safety(con, marketplace)
        return {'applied_count': applied_count, 'conflict_count': conflict_count, 'at': timestamp}

    if manage_transaction:
        with _atomic(con):
            return apply()
    if not con.in_transaction:
        raise ValueError('Caller must open a transaction before refreshing scheduled stops')
    return apply()


def import_stop_list(con, path, reviewer, *, observed_at=None):
    """Validate a full file before atomically importing explicit human stops.

    Unapproved rows and FALSE/blank rows are audited NO_CHANGE, including for
    currently stopped SKUs. Family text alone never approves a SKU mapping.
    """
    ensure_schema(con)
    reviewer, timestamp = _text(reviewer, 'stop reviewer'), _stamp(observed_at)
    source_rows = _read_stop_rows(path)
    rows = [normalize_stop_row(row) for row in source_rows]
    import_id = uuid.uuid4().hex
    planned, ignored, seen, duplicate_count = [], [], set(), 0
    with _atomic(con):
        for row in rows:
            marketplace = _text(row['marketplace'], 'stop marketplace')
            sku = _sku(row['sku']) if row['sku'] else None
            family = _text(row['family'], 'stop family') if row['family'] else None
            if row['review_status'] != 'APPROVED' or row['stop'] != 'TRUE':
                ignored.append((marketplace, sku or '', row))
                continue
            reason = _text(row['reason'], 'approved stop reason')
            effective = _effective_day(row['effective_date'])
            if row['reviewer'] and _text(row['reviewer'], 'row reviewer') != reviewer:
                raise ValueError('Stop row reviewer differs from the explicit importing reviewer')
            if sku:
                target = _record(con, marketplace, sku)
                if target['mapping_status'] != 'APPROVED' or not target['family']:
                    raise ValueError('Stop import requires an approved formal family mapping')
                if family and (target['mapping_status'] != 'APPROVED' or target['family'] != family):
                    raise ValueError('Stop family must match an approved SKU mapping')
                family = family or (target['family'] if target['mapping_status'] == 'APPROVED' else None)
            elif family:
                target = con.execute('''SELECT * FROM discovery_sku
                    WHERE marketplace=? AND family=? AND mapping_status='APPROVED'
                    ORDER BY sku LIMIT 1''', (marketplace, family)).fetchone()
                if target is None:
                    raise ValueError('Series stop requires at least one approved family member')
                target = dict(target)
                sku = target['sku']
            else:
                raise ValueError('Approved stop requires a discovered SKU or confirmed family')
            api_marketplace = target['api_marketplace']
            _site(con, marketplace, api_marketplace)
            if row['api_marketplace'] and row['api_marketplace'] != api_marketplace:
                raise ValueError('Stop row API marketplace conflicts with the discovered SKU')
            identity = (marketplace, sku, effective)
            if identity in seen:
                raise ValueError('Duplicate stop decisions in the same import')
            seen.add(identity)
            decision_sha = _sha({'marketplace': marketplace, 'api_marketplace': api_marketplace,
                                 'sku': sku, 'family': family, 'reason': reason,
                                 'effective_date': effective, 'reviewer': reviewer})
            if con.execute('SELECT 1 FROM batch_scheduled_stop WHERE decision_sha=?',
                           (decision_sha,)).fetchone():
                duplicate_count += 1
                ignored.append((marketplace, sku, row))
                continue
            planned.append({'id': uuid.uuid4().hex, 'import_id': import_id,
                            'marketplace': marketplace, 'api_marketplace': api_marketplace,
                            'sku': sku, 'family': family, 'reason': reason,
                            'effective_date': effective, 'reviewer': reviewer,
                            'reviewed_at': timestamp, 'status': 'PENDING', 'applied_at': None,
                            'decision_sha': decision_sha, 'expected_mapping': _json(_stop_mapping(con, target)),
                            'effective_timezone': _stop_timezone(con, marketplace)})
        for item in planned:
            columns = ','.join(item)
            con.execute('INSERT INTO batch_scheduled_stop(' + columns + ') VALUES(' +
                        ','.join('?' for _ in item) + ')', tuple(item.values()))
            original = _stop_original(con, item['marketplace'], item['sku'], item['family'])
            _review_audit(con, timestamp, reviewer, 'STOP_IMPORT_REVIEWED', import_id,
                          item['marketplace'], item['sku'], original,
                          {'stop': 'TRUE', 'family': item['family'], 'reason': item['reason'],
                           'effective_date': item['effective_date'], 'effective_timezone': item['effective_timezone'],
                           'status': 'PENDING'})
        for marketplace, sku, row in ignored:
            original = _stop_original(con, marketplace, sku, row['family'])
            _review_audit(con, timestamp, reviewer, 'STOP_IMPORT_NO_CHANGE', import_id,
                          marketplace, sku, original,
                          {'decision': 'NO_CHANGE', 'submitted_stop': row['stop'],
                           'submitted_review_status': row['review_status']})
        applied = refresh_pending_stops(con, timestamp, manage_transaction=False)['applied_count']
        pending_count = con.execute("SELECT COUNT(*) FROM batch_scheduled_stop WHERE import_id=? AND status='PENDING'",
                                    (import_id,)).fetchone()[0]
    return {'import_id': import_id, 'row_count': len(rows), 'reviewed_stop_count': len(planned),
            'applied_count': applied, 'pending_count': pending_count,
            'no_change_count': len(ignored), 'duplicate_count': duplicate_count,
            'reviewer': reviewer, 'reviewed_at': timestamp}


def review_audit_rows(con, marketplace=None):
    ensure_schema(con)
    query = 'SELECT * FROM batch_review_audit'
    args = ()
    if marketplace is not None:
        query += ' WHERE marketplace=?'
        args = (marketplace,)
    result = []
    for row in con.execute(query + ' ORDER BY id', args):
        item = dict(row)
        item['original_values'] = json.loads(item['original_values'])
        item['changed_values'] = json.loads(item['changed_values'])
        result.append(item)
    return result
