"""Read explicit SKU/model columns as unapproved reference data, never instructions."""
import hashlib
from pathlib import Path

from openpyxl import load_workbook


def load_sku_reference(path):
    source = Path(path)
    wb = load_workbook(source, read_only=True, data_only=False)
    try:
        if '全部SKU' not in wb.sheetnames:
            raise ValueError('SKU reference requires the 全部SKU sheet')
        rows = wb['全部SKU'].iter_rows(values_only=True)
        header = list(next(rows))
        if not {'原始SKU', '产品型号'}.issubset(header):
            raise ValueError('SKU reference is missing exact SKU/model columns')
        sku_index, family_index = header.index('原始SKU'), header.index('产品型号')
        format_index = header.index('格式检查') if '格式检查' in header else None
        candidates, row_count, flagged = {}, 0, 0
        for row in rows:
            sku, family = row[sku_index], row[family_index]
            if sku is None and family is None:
                continue
            if not isinstance(sku, str) or not sku.strip() or sku.startswith('='):
                raise ValueError('SKU reference has an invalid or formula SKU')
            if family is not None and (not isinstance(family, str) or family.startswith('=')):
                raise ValueError('SKU reference has an invalid or formula model')
            family = family.strip() if isinstance(family, str) and family.strip() else None
            issue = row[format_index] if format_index is not None else None
            if issue is not None and not isinstance(issue, str):
                raise ValueError('SKU reference has invalid format metadata')
            row_count += 1
            flagged += bool(issue)
            prior = candidates.get(sku)
            if prior and prior['candidate_family'] != family:
                choices = set(prior.get('conflicting_families', []))
                choices.update(f for f in (prior['candidate_family'], family) if f)
                candidates[sku] = {'candidate_family': None, 'basis': 'USER_WORKBOOK_CONFLICT',
                                   'format_issue': issue, 'conflicting_families': sorted(choices)}
            elif not prior:
                candidates[sku] = {'candidate_family': family, 'basis': 'USER_WORKBOOK_EXACT_SKU',
                                   'format_issue': issue}
        return candidates, {'source_name': source.name, 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
                            'source_rows': row_count, 'source_unique_skus': len(candidates),
                            'source_models': len({r['candidate_family'] for r in candidates.values() if r['candidate_family']}),
                            'source_format_warnings': flagged,
                            'source_conflicts': sum(bool(r.get('conflicting_families')) for r in candidates.values()),
                            'approval_status': 'REFERENCE_ONLY_UNAPPROVED'}
    finally:
        wb.close()
