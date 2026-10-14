"""Auditable Excel workbook containing actual planning inputs, never formulas/PII."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = ['marketplace', 'sku', 'family', 'asin', 'sales_30d', 'observed_sales_days',
           'daily_sales', 'fulfillable', 'reserved', 'unfulfillable', 'stock_observed_at',
           'stock_age_days', 'eligible_inbound', 'unknown_eta_inbound', 'coverage_days',
           'status', 'stop_reason', 'recommended_qty', 'asof', 'window_start',
           'lead_days', 'target_days', 'safety_days', 'moq', 'case_pack', 'sales_timezone']


def sheet(wb, name, columns, rows):
    ws = wb.create_sheet(name)
    ws.append(columns)
    for row in rows:
        ws.append([row.get(key) for key in columns])
        # Treat all external strings as literal text, preventing formula injection.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
    for index, key in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = max(16, min(42, len(key) + 6))
    return ws


def export(rows, path, con=None, marketplace=None, settings=None):
    rows = list(rows)
    marketplace = marketplace or (rows[0]['marketplace'] if rows else 'US')
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, 'Replenishment', COLUMNS, rows)
    alerts = [dict(row) for row in rows if row['status'] != 'ACTIVE']
    sheet(wb, 'Alerts', ['sku', 'family', 'status', 'stop_reason', 'recommended_qty'], alerts)
    tables = {
        'Sales': ('sales', ['marketplace', 'sku', 'day', 'units']),
        'Inventory': ('stock', ['marketplace', 'sku', 'fulfillable', 'reserved', 'unfulfillable', 'observed_at']),
        'Inbound': ('inbound', ['marketplace', 'sku', 'shipment_id', 'units', 'eta', 'status']),
        'Orders': ('order_line', ['marketplace', 'order_id', 'item_id', 'sku', 'day', 'status', 'units']),
        'Sync Log': ('sync_log', ['id', 'at', 'marketplace', 'kind', 'status', 'row_count', 'detail']),
        'Audit': ('audit', ['id', 'at', 'action', 'marketplace', 'sku', 'detail']),
    }
    for name, (table, columns) in tables.items():
        data = [dict(r) for r in con.execute(f'SELECT * FROM {table} WHERE marketplace=?', (marketplace,))] if con is not None else []
        sheet(wb, name, columns, data)
    parameters = dict(settings or {})
    if rows:
        for key in ['asof', 'window_start', 'lead_days', 'target_days', 'safety_days', 'moq', 'case_pack', 'sales_timezone']:
            if parameters.get(key) is None:
                parameters[key] = rows[0][key]
    parameters.update({'marketplace': marketplace, 'mode': 'READ_ONLY',
                       'sales_basis': 'Daily gross ordered AFN units; canceled orders excluded; not net returns',
                       'formula': 'ceil to case pack(max(MOQ, daily_sales * horizon - fulfillable - eligible_inbound)) if positive',
                       'inbound_policy': 'Unknown ETA blocks planning; terminal/working shipments excluded',
                       'snapshot_history': 'SQLite snapshots table; no customer PII stored'})
    sheet(wb, 'Parameters', ['parameter', 'value'], [{'parameter': k, 'value': v} for k, v in parameters.items()])
    wb.save(path)
