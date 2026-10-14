"""Independent-SKU, read-only Excel report with explicit unknowns and provenance.

The family field is descriptive in this report. This exporter never invokes the
family planner, propagates family stops, writes to SQLite, or certifies an API
response merely because a sync was marked successful.
"""
from collections.abc import Mapping
from datetime import datetime
import math
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from .export import sheet

WINDOWS = (7, 30, 60, 90)
STOCK_FIELDS = ('fulfillable', 'reserved', 'unfulfillable')
INBOUND_FIELDS = ('inbound_working', 'inbound_shipped', 'inbound_receiving')
INVENTORY_FIELDS = STOCK_FIELDS + INBOUND_FIELDS
PARAMETER_KEYS = (
    'asof', 'sales_timezone', 'lead_days', 'target_days', 'safety_days',
    'moq', 'case_pack', 'max_stock_age_days', 'inbound_rule',
    'parameters_confirmed', 'confirmed_by', 'confirmed_at',
)
SUMMARY_AUDIT_KEYS = {
    'collection_status', 'data_source', 'window_start', 'window_end', 'asof',
    'orders_collected', 'orders_canceled', 'order_lines_collected',
    'inbound_shipments_collected', 'inbound_v2024_status', 'inbound_v0_status',
    'quality_gate', 'sample_orders', 'sample_order_lines', 'sample_inventory_skus',
    'inventory_pages', 'orders_pages', 'deduplicated_orders', 'deduplicated_lines',
}
AUDIT_NUMBERS = {
    'orders', 'items', 'deduplicated', 'row_count', 'pages', 'shipments',
    'shipment_items', 'rows', 'skus', 'retries', 'count',
}
AUDIT_STRINGS = {'stage', 'kind', 'status', 'source', 'code'}
AUDIT_DATES = {'at', 'begin', 'finish', 'window_start', 'window_end', 'observed_at'}
AUDIT_COLUMNS = (
    'stage', 'kind', 'status', 'source', 'code', 'at', 'begin', 'finish',
    'observed_at', 'window_start', 'window_end', 'row_count', 'rows', 'skus',
    'orders', 'items', 'shipments', 'shipment_items', 'pages', 'deduplicated',
    'retries', 'count', 'complete',
)
IDENTITY_COLUMNS = ('marketplace', 'sku', 'asin', 'family')
WINDOW_COLUMNS = tuple(
    f'{name}_{days}d'
    for days in WINDOWS
    for name in ('sales', 'daily_sales', 'observed_sales', 'observed_sales_days',
                 'sales_complete', 'sales_completeness')
)
MAIN_COLUMNS = IDENTITY_COLUMNS + (
    'status', 'parameter_status', 'sales_status', 'sales_30d', 'daily_sales',
    'fulfillable', 'reserved', 'unfulfillable', 'inventory_inbound',
    'eligible_inbound', 'unknown_eta_inbound', 'coverage_days', 'stockout_risk',
    'recommended_qty', 'recommendation_label', 'anomaly_codes', 'stop_reason',
    'sales_7d', 'sales_60d', 'sales_90d', 'daily_sales_7d', 'daily_sales_60d',
    'daily_sales_90d', 'inventory_complete', 'inbound_complete', 'asof',
    'window_start', 'stock_observed_at', 'lead_days', 'target_days', 'safety_days',
    'inbound_rule', 'moq', 'case_pack', 'sales_timezone',
    'sales_source', 'inventory_source', 'inbound_source',
)
RECOMMENDATION_COLUMNS = IDENTITY_COLUMNS + (
    'recommended_qty', 'recommendation_label', 'status', 'parameter_status',
    'sales_30d', 'daily_sales', 'fulfillable', 'eligible_inbound',
    'inventory_inbound', 'unknown_eta_inbound', 'coverage_days', 'lead_days',
    'target_days', 'safety_days', 'inbound_rule', 'moq', 'case_pack', 'stop_reason',
)
ANOMALY_COLUMNS = ('marketplace', 'sku', 'asin', 'code', 'reason', 'blocking', 'source')

COMMENTS = {
    'family': '辅助信息：本报表按 Marketplace + Seller SKU 独立计算；系列审批与系列停补不参与本次计算，原配置保留。',
    'sales_status': 'SOLD/NO_SALES 仅来自完整的90日窗口；缺失数据为 UNKNOWN，不代表零销量。',
    'sales_30d': '默认计算基准。只有完整30日销量可用于计算；取消订单不计入，未知数值保持空白。',
    'fulfillable': 'FBA Inventory API fulfillableQuantity，可售数量；补货抵扣仅使用此项及已核对的 eligible_inbound。',
    'reserved': 'totalReservedQuantity，包括客户订单预留、FC转运及处理；单独展示，不重复加到可售或在途。',
    'unfulfillable': 'totalUnfulfillableQuantity，不可售；不计为补货可用库存。',
    'inbound_working': 'inboundWorkingQuantity，创建或准备中的入仓库存摘要；不与货件明细重复抵扣。',
    'inbound_shipped': 'inboundShippedQuantity，已发运入仓库存摘要；仅用于货件核对。',
    'inbound_receiving': 'inboundReceivingQuantity，接收中的入仓库存摘要；仅用于货件核对。',
    'inventory_inbound': '上述三项库存摘要之和；与货件数量重叠，禁止与 eligible_inbound 再次相加抵扣。',
    'eligible_inbound': '按已配置的入仓计入规则核对后的唯一货件剩余数量；只抵扣一次。',
    'unknown_eta_inbound': '未知ETA的货件剩余数量；不将ETA未知视为数量零。根据规则可能阻断建议。',
    'recommended_qty': '明确SKU停补为0；数据缺失或异常导致无法计算时保持空白。未确认参数的正数仅为试算。',
    'parameter_status': 'CONFIRMED 表示参数已明确确认；SCENARIO_ONLY 表示参数未确认，仅供试算，不得作为正式采购依据。',
    'units': '货件发运减已接收的剩余量；并非库存摘要之外的一份额外在途库存。',
    'remaining': '货件发运减已接收的剩余量；已接收大于已发运或其他数量冲突时保持空白，不能以0抵扣。',
    'quantity_shipped': 'Amazon货件行的已发运数量；只展示已验证的实际值，不填补缺失值。',
    'quantity_received': 'Amazon货件行的已接收数量；若大于已发运则数量冲突，剩余在途未知。',
    'eta': '仅允许人工确认的ETA；未知为空，不使用 Amazon need-by/arrival 字段伪装已确认ETA。',
}


def _quantity(value):
    return (type(value) in (int, float) and math.isfinite(value)
            and value >= 0 and int(value) == value)


def _known_number(value):
    return type(value) in (int, float) and math.isfinite(value) and value >= 0


def _table_exists(con, table):
    return con is not None and con.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)).fetchone() is not None


def _dict_rows(cursor):
    names = [item[0] for item in cursor.description]
    return [dict(zip(names, tuple(row))) for row in cursor]


def _inventory_details(con, marketplace):
    if not _table_exists(con, 'inventory_details'):
        return {}
    values = _dict_rows(con.execute(
        'SELECT sku,inbound_working,inbound_shipped,inbound_receiving,observed_at '
        'FROM inventory_details WHERE marketplace=?', (marketplace,)))
    return {row['sku']: row for row in values}


def _inbound_details(con, marketplace, known_skus):
    if not _table_exists(con, 'inbound'):
        return []
    values = _dict_rows(con.execute(
        'SELECT marketplace,sku,shipment_id,units,eta,status FROM inbound '
        'WHERE marketplace=? ORDER BY sku,shipment_id', (marketplace,)))
    return [row for row in values if row['sku'] in known_skus]


def _inbound_exceptions(con, marketplace, known_skus):
    if not _table_exists(con, 'sku_inbound_exception'):
        return []
    values = _dict_rows(con.execute(
        'SELECT marketplace,sku,shipment_id,fnsku,status,quantity_shipped,quantity_received,remaining,code '
        'FROM sku_inbound_exception WHERE marketplace=? ORDER BY sku,shipment_id,fnsku', (marketplace,)))
    return [row for row in values if row['sku'] in known_skus]


def _apply_inbound_exceptions(rows, exceptions, anomalies):
    by_sku = {row['sku']: row for row in rows}
    existing = {(row['sku'], row['code']) for row in anomalies}
    for exception in exceptions:
        row = by_sku[exception['sku']]
        code = _code(exception.get('code')) or 'INBOUND_QUANTITY_CONFLICT'
        row['inbound_complete'], row['data_complete'] = False, False
        row['eligible_inbound'] = row['unknown_eta_inbound'] = None
        row['recommended_qty'] = 0 if row['status'] == 'STOPPED' else None
        if row['status'] != 'STOPPED':
            row['status'] = 'DATA_REVIEW_REQUIRED'
            row['recommendation_label'] = '数据异常：无法计算'
            row['stop_reason'] = row.get('stop_reason') or '货件数量冲突；剩余在途为UNKNOWN，须人工核实。'
        codes = set(row['anomaly_codes'].split('; ')) - {''}
        row['anomaly_codes'] = '; '.join(sorted(codes | {code}))
        if (row['sku'], code) not in existing:
            anomalies.append({'marketplace': row['marketplace'], 'sku': row['sku'], 'asin': row.get('asin'),
                              'code': code, 'reason': '货件已接收与已发运数量冲突；剩余在途未知，须核实。',
                              'blocking': True, 'source': 'Inbound quantity exception evidence'})
            existing.add((row['sku'], code))


def _code(value):
    return value if isinstance(value, str) and re.fullmatch(r'[A-Z][A-Z0-9_]{0,99}', value) else None


def _same_snapshot(first, second):
    try:
        left = datetime.fromisoformat(first.replace('Z', '+00:00'))
        right = datetime.fromisoformat(second.replace('Z', '+00:00'))
        return left.tzinfo is not None and right.tzinfo is not None and abs((left - right).total_seconds()) <= 60
    except (AttributeError, TypeError, ValueError):
        return False


def _complete_window(row, days):
    return (row.get(f'sales_complete_{days}d') is True
            and row.get(f'observed_sales_days_{days}d') == days
            and _quantity(row.get(f'sales_{days}d')))


def _sanitize_rows(rows, marketplace, details):
    result, anomalies, seen, anomaly_seen = [], [], set(), set()
    for original in rows:
        if not isinstance(original, Mapping):
            raise ValueError('SKU report requires mapping rows')
        market, sku = original.get('marketplace'), original.get('sku')
        if market != marketplace or not isinstance(sku, str) or not sku.strip():
            raise ValueError('SKU report marketplace or SKU identity is invalid')
        if (market, sku) in seen:
            raise ValueError('Duplicate Marketplace + Seller SKU in report')
        seen.add((market, sku))
        # Keep only report fields; arbitrary payload and buyer fields cannot reach
        # the workbook. Copy before normalization so the planner is unmodified.
        allowed = set(MAIN_COLUMNS + WINDOW_COLUMNS + INVENTORY_FIELDS)
        allowed.update({'stock_age_days', 'inventory_observed_at', 'parameters_confirmed', 'data_complete',
                        'sales_source', 'inventory_source', 'inbound_source'})
        row = {key: original.get(key) for key in allowed if key in original}
        extra = details.get(sku, {})
        if not _same_snapshot(row.get('stock_observed_at'), extra.get('observed_at')):
            extra = {}
        for key in INBOUND_FIELDS:
            if key not in row:
                row[key] = extra.get(key)
        row.setdefault('inventory_observed_at', extra.get('observed_at'))
        codes = set()
        source_codes = original.get('anomaly_codes') or []
        if isinstance(source_codes, (list, tuple)):
            codes.update(code for value in source_codes if (code := _code(value)))
        for item in original.get('anomalies') or []:
            if not isinstance(item, Mapping) or not _code(item.get('code')):
                continue
            codes.add(item['code'])
            if (sku, item['code']) not in anomaly_seen:
                anomalies.append({
                    'marketplace': market, 'sku': sku, 'asin': row.get('asin'),
                    'code': item['code'], 'reason': item.get('reason'),
                    'blocking': item.get('blocking') is True, 'source': item.get('source'),
                })
                anomaly_seen.add((sku, item['code']))
        for days in WINDOWS:
            complete = _complete_window(original, days)
            row[f'sales_complete_{days}d'] = complete
            row[f'sales_{days}d'] = original.get(f'sales_{days}d') if complete else None
            row[f'daily_sales_{days}d'] = row[f'sales_{days}d'] / days if complete else None
            count = original.get(f'observed_sales_days_{days}d')
            row[f'sales_completeness_{days}d'] = min(1, count / days) if _quantity(count) else None
            for field in (f'observed_sales_{days}d', f'observed_sales_days_{days}d'):
                if not _quantity(row.get(field)):
                    row[field] = None
        row['daily_sales'] = row['daily_sales_30d']
        if row['sales_complete_90d']:
            row['sales_status'] = 'SOLD' if row['sales_90d'] > 0 else 'NO_SALES'
        else:
            row['sales_status'] = 'UNKNOWN'
        for key in INVENTORY_FIELDS + ('eligible_inbound', 'unknown_eta_inbound', 'inventory_inbound'):
            if not _quantity(row.get(key)):
                row[key] = None
        row['inventory_complete'] = (original.get('inventory_complete') is True
                                     and all(_quantity(row.get(k)) for k in INVENTORY_FIELDS))
        row['inbound_complete'] = original.get('inbound_complete') is True
        row['data_complete'] = (row['sales_complete_90d'] and row['inventory_complete']
                                and row['inbound_complete'] and original.get('data_complete', True) is True)
        if not row['sales_complete_90d']:
            codes.add('SALES_90D_INCOMPLETE')
        if not row['inventory_complete']:
            codes.add('INVENTORY_INCOMPLETE')
        if not row['inbound_complete']:
            codes.add('INBOUND_COLLECTION_INCOMPLETE')
        status = str(original.get('status') or original.get('replenishment_status') or 'UNKNOWN')
        row['status'] = status
        confirmed = (original.get('parameters_confirmed') is True
                     and original.get('parameter_status') == 'CONFIRMED')
        row['parameter_status'] = 'CONFIRMED' if confirmed else 'SCENARIO_ONLY'
        row['parameters_confirmed'] = confirmed
        qty = original.get('recommended_qty')
        row['recommended_qty'] = qty if _quantity(qty) else None
        if status == 'STOPPED':
            row['recommended_qty'] = 0
        elif (not row['sales_complete_30d'] or not row['sales_complete_90d'] or not row['inventory_complete']
              or not row['inbound_complete'] or status not in {'ACTIVE', 'SCENARIO_ONLY'}):
            if _quantity(qty) and qty > 0:
                codes.add('UNVERIFIED_RECOMMENDATION_BLOCKED')
            row['recommended_qty'] = None
        row['recommendation_label'] = ('已明确SKU停补' if status == 'STOPPED' else
            '数据异常：无法计算' if row['recommended_qty'] is None else
            '参数未确认：仅供试算，不得作为正式采购依据' if not confirmed else
            '人工审核建议；不执行采购或Amazon写入')
        row['anomaly_codes'] = '; '.join(sorted(codes))
        for code in sorted(codes):
            if (sku, code) not in anomaly_seen:
                anomalies.append({'marketplace': market, 'sku': sku, 'asin': row.get('asin'),
                                  'code': code, 'reason': '请核对数据异常和采集完整率；未知数值不是零。',
                                  'blocking': True, 'source': 'Workbook validation'})
                anomaly_seen.add((sku, code))
        result.append(row)
    result.sort(key=lambda row: (
        0 if _quantity(row.get('recommended_qty')) and row['recommended_qty'] > 0 else 1,
        0 if row.get('stockout_risk') not in (None, '', 'NONE', 'NO_SALES') else 1,
        -(row.get('sales_30d') or 0), row['sku'],
    ))
    return result, anomalies


def _safe_audit(collection_audit):
    supplied = collection_audit or {}
    if isinstance(supplied, Mapping):
        entries = supplied.get('audit') or []
        summary = supplied.get('summary') or {}
    elif isinstance(supplied, (list, tuple)):
        entries, summary = supplied, {}
    else:
        entries, summary = [], {}
    clean = []
    for item in entries:
        if not isinstance(item, Mapping):
            continue
        value = {}
        for key in AUDIT_NUMBERS:
            if _quantity(item.get(key)):
                value[key] = item[key]
        for key in AUDIT_STRINGS:
            text = item.get(key)
            if isinstance(text, str) and len(text) <= 120 and '\n' not in text:
                value[key] = text
        for key in AUDIT_DATES:
            text = item.get(key)
            if isinstance(text, str):
                try:
                    datetime.fromisoformat(text.replace('Z', '+00:00'))
                except ValueError:
                    continue
                value[key] = text
        if type(item.get('complete')) is bool:
            value['complete'] = item['complete']
        if value:
            clean.append(value)
        issues = item.get('issues')
        if isinstance(issues, Mapping):
            clean.extend({'kind': 'quality_issue', 'code': code, 'count': count}
                         for code, count in issues.items() if _code(code) and _quantity(count))
    safe_summary = {key: summary[key] for key in SUMMARY_AUDIT_KEYS
                    if isinstance(summary, Mapping) and key in summary
                    and type(summary[key]) in (str, int, float, bool)}
    return clean, safe_summary


def _summary(rows, anomalies):
    count = len(rows)
    sales_complete = sum(r['sales_complete_90d'] for r in rows)
    inventory_complete = sum(r['inventory_complete'] for r in rows)
    inbound_complete = sum(r['inbound_complete'] for r in rows)
    complete = sum(r['data_complete'] for r in rows)
    positive = [r for r in rows if _quantity(r['recommended_qty']) and r['recommended_qty'] > 0]
    eta_unknown = sum(_quantity(r.get('unknown_eta_inbound')) and r['unknown_eta_inbound'] > 0 for r in rows)
    unknown = sum(not r['data_complete'] or r['recommended_qty'] is None
                  or (_quantity(r.get('unknown_eta_inbound')) and r['unknown_eta_inbound'] > 0) for r in rows)
    return {
        'sku_count': count,
        'sold_sku_count': sum(r['sales_status'] == 'SOLD' for r in rows),
        'no_sales_sku_count': sum(r['sales_status'] == 'NO_SALES' for r in rows),
        'sales_unknown_sku_count': sum(r['sales_status'] == 'UNKNOWN' for r in rows),
        'unknown_data_sku_count': unknown,
        'eta_unknown_sku_count': eta_unknown,
        'unknown_data_definition': '数据完整率交集未通过、无法计算建议或存在正数量未知ETA的SKU并集；参数未确认但有有效试算数量不单独计为数据未知。',
        'inventory_positive_sku_count': sum(any(_quantity(r.get(k)) and r[k] > 0 for k in STOCK_FIELDS) for r in rows),
        'recommended_sku_count': len(positive),
        'recommended_units': sum(r['recommended_qty'] for r in positive),
        'anomaly_sku_count': len({(a['marketplace'], a['sku']) for a in anomalies}),
        'stopped_sku_count': sum(r['status'] == 'STOPPED' for r in rows),
        'sales_complete_sku_count': sales_complete,
        'inventory_complete_sku_count': inventory_complete,
        'inbound_complete_sku_count': inbound_complete,
        'complete_sku_count': complete,
        'data_collection_completeness': complete / count if count else None,
        'sales_completeness': sales_complete / count if count else None,
        'inventory_completeness': inventory_complete / count if count else None,
        'inbound_completeness': inbound_complete / count if count else None,
        'completeness_definition': '完整90日销量、六类库存量、完整货件采集均已验证且无身份/数据质量冲突的SKU数 / 本表全部SKU；未知ETA另行异常拦截。',
        'recommendation_basis_days': 30,
        'parameter_status': 'CONFIRMED' if rows and all(r['parameters_confirmed'] for r in rows) else 'SCENARIO_ONLY',
        'confirmed_recommended_units': sum(r['recommended_qty'] for r in positive if r['parameters_confirmed']),
        'scenario_recommended_units': sum(r['recommended_qty'] for r in positive if not r['parameters_confirmed']),
        'mode': 'INDEPENDENT_SKU_READ_ONLY',
        'family_mapping_required': False,
        'family_stop_applied': False,
        'amazon_write_executed': False,
        'purchase_executed': False,
    }


def _report_sheet(wb, name, columns, rows):
    ws = sheet(wb, name, list(columns), rows)
    for cell in ws[1]:
        if cell.value in COMMENTS:
            cell.comment = Comment(COMMENTS[cell.value], 'Uloveido')
    return ws


def export_sku_report(rows, path, con, marketplace, settings, collection_audit=None):
    """Export exactly the supplied SKU cohort without importing or changing data.

    Only allowlisted planning values and numeric collection metadata are exported.
    No order identities, buyer fields, request/response bodies or credential fields
    are read. SQLite queries are batched and never mutate the database.
    """
    if not isinstance(settings, Mapping):
        raise ValueError('SKU report settings must be a mapping')
    details = _inventory_details(con, marketplace)
    values, anomalies = _sanitize_rows(rows, marketplace, details)
    cohort = {row['sku'] for row in values}
    inbound = _inbound_details(con, marketplace, cohort)
    inbound_exceptions = _inbound_exceptions(con, marketplace, cohort)
    _apply_inbound_exceptions(values, inbound_exceptions, anomalies)
    summary = _summary(values, anomalies)
    summary['inbound_quantity_conflict_sku_count'] = len({row['sku'] for row in inbound_exceptions})
    summary['inbound_exception_line_count'] = len(inbound_exceptions)
    audit_rows, audit_summary = _safe_audit(collection_audit)
    summary.update(audit_summary)
    summary['excel_path'] = str(Path(path))
    wb = Workbook()
    wb.remove(wb.active)
    _report_sheet(wb, 'SKU补货总表', MAIN_COLUMNS, values)
    _report_sheet(wb, '7_30_60_90天销量', IDENTITY_COLUMNS + ('sales_status',) + WINDOW_COLUMNS +
                  ('asof', 'window_start', 'sales_timezone', 'sales_source'), values)
    _report_sheet(wb, 'FBA库存明细', IDENTITY_COLUMNS + INVENTORY_FIELDS + (
        'inventory_inbound', 'inventory_complete', 'stock_observed_at',
        'inventory_observed_at', 'stock_age_days', 'inventory_source'), values)
    inbound_rows = []
    by_sku = {row['sku']: row for row in values}
    exception_keys = {(row['sku'], row['shipment_id']) for row in inbound_exceptions}
    for original in inbound:
        if (original['sku'], original['shipment_id']) in exception_keys:
            continue
        row = dict(original)
        if not _quantity(row.get('units')):
            row['units'] = None
        row['remaining'] = row['units']
        plan = by_sku[row['sku']]
        row.update(asin=plan.get('asin'), family=plan.get('family'),
                   inbound_rule=plan.get('inbound_rule'), inbound_complete=plan['inbound_complete'],
                   inbound_source=plan.get('inbound_source'))
        row['eta'] = row.get('eta') or None
        row['inventory_note'] = '货件余额与FBA库存入仓摘要核对；不得将两者重复相加抵扣。'
        inbound_rows.append(row)
    for original in inbound_exceptions:
        row = dict(original)
        plan = by_sku[row['sku']]
        for key in ('quantity_shipped', 'quantity_received'):
            if not _quantity(row.get(key)):
                row[key] = None
        # A quantity exception cannot be resolved by clamping the subtraction
        # to zero, even if a stale database row stored that fabricated value.
        row.update(units=None, remaining=None, eta=None,
                   asin=plan.get('asin'), family=plan.get('family'),
                   inbound_rule=plan.get('inbound_rule'), inbound_complete=False,
                   inbound_source=plan.get('inbound_source'),
                   code=_code(row.get('code')) or 'INBOUND_QUANTITY_CONFLICT',
                   inventory_note='数量冲突：已接收超过已发运；剩余在途为UNKNOWN，须人工核实，不填0。')
        inbound_rows.append(row)
    inbound_rows.sort(key=lambda row: (0 if row.get('code') else 1, row['sku'], row['shipment_id'], row.get('fnsku') or ''))
    _report_sheet(wb, '在途库存明细', IDENTITY_COLUMNS + (
        'shipment_id', 'fnsku', 'quantity_shipped', 'quantity_received', 'remaining', 'units', 'eta', 'status',
        'code', 'inbound_rule', 'inbound_complete', 'inbound_source', 'inventory_note'), inbound_rows)
    _report_sheet(wb, '入仓库存核对', IDENTITY_COLUMNS + INBOUND_FIELDS + (
        'inventory_inbound', 'eligible_inbound', 'unknown_eta_inbound',
        'inbound_complete', 'inbound_rule', 'status', 'anomaly_codes'), values)
    _report_sheet(wb, '建议补货数量', RECOMMENDATION_COLUMNS, values)
    risk_rows = [row for row in values if row.get('stockout_risk') not in ('NONE', 'NO_SALES', 'NO_RECENT_DEMAND', 'COVERED')]
    _report_sheet(wb, '缺货风险', IDENTITY_COLUMNS + (
        'stockout_risk', 'status', 'sales_status', 'daily_sales', 'fulfillable',
        'coverage_days', 'lead_days', 'recommended_qty', 'recommendation_label', 'stop_reason'), risk_rows)
    _report_sheet(wb, '数据异常', ANOMALY_COLUMNS, anomalies)
    _report_sheet(wb, '已停补SKU', RECOMMENDATION_COLUMNS, [row for row in values if row['status'] == 'STOPPED'])
    parameters = [{'parameter': 'report_mode', 'value': 'INDEPENDENT_SKU_READ_ONLY'},
                  {'parameter': 'parameter_status', 'value': summary['parameter_status']},
                  {'parameter': 'warning', 'value': '参数未确认时仅为试算，不得作为正式采购依据；所有建议仅供人工审核。'},
                  {'parameter': 'sales_basis_days', 'value': 30},
                  {'parameter': 'family_mapping_required', 'value': False},
                  {'parameter': 'family_stop_applied', 'value': False},
                  {'parameter': 'direct_sku_stop_applied', 'value': True},
                  {'parameter': 'reserved_policy', 'value': '预留库存单独显示；不计为可售、不与在途重复计算。'},
                  {'parameter': 'inbound_policy', 'value': '库存摘要只核对；仅按确认规则对货件余额抵扣一次。'},
                  {'parameter': 'unknown_policy', 'value': '未知销量、库存、ETA保留为空或UNKNOWN，不填0。'},
                  {'parameter': 'formula', 'value': 'max(0, ceil(30日销量/30 × (运输入仓天数+目标覆盖天数+安全天数) - 可售 - 唯一可抵扣在途)); 然后按MOQ及装箱数取整。'}]
    parameters.extend({'parameter': key, 'value': settings[key]}
                      for key in PARAMETER_KEYS if key in settings)
    parameter_sheet = _report_sheet(wb, '计算参数', ('parameter', 'value'), parameters)
    parameter_sheet.column_dimensions['B'].width = 100
    warning_fill = PatternFill('solid', fgColor='FFF2CC')
    for row in parameter_sheet.iter_rows(min_row=2, max_row=4):
        for cell in row:
            cell.fill = warning_fill
            cell.font = Font(bold=True, color='9C5700')
    _report_sheet(wb, '采集审计', AUDIT_COLUMNS, audit_rows)
    statistics = _report_sheet(wb, '报表统计', ('metric', 'value'),
                               [{'metric': key, 'value': value} for key, value in summary.items()])
    statistics.column_dimensions['B'].width = 100
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return summary
