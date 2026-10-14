"""First-batch Excel preparation for finite, unapproved SKU review sets.

No API, database write or authoritative approval is performed here. Only
allowlisted operational fields are exported; unknown quantities remain blank.
"""
from collections import defaultdict
from pathlib import Path
import math

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .batch_export import ANOMALY_COLUMNS, _items, _value


DECISION_SHEET = '审核决定'
MEMBERS_SHEET = 'SKU完整明细'
ANOMALIES_SHEET = '异常待核'
STOP_SHEET = '停补标记'
QUANTITY_FIELDS = ('fulfillable', 'reserved', 'unfulfillable',
                   'inbound_working', 'inbound_shipped', 'inbound_receiving')
DECISION_COLUMNS = (
    'priority', 'candidate_family', 'sku_count',
    'decision', 'approved_family', 'reviewer', 'reason',
    'stock_positive_count',
    'recent_sales_positive_count', 'sample_units_observed', 'sales_observed_sku_count',
    'inventory_complete_sku_count', *QUANTITY_FIELDS, 'anomaly_count', 'anomaly_codes',
    'mapping_basis', 'batch_eligible', 'priority_reasons',
    'marketplace', 'api_marketplace', 'snapshot_id', 'group_id', 'manifest_id', 'member_digest',
)
DETAIL_COLUMNS = (
    'snapshot_id', 'group_id', 'member_digest', 'sku', 'asin', 'marketplace',
    'api_marketplace', 'candidate_family', 'family', 'mapping_basis',
    'mapping_status', 'replenishment_status', 'listing_status', 'listing_evidence',
    'recent_sales_status', 'recent_sales_source', 'recent_units_observed',
    'fba_stock_status', *QUANTITY_FIELDS, 'inventory_observed_at',
    'group_member_count', 'batch_eligible', 'priority', 'priority_reasons',
    'anomaly_codes', 'anomaly_messages', 'first_discovered_at', 'last_seen_at',
    'history_status', 'recommended_qty',
)
STOP_COLUMNS = (
    'snapshot_id', 'group_id', 'member_digest', 'marketplace', 'api_marketplace',
    'sku', 'family', 'candidate_family', 'existing_replenishment_status',
    'stop', 'reason', 'effective_date', 'review_status', 'reviewer',
)
VALIDATION_COLUMNS = ('check', 'status', 'severity', 'group_id', 'sku', 'message')
SUMMARY_COLUMNS = ('metric', 'value', 'description')
SOURCE_COLUMNS = ('source_field', 'value')
SUMMARY_KEYS = (
    'sku_count', 'group_count', 'batch_eligible_group_count', 'exception_group_count',
    'stock_positive_sku_count', 'stock_or_inbound_positive_sku_count',
    'inbound_only_sku_count', 'recent_sales_positive_sku_count',
    'anomaly_sku_count', 'mapping_conflict_sku_count', 'format_warning_sku_count',
    'exact_match_sku_count', 'review_action_count', 'review_actions_saved',
    'review_reduction_pct', 'validation_structure', 'import_permitted',
    'require_confirmation', 'snapshot_id', 'source_snapshot_id', 'generated_at',
    'full_history_sync', 'daily_schedule_deployed',
)
SOURCE_KEYS = (
    'file_name', 'source_file_name', 'sha256', 'source_sha256', 'snapshot_id',
    'source_snapshot_id', 'source_member_digest', 'generated_at', 'asof',
)
LABELS = {
    'priority': '审核优先级', 'priority_reasons': '优先级依据',
    'snapshot_id': '冻结快照标识', 'group_id': '明确审核组标识',
    'manifest_id': '固定成员清单标识（与 group_id 一致）',
    'member_digest': '固定成员校验值', 'marketplace': '站点',
    'api_marketplace': 'Amazon Marketplace ID', 'candidate_family': '候选系列（待确认）',
    'mapping_basis': '映射来源与依据', 'sku_count': '明确 SKU 成员数量',
    'batch_eligible': '本组是否允许批量审核', 'stock_positive_count': '有库存或在途的 SKU 数',
    'recent_sales_positive_count': '有样本正销量证据的 SKU 数',
    'sample_units_observed': '已知样本观察销量之和（不是完整窗口销量）',
    'sales_observed_sku_count': '有已知样本销量数量的 SKU 数',
    'inventory_complete_sku_count': '六项库存数量均已知的 SKU 数',
    'fulfillable': '可售库存', 'reserved': '预留库存', 'unfulfillable': '不可售库存',
    'inbound_working': '入仓处理中数量', 'inbound_shipped': '入仓已发运数量',
    'inbound_receiving': '入仓接收中数量', 'anomaly_count': '异常 SKU 数',
    'anomaly_codes': '异常代码', 'anomaly_messages': '异常说明',
    'decision': '审核决定：APPROVE / REJECT / NEEDS_REVIEW，默认待审核',
    'approved_family': '人工确认的正式系列：APPROVE 时填写',
    'reviewer': '审核人：明确决定或停补时必须填写', 'reason': '审核或停补原因',
    'sku': 'Seller SKU', 'asin': 'ASIN', 'family': '现有正式系列（未确认时留空）',
    'mapping_status': '映射审核状态，独立于补货状态',
    'replenishment_status': '当前补货安全状态（不会因本工作簿而改变）',
    'existing_replenishment_status': '现有补货状态，仅供查看，不会解除既有停补',
    'listing_status': '销售状态证据：UNKNOWN 不等于下架',
    'listing_evidence': '销售状态证据来源', 'recent_sales_status': '样本销量证据状态',
    'recent_sales_source': '样本销量证据来源',
    'recent_units_observed': '小批量样本观察销量：空白不表示零销量',
    'fba_stock_status': 'FBA 库存证据状态', 'inventory_observed_at': '库存观察时间（保留原时区）',
    'group_member_count': '所属审核组的明确成员数',
    'first_discovered_at': '首次发现时间（保留原时区）',
    'last_seen_at': '最近发现时间（保留原时区）', 'history_status': '历史存在状态证据',
    'recommended_qty': '本轮仅审核准备；建议数量强制为 0',
    'stop': 'TRUE 表示申请停补；空白表示未决定；FALSE 不会解除既有停补',
    'effective_date': '停补生效日期 YYYY-MM-DD，按美国站当地日期核对',
    'review_status': '停补审核决定：APPROVE / REJECT / NEEDS_REVIEW；默认待审核',
    'severity': '异常严重性', 'code': '异常代码', 'message': '说明', 'source': '异常来源',
    'check': '模拟校验项目', 'status': '校验结果：READY / PENDING / BLOCKED / WARNING',
}
SUMMARY_DESCRIPTIONS = {
    'sku_count': '本批明确 SKU 数量，精确匹配也未批准。',
    'group_count': '本批审核决定总数，包括必须单独解决的异常组。',
    'batch_eligible_group_count': '可审核固定成员集合的系列组，不包括未来同前缀 SKU。',
    'exception_group_count': '异常单独审核组，解决问题前不允许批准。',
    'stock_positive_sku_count': '当前保存的观察中可售、预留、不可售任一库存数量大于零。',
    'stock_or_inbound_positive_sku_count': '六项库存或在途数量中至少一项大于零。',
    'inbound_only_sku_count': '有已知正入仓数量，但未观察到可售、预留、不可售正库存；未知数量不能据此确认零库存。',
    'recent_sales_positive_sku_count': '小批量样本中观察到正销量，未完成 90 天采集。',
    'anomaly_sku_count': '本批存在映射或格式异常的唯一 SKU 数。',
    'review_actions_saved': '估算减少的审核决定次数，不代表人工耗时减少。',
    'review_reduction_pct': '审核决定次数减少百分比，不代表人工耗时减少。',
    'import_permitted': '本轮不会正式导入；需要用户审核完毕并明确确认。',
    'require_confirmation': '正式导入之前必须取得用户明确确认并重新执行模拟校验。',
}


def _known_quantity(value):
    return (isinstance(value, (int, float)) and not isinstance(value, bool)
            and math.isfinite(value) and value >= 0)


def _positive_stock(row):
    return any(_known_quantity(row.get(key)) and row[key] > 0 for key in QUANTITY_FIELDS)


def _positive_fba_stock(row):
    return any(_known_quantity(row.get(key)) and row[key] > 0 for key in QUANTITY_FIELDS[:3])


def _positive_sales(row):
    value = row.get('recent_units_observed')
    return _known_quantity(value) and value > 0


def _strict_sum(rows, key):
    values = [row.get(key) for row in rows]
    return sum(values) if values and all(_known_quantity(v) for v in values) else None


def _safe_sheet(wb, name, columns, rows, *, editable=()):
    ws = wb.create_sheet(name)
    ws.append(list(columns))
    for row_index, row in enumerate(rows, 2):
        ws.append([_value(row.get(key)) for key in columns])
        for col_index, key in enumerate(columns, 1):
            cell = ws.cell(row_index, col_index)
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if key in editable:
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
                cell.protection = Protection(locked=False)
            cell.alignment = Alignment(vertical='top', wrap_text=key in (
                'priority_reasons', 'anomaly_messages', 'message', '操作说明', 'description'))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for index, key in enumerate(columns, 1):
        header = ws.cell(1, index)
        header.font = Font(bold=True, color='FFFFFF')
        header.fill = PatternFill('solid', fgColor='1F4E78')
        header.alignment = Alignment(wrap_text=True, vertical='center')
        if key in LABELS:
            header.comment = Comment(LABELS[key], 'Uloveido')
        width = 45 if key in ('priority_reasons', 'mapping_basis', 'anomaly_messages', 'message') else 22
        if key in ('snapshot_id', 'group_id', 'manifest_id', 'member_digest'):
            width = 36
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 36
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.zoomScale = 85
    return ws


def _dropdown(ws, columns, key, options, *, rows=None, allow_blank=False):
    validation = DataValidation(type='list', formula1='"' + ','.join(options) + '"',
                                allow_blank=allow_blank)
    validation.showErrorMessage = True
    validation.errorTitle = '请使用允许的审核值'
    validation.error = '请使用下拉列表；不得粘贴其他审核状态。'
    validation.errorStyle = 'stop'
    ws.add_data_validation(validation)
    col = columns.index(key) + 1
    if rows is None:
        validation.add(f'{get_column_letter(col)}2:{get_column_letter(col)}{max(2, ws.max_row)}')
    else:
        for row_index in rows:
            validation.add(ws.cell(row_index, col))
    return validation


def _protect(ws):
    # Protection avoids accidental edits; import validation is authoritative.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def _prepare(payload):
    groups = list(payload.get('groups') or [])
    if not groups:
        raise ValueError('First-batch review requires a fixed set of groups')
    snapshot_id = payload.get('snapshot_id') or (payload.get('summary') or {}).get('snapshot_id')
    normalized = []
    for group in groups:
        row = dict(group)
        row['snapshot_id'] = row.get('snapshot_id') or snapshot_id
        for key in ('snapshot_id', 'group_id', 'member_digest'):
            if not isinstance(row.get(key), str) or not row[key]:
                raise ValueError('First-batch review is missing a frozen scope identifier')
        if row['snapshot_id'] != snapshot_id:
            raise ValueError('Review groups do not share the source snapshot')
        normalized.append(row)
    indexed = {str(group['group_id']): group for group in normalized}
    if len(indexed) != len(normalized):
        raise ValueError('Duplicate review group identity')
    anomalies = list(payload.get('anomalies') or [])
    anomaly_by_identity = defaultdict(list)
    for anomaly in anomalies:
        anomaly_by_identity[(str(anomaly.get('group_id')), anomaly.get('sku'))].append(anomaly)
    # Reuse the ordinary export membership checks, without changing the source.
    items = _items(dict(payload, groups=normalized), normalized)
    originals = {}
    for original in payload.get('items') or []:
        group = indexed[str(original['group_id'])]
        if original.get('snapshot_id') not in (None, group['snapshot_id']):
            raise ValueError('Review member snapshot differs from its fixed group')
        if original.get('member_digest') not in (None, group['member_digest']):
            raise ValueError('Review member checksum differs from its fixed group')
        if original.get('exact_reference_match') is not True and original.get('mapping_basis') != 'USER_WORKBOOK_EXACT_SKU':
            raise ValueError('First-batch review only accepts exact source-workbook matches')
        originals[(str(original['group_id']), original['sku'])] = original
    rows = []
    for item in items:
        identity = (str(item['group_id']), item['sku'])
        original = originals[identity]
        group = indexed[identity[0]]
        row = {key: item.get(key) for key in DETAIL_COLUMNS}
        row.update(recent_sales_source=original.get('recent_sales_source') or 'NO_COMPLETE_SALES_EVIDENCE',
                   group_member_count=group['sku_count'], batch_eligible=group.get('batch_eligible') is True,
                   recommended_qty=0)
        for key in QUANTITY_FIELDS + ('recent_units_observed',):
            row[key] = original.get(key) if _known_quantity(original.get(key)) else None
        relevant = anomaly_by_identity.get(identity) or []
        row['anomaly_codes'] = '; '.join(sorted({a['code'] for a in relevant if isinstance(a.get('code'), str)})) or None
        row['anomaly_messages'] = '; '.join(a['message'] for a in relevant if isinstance(a.get('message'), str)) or None
        rows.append(row)
    by_group = defaultdict(list)
    for row in rows:
        by_group[str(row['group_id'])].append(row)
    # Both sales and stock/inbound evidence appear before unknown/zero evidence.
    def risk_key(group):
        children = by_group[str(group['group_id'])]
        return (-int(any(_positive_sales(r) for r in children)),
                -int(any(_positive_stock(r) for r in children)),
                str(group.get('priority') or 'UNKNOWN'),
                -float(group.get('priority_score') or 0), str(group['group_id']))
    normalized.sort(key=risk_key)
    order = {str(group['group_id']): i for i, group in enumerate(normalized)}
    rows.sort(key=lambda row: (order[str(row['group_id'])],
                              -int(_positive_sales(row)), -int(_positive_stock(row)), str(row['sku'])))
    group_rows = []
    for group in normalized:
        children = by_group[str(group['group_id'])]
        row = {key: group.get(key) for key in DECISION_COLUMNS}
        row.update(manifest_id=group['group_id'], decision='NEEDS_REVIEW',
                   approved_family=None, reviewer=None, reason=None,
                   stock_positive_count=sum(_positive_stock(r) for r in children),
                   recent_sales_positive_count=sum(_positive_sales(r) for r in children),
                   sales_observed_sku_count=sum(_known_quantity(r.get('recent_units_observed')) for r in children),
                   inventory_complete_sku_count=sum(all(_known_quantity(r.get(k)) for k in QUANTITY_FIELDS) for r in children),
                   anomaly_count=sum(bool(r['anomaly_codes']) for r in children),
                   anomaly_codes='; '.join(sorted({code for r in children for code in (r.get('anomaly_codes') or '').split('; ') if code})) or None)
        observed = [r['recent_units_observed'] for r in children if _known_quantity(r.get('recent_units_observed'))]
        row['sample_units_observed'] = sum(observed) if observed else None
        for key in QUANTITY_FIELDS:
            row[key] = _strict_sum(children, key)
        group_rows.append(row)
    identities = set(originals)
    if any((str(a.get('group_id')), a.get('sku')) not in identities for a in anomalies):
        raise ValueError('Anomaly refers to a SKU outside the fixed review set')
    return group_rows, rows, anomalies


def _instructions(wb):
    lines = [
        ('审核入口', '先看“审核决定”的优先级、库存/在途及样本销量，再按 group_id 筛选“SKU完整明细”查看每组完整成员。英文表头含中文批注，黄色单元格可填写。'),
        ('精确匹配仍未批准', '本批来自上传表精确匹配，仍只是候选；全部 decision 默认 NEEDS_REVIEW。来源精确匹配不是对真实系列或停补规则的自动批准。'),
        ('完整查看成员', '使用 group_id 过滤 SKU完整明细，核对每条 SKU、ASIN、候选系列、映射依据和异常。sku_count 与 group_member_count 仅包含当前冻结的明确 SKU 集合。'),
        ('批量批准系列', '确认本组每个成员都属于同一系列后，在“审核决定”将 decision 改为 APPROVE，填写 approved_family、reviewer、reason。approved_family 必须人工填写；不得将候选直接视为正式系列。'),
        ('拒绝或继续核对', '映射不正确时选 REJECT 并填写 reviewer、reason；证据不足时保留 NEEDS_REVIEW。不能批准未来所有同前缀 SKU。'),
        ('异常单独处理', 'batch_eligible=FALSE 的异常组可选 REJECT 或 NEEDS_REVIEW，并填写 reviewer、reason 记录意见。异常可以拒绝，批准前需先核实并重新生成单体审核集合；当前不可直接批量批准，approved_family 保持锁定空白。请在“异常待核”核对格式或系列冲突，当前文件不能强行批准这些组。'),
        ('禁止修改固定集合', '不得改动 snapshot_id、group_id、manifest_id、member_digest、SKU/ASIN、候选系列或成员数量，也不得增删成员行。需要变更集合时重新生成文件。Excel 锁定只是防误操作，模拟校验仍会检查实际集合。'),
        ('样本销量口径', 'recent_units_observed / sample_units_observed 仅是已保存小批量样本的观察数量，不是完整 7/30/60/90 天销量。sales_observed_sku_count 显示已知数量覆盖；没有证据的成员保持空白，不当作零。'),
        ('库存口径', '六项库存数量分别为可售、预留、不可售、入仓处理中、已发运、接收中。分组总数只有该字段所有成员都已知时才显示；空白不等于零。库存观察时间保留原始时区，库存存在也不证明 Listing 在售。'),
        ('停补标记', '在“停补标记”定位明确 SKU；若需停补，填写 stop=TRUE、reason、effective_date、reviewer，并将 review_status 改为 APPROVE。生效日期按美国站当地日期核对；候选系列只供参考。'),
        ('停补联动', '同站点同一已批准系列任意成员有效停补，会联动该系列全部已确认子体，包括本批之外的已确认成员。模拟校验会报告影响范围；未知系列先拦截，不能猜测联动。'),
        ('已有停补保护', '空白、FALSE、REJECT 或销量增长、缺货、新增子体都不会恢复已有停补。恢复须另走明确人工确认并保留审计。所有 STOPPED SKU 建议量保持 0。'),
        ('模拟校验', '本次“模拟校验”检查来源快照、成员集合、映射冲突及停补风险。编辑后须再次运行离线模拟，确认成员未变、正式映射无冲突，并核对每项停补的系列联动范围。'),
        ('正式导入前确认', '填好审核表后提交给助手复核。只有你明确确认该审核结果并允许导入后，才能写入正式映射或停补配置。NEEDS_REVIEW 行保持不变；本次生成文件不导入任何决定。'),
        ('安全边界', '本轮不请求 Amazon API，不采集完整 90 天销量，不部署调度，不创建货件或修改 Listing、库存、价格。全部补货建议为 0，只供人工审核。'),
    ]
    ws = _safe_sheet(wb, '中文说明', ('项目', '操作说明'),
                     [{'项目': title, '操作说明': text} for title, text in lines])
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 115
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row[0].row].height = 58


def export_first_batch(payload, output_dir):
    """Create a literal-only first-batch workbook, with no preapproved decisions."""
    groups, items, anomalies = _prepare(payload)
    wb = Workbook()
    wb.remove(wb.active)
    editable = ('decision', 'approved_family', 'reviewer', 'reason')
    ws = _safe_sheet(wb, DECISION_SHEET, DECISION_COLUMNS, groups, editable=editable)
    ws.freeze_panes = 'D2'
    eligible_rows = []
    exception_rows = []
    for row_index, group in enumerate(groups, 2):
        if group.get('batch_eligible') is True:
            eligible_rows.append(row_index)
        else:
            exception_rows.append(row_index)
            for key in editable:
                cell = ws.cell(row_index, DECISION_COLUMNS.index(key) + 1)
                cell.protection = Protection(locked=key == 'approved_family')
                if key == 'approved_family':
                    cell.fill = PatternFill('solid', fgColor='E7E6E6')
                cell.comment = Comment('异常组可以 REJECT 或 NEEDS_REVIEW 并记录审核人和意见；不能 APPROVE。批准前须先核实并重新生成单体审核集合，正式系列保持空白。', 'Uloveido')
    _dropdown(ws, DECISION_COLUMNS, 'decision', ('APPROVE', 'REJECT', 'NEEDS_REVIEW'), rows=eligible_rows)
    _dropdown(ws, DECISION_COLUMNS, 'decision', ('REJECT', 'NEEDS_REVIEW'), rows=exception_rows)
    _protect(ws)
    ws = _safe_sheet(wb, MEMBERS_SHEET, DETAIL_COLUMNS, items)
    _protect(ws)
    ws = _safe_sheet(wb, ANOMALIES_SHEET, ANOMALY_COLUMNS, anomalies)
    _protect(ws)
    stops = []
    for item in items:
        row = {key: item.get(key) for key in STOP_COLUMNS}
        row.update(existing_replenishment_status=item.get('replenishment_status'), stop=None,
                   reason=None, effective_date=None, review_status='NEEDS_REVIEW', reviewer=None)
        stops.append(row)
    ws = _safe_sheet(wb, STOP_SHEET, STOP_COLUMNS, stops,
                     editable=('stop', 'reason', 'effective_date', 'review_status', 'reviewer'))
    _dropdown(ws, STOP_COLUMNS, 'stop', ('TRUE', 'FALSE'), allow_blank=True)
    _dropdown(ws, STOP_COLUMNS, 'review_status', ('APPROVE', 'REJECT', 'NEEDS_REVIEW'))
    _protect(ws)
    ws = _safe_sheet(wb, '模拟校验', VALIDATION_COLUMNS, payload.get('validation') or [])
    _protect(ws)
    summary = payload.get('summary') or {}
    computed = dict(summary, sku_count=len(items), group_count=len(groups),
                    batch_eligible_group_count=sum(g.get('batch_eligible') is True for g in groups),
                    exception_group_count=sum(g.get('batch_eligible') is not True for g in groups),
                    stock_positive_sku_count=sum(_positive_fba_stock(item) for item in items),
                    stock_or_inbound_positive_sku_count=sum(_positive_stock(item) for item in items),
                    inbound_only_sku_count=sum(_positive_stock(item) and not _positive_fba_stock(item) for item in items),
                    recent_sales_positive_sku_count=sum(_positive_sales(item) for item in items),
                    anomaly_sku_count=len({(a.get('api_marketplace') or a.get('marketplace'), a.get('sku')) for a in anomalies}),
                    import_permitted=False, require_confirmation=True,
                    full_history_sync=False, daily_schedule_deployed=False)
    _safe_sheet(wb, '审核统计', SUMMARY_COLUMNS,
                [{'metric': key, 'value': computed[key], 'description': SUMMARY_DESCRIPTIONS.get(key)}
                 for key in SUMMARY_KEYS if key in computed])
    source = payload.get('source') or {}
    _safe_sheet(wb, '来源快照', SOURCE_COLUMNS,
                [{'source_field': key, 'value': source[key]} for key in SOURCE_KEYS if key in source])
    _instructions(wb)
    dest = Path(output_dir)
    dest.mkdir(parents=True, exist_ok=True)
    path = dest / 'first_batch_sku_mapping_review.xlsx'
    wb.save(path)
    return {'review_path': str(path)}
