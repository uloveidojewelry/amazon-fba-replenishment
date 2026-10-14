"""Offline review preparation and explicit human workbook decisions."""
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from .batch_decisions import (_atomic, approve_batch, reject_batch, review_manifest_members,
                              review_audit_rows, _review_audit)
from .batch_export import export_batch_review
from .mock import is_mock_database
from .review_queue import build_review_queue


def prepare_batch_review(con, output_dir, marketplace='US', sample_path=None, asof=None):
    if is_mock_database(con):
        raise ValueError('Real batch review cannot use a mock database')
    sample = None
    if sample_path:
        with Path(sample_path).open(encoding='utf-8') as stream:
            sample=json.load(stream)
        if not isinstance(sample,dict) or not isinstance(sample.get('orders'),list) or not isinstance(sample.get('items'),dict):
            raise ValueError('Sales priority evidence requires a sanitized sample structure')
    queue=build_review_queue(con,marketplace,sample=sample,asof=asof)
    # Full before/after decisions remain in SQLite. Workbook includes selected
    # non-secret operational fields only, never arbitrary serialized responses.
    for event in review_audit_rows(con,marketplace):
        before,after=event['original_values'],event['changed_values']
        queue['audit'].append({'at':event['at'],'action':event['action'],'marketplace':event['marketplace'],
                               'sku':event['sku'],'group_id':event['operation_id'],'reviewer':event['reviewer'],
                               'original_family':before.get('family'),'approved_family':after.get('family'),
                               'original_mapping_status':before.get('mapping_status'),'mapping_status':after.get('mapping_status'),
                               'original_replenishment_status':before.get('replenishment_status'),
                               'replenishment_status':after.get('replenishment_status'),'reason':after.get('reason')})
    paths=export_batch_review(queue,output_dir)
    dest=Path(output_dir)
    with (dest/'batch-review-summary.json').open('w',encoding='utf-8') as stream:
        json.dump(queue['summary'],stream,ensure_ascii=False,indent=2)
    return dict(paths,summary=queue['summary'])


def _sheet_rows(workbook,name):
    if name not in workbook.sheetnames:
        raise ValueError('Review workbook is missing the fixed decision/member sheet')
    iterator=workbook[name].iter_rows(values_only=True)
    headers=next(iterator,())
    if not headers or len(set(headers))!=len(headers):
        raise ValueError('Review workbook has missing or duplicate headers')
    result=[]
    for row in iterator:
        if not any(value is not None for value in row):
            continue
        result.append(dict(zip(headers,row)))
    return result


def import_batch_review(con,path,reviewer,*,confirm=False,marketplace='US',
                        allow_remap=False,resolve_conflicts=False):
    """Validate every selected row and its visible fixed membership atomically."""
    if confirm is not True or not isinstance(reviewer,str) or not reviewer.strip():
        raise ValueError('Explicit reviewed-workbook confirmation and reviewer are required')
    workbook=load_workbook(path,read_only=True,data_only=False)
    try:
        decisions=_sheet_rows(workbook,'批量审核决定')
        members=_sheet_rows(workbook,'SKU明确集合')
    finally:
        workbook.close()
    selected=[]
    for row in decisions:
        action=row.get('decision')
        if action is None or action in ('PENDING','UNAPPROVED'):
            continue
        if action not in ('APPROVED','REJECTED'):
            raise ValueError('Review decision must be explicit APPROVED, REJECTED or PENDING')
        if row.get('reviewer')!=reviewer or row.get('marketplace')!=marketplace:
            raise ValueError('Review identity or marketplace differs from the importing reviewer/scope')
        if not isinstance(row.get('reason'),str) or not row['reason'].strip() or row['reason'].startswith('='):
            raise ValueError('Explicit human review reason is required')
        selected.append(row)
    if not selected:
        return {'approved_groups':0,'rejected_groups':0,'approved_skus':0,'rejected_skus':0,
                'skipped_groups':len(decisions),'status':'NO_APPROVED_DECISIONS'}
    required=('snapshot_id','group_id','member_digest','api_marketplace')
    ids=set()
    planned=[]
    for row in selected:
        if any(not isinstance(row.get(k),str) or not row[k] or row[k].startswith('=') for k in required):
            raise ValueError('Review decision is missing a literal frozen scope identifier')
        if row['group_id'] in ids:
            raise ValueError('Duplicate review scope cannot be decided twice in a workbook')
        ids.add(row['group_id'])
        if row.get('manifest_id') not in (None,row['group_id']):
            raise ValueError('Workbook scope identifiers disagree')
        saved=con.execute('SELECT * FROM review_queue_group WHERE snapshot_id=? AND group_id=?',
                          (row['snapshot_id'],row['group_id'])).fetchone()
        snapshot=con.execute('SELECT * FROM review_queue_snapshot WHERE snapshot_id=?',(row['snapshot_id'],)).fetchone()
        if not saved or not snapshot or snapshot['marketplace']!=marketplace or snapshot['api_marketplace']!=row['api_marketplace']:
            raise ValueError('Unknown or cross-marketplace review snapshot')
        if saved['member_digest']!=row['member_digest'] or saved['family']!=row.get('candidate_family'):
            raise ValueError('Frozen review checksum or candidate family was changed')
        manifest_members=review_manifest_members(con,row['group_id'])
        visible=[r for r in members if r.get('group_id')==row['group_id']]
        if type(row.get('sku_count')) is not int or row['sku_count']!=len(manifest_members):
            raise ValueError('Decision SKU count differs from its frozen set')
        expected={r['sku']:r['expected_state']['asin'] for r in manifest_members}
        if (len(visible)!=len(expected) or len({r.get('sku') for r in visible})!=len(visible)
                or {r.get('sku'):r.get('asin') for r in visible}!=expected):
            raise ValueError('Visible workbook membership changed; no hidden or future SKU may be approved')
        if any(r.get('snapshot_id')!=row['snapshot_id'] or r.get('member_digest')!=row['member_digest']
               or r.get('marketplace')!=marketplace or r.get('api_marketplace')!=row['api_marketplace'] for r in visible):
            raise ValueError('Visible member identity/checksum differs from the decision scope')
        if row['decision']=='APPROVED' and (not saved['batch_eligible'] or not isinstance(row.get('approved_family'),str)
                                         or row['approved_family']!=saved['family']):
            raise ValueError('Unknown/conflicting family or changed target requires a new explicit review scope')
        planned.append(row)
    summary={'approved_groups':0,'rejected_groups':0,'approved_skus':0,'rejected_skus':0,
             'skipped_groups':len(decisions)-len(selected),'status':'SUCCESS'}
    with _atomic(con):
        previous={r['sku']:r['replenishment_status'] for r in con.execute('SELECT sku,replenishment_status FROM discovery_sku WHERE marketplace=?',(marketplace,))}
        for row in planned:
            if row['decision']=='APPROVED':
                result=approve_batch(con,row['group_id'],row['member_digest'],row['approved_family'],reviewer,
                                     confirm=True,allow_remap=allow_remap,resolve_conflicts=resolve_conflicts,
                                     reason=row['reason'],refresh_safety=False)
                summary['approved_groups']+=1;summary['approved_skus']+=result['approved_count']
            else:
                result=reject_batch(con,row['group_id'],row['member_digest'],reviewer,row['reason'],confirm=True,refresh_safety=False)
                summary['rejected_groups']+=1;summary['rejected_skus']+=row['sku_count']
        from .discovery import _refresh_safety
        from .core import now
        _refresh_safety(con,marketplace)
        operation_id='workbook-'+hashlib.sha256(Path(path).read_bytes()).hexdigest()[:24]
        timestamp=now()
        for current in con.execute('SELECT sku,family,replenishment_status FROM discovery_sku WHERE marketplace=?',(marketplace,)):
            if previous.get(current['sku'])!=current['replenishment_status']:
                _review_audit(con,timestamp,reviewer,'BATCH_REPLENISHMENT_STATUS_UPDATED',operation_id,marketplace,current['sku'],
                              {'replenishment_status':previous.get(current['sku'])},
                              {'family':current['family'],'replenishment_status':current['replenishment_status']})
    return summary


def import_confirmed_stops(con,path,reviewer,marketplace='US'):
    from .batch_decisions import _read_stop_rows, normalize_stop_row, import_stop_list
    rows=[normalize_stop_row(row) for row in _read_stop_rows(path)]
    if any(row['review_status']=='APPROVED' and row['stop']=='TRUE' and row['marketplace']!=marketplace for row in rows):
        raise ValueError('Approved stop file contains a different marketplace than the selected import scope')
    return import_stop_list(con,path,reviewer)
