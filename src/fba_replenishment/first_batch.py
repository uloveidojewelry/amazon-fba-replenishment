"""First exact-reference review preparation. This module has no import operation."""
from collections import defaultdict
from contextlib import contextmanager
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import sqlite3
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from .batch_export import GROUP_COLUMNS, MEMBER_COLUMNS, ANOMALY_COLUMNS
from .first_batch_export import export_first_batch
from .first_batch_validation import validate_first_batch

GROUP_SHEET='系列优先审核'
MEMBER_SHEET='SKU明确集合'
EXACT_SHEET='精确匹配优先'
DECISION_SHEET='批量审核决定'
FIRST_DECISION_SHEET='审核决定'
FIRST_MEMBER_SHEET='SKU完整明细'
FIRST_STOP_SHEET='停补标记'
FORMAT='FIRST_BATCH_EXACT_REVIEW_V1'


@contextmanager
def readonly_database(path):
    # mode=ro prevents schema creation, migrations, approvals, audit writes and
    # "temporary approval then rollback" simulations from touching the store.
    uri=Path(path).resolve().as_uri()+'?mode=ro'
    con=sqlite3.connect(uri,uri=True)
    con.row_factory=sqlite3.Row
    con.execute('PRAGMA query_only=ON')
    try:yield con
    finally:con.close()


def _rows(workbook,name,columns=None):
    if name not in workbook.sheetnames:
        raise ValueError('Required source review worksheet is missing')
    iterator=workbook[name].iter_rows(values_only=True)
    headers=next(iterator,())
    if not headers or len(set(headers))!=len(headers):
        raise ValueError('Review worksheet has duplicate or missing headers')
    allowed=set(columns) if columns is not None else set(headers)
    return [{key:value for key,value in zip(headers,row) if key in allowed}
            for row in iterator if any(value is not None for value in row)]


def _key(row):
    return row.get('api_marketplace'),row.get('sku')


def _known_sum(rows,key):
    values=[r.get(key) for r in rows]
    return sum(values) if values and all(type(n) is int and n>=0 for n in values) else None


def _positive(row,keys):
    return any(type(row.get(key)) is int and row[key]>0 for key in keys)


def extract_first_batch(source_path,*,marketplace='US',expected_skus=852,expected_groups=336):
    path=Path(source_path)
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        exact=_rows(wb,EXACT_SHEET,MEMBER_COLUMNS)
        all_members=_rows(wb,MEMBER_SHEET,MEMBER_COLUMNS)
        group_rows=_rows(wb,GROUP_SHEET,GROUP_COLUMNS)
        source_decisions=_rows(wb,DECISION_SHEET)
        source_anomalies=_rows(wb,'数据异常',ANOMALY_COLUMNS)
    finally:wb.close()
    if len(exact)!=expected_skus or len({_key(r) for r in exact})!=len(exact):
        raise ValueError('First cohort exact-reference count or unique identity differs from the approved preparation scope')
    if any(r.get('mapping_basis')!='USER_WORKBOOK_EXACT_SKU' or r.get('marketplace')!=marketplace for r in exact):
        raise ValueError('The first cohort may contain only the selected marketplace exact-reference records')
    identities={_key(r):r for r in all_members}
    if len(identities)!=len(all_members) or any(identities.get(_key(r))!=r for r in exact):
        raise ValueError('Exact queue differs from the full visible source member records')
    by_group=defaultdict(list)
    for item in exact:by_group[item['group_id']].append(item)
    if len(by_group)!=expected_groups:
        raise ValueError('First cohort group count differs from the frozen review scope')
    group_map={r['group_id']:r for r in group_rows}
    decision_map={r['group_id']:r for r in source_decisions}
    if len(group_map)!=len(group_rows) or len(decision_map)!=len(source_decisions):
        raise ValueError('Duplicate source review group identifier')
    all_by_group=defaultdict(list)
    for r in all_members:all_by_group[r['group_id']].append(r)
    groups=[]
    selected_keys=set(identities)&{_key(r) for r in exact}
    anomalies=[r for r in source_anomalies if _key(r) in selected_keys]
    per_sku_anomalies=defaultdict(list)
    for issue in anomalies:per_sku_anomalies[_key(issue)].append(issue)
    for group_id,items in by_group.items():
        group=group_map.get(group_id);decision=decision_map.get(group_id)
        if not group or not decision or group.get('sku_count')!=len(items) or len(all_by_group[group_id])!=len(items):
            raise ValueError('Source group has mixed or hidden members outside the exact first cohort')
        if any(group.get(k)!=decision.get(k) for k in ('snapshot_id','member_digest','marketplace','api_marketplace','candidate_family','sku_count','batch_eligible')):
            raise ValueError('Source decision scope differs from its group definition')
        if any(item.get(k)!=group.get(k) for item in items for k in ('snapshot_id','member_digest','marketplace','api_marketplace','candidate_family')):
            raise ValueError('Member identity or checksum differs from the source group')
        enriched=dict(group,manifest_id=group_id,decision='NEEDS_REVIEW',approved_family=None,
                      reviewer=None,reason=None,review_status='NEEDS_REVIEW')
        for key in ('fulfillable','reserved','unfulfillable','inbound_working','inbound_shipped','inbound_receiving'):
            enriched[key]=_known_sum(items,key)
        positive_sales=[r for r in items if str(r.get('recent_sales_status','')).startswith('POSITIVE')]
        enriched.update(stock_sku_count=sum(_positive(r,('fulfillable','reserved','unfulfillable')) for r in items),
                        stock_or_inbound_sku_count=sum(_positive(r,('fulfillable','reserved','unfulfillable','inbound_working','inbound_shipped','inbound_receiving')) for r in items),
                        recent_sales_sku_count=len(positive_sales),
                        sample_units_observed=_known_sum(positive_sales,'recent_units_observed'),
                        sample_evidence='POSITIVE_SAMPLE_ONLY; incomplete sales' if positive_sales else 'UNKNOWN',
                        anomaly_count=sum(len(per_sku_anomalies[_key(r)]) for r in items),
                        anomaly_message='; '.join(sorted({issue['code'] for r in items for issue in per_sku_anomalies[_key(r)]})))
        groups.append(enriched)
        for item in items:
            # Do not copy prior approval fields into a newly prepared cohort.
            item.update(decision='NEEDS_REVIEW',review_status='NEEDS_REVIEW',recommended_qty=0,
                        anomaly_codes='; '.join(issue['code'] for issue in per_sku_anomalies[_key(item)]))
            item['recent_sales_source']='原批量审核表的小批量订单样本，非完整销量窗口' if str(item.get('recent_sales_status','')).startswith('POSITIVE') else 'UNKNOWN：没有完整窗口销量证据'
    # Positive sales first, then stock/transit evidence; exceptions remain visible
    # on their own sheet rather than hiding all live-risk groups below warnings.
    groups.sort(key=lambda g:(-bool(g['recent_sales_sku_count']),-bool(g['stock_or_inbound_sku_count']),
                             g['priority'],-g['priority_score'],g['candidate_family'] or '',g['group_id']))
    ordering={g['group_id']:i for i,g in enumerate(groups)}
    exact.sort(key=lambda r:(ordering[r['group_id']],-str(r.get('recent_sales_status','')).startswith('POSITIVE'),
                             -_positive(r,('fulfillable','reserved','unfulfillable','inbound_working','inbound_shipped','inbound_receiving')),r['sku']))
    summary={'format':FORMAT,'sku_count':len(exact),'group_count':len(groups),
             'batch_reviewable_group_count':sum(g['batch_eligible'] is True for g in groups),
             'exception_group_count':sum(g['batch_eligible'] is not True for g in groups),
             'stock_positive_sku_count':sum(_positive(r,('fulfillable','reserved','unfulfillable')) for r in exact),
             'stock_or_inbound_positive_sku_count':sum(_positive(r,('fulfillable','reserved','unfulfillable','inbound_working','inbound_shipped','inbound_receiving')) for r in exact),
             'inbound_only_sku_count':sum(not _positive(r,('fulfillable','reserved','unfulfillable')) and _positive(r,('inbound_working','inbound_shipped','inbound_receiving')) for r in exact),
             'sample_sales_positive_sku_count':sum(str(r.get('recent_sales_status','')).startswith('POSITIVE') for r in exact),
             'sales_unknown_sku_count':sum(r.get('recent_sales_status')=='UNKNOWN' for r in exact),
             'mapping_anomaly_sku_count':len({_key(r) for r in anomalies}),
             'mapping_conflict_sku_count':sum(r.get('code')=='FAMILY_CONFLICT' for r in anomalies),
             'approval_count':0,'default_decision':'NEEDS_REVIEW','import_permitted':False,
             'requires_explicit_user_confirmation':True,'generated_at':datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),
             'source_workbook':path.name,'source_sha256':sha256(path.read_bytes()).hexdigest()}
    snapshots={r['snapshot_id'] for r in groups}
    if len(snapshots)!=1:
        raise ValueError('First cohort contains different frozen source snapshots')
    summary.update(snapshot_id=next(iter(snapshots)),source_snapshot_id=next(iter(snapshots)),
                   batch_eligible_group_count=summary['batch_reviewable_group_count'],
                   recent_sales_positive_sku_count=summary['sample_sales_positive_sku_count'],
                   anomaly_sku_count=summary['mapping_anomaly_sku_count'],
                   format_warning_sku_count=sum(r.get('code')=='REFERENCE_FORMAT_WARNING' for r in anomalies),
                   exact_match_sku_count=len(exact),review_action_count=len(groups),
                   require_confirmation=True,full_history_sync=False,daily_schedule_deployed=False)
    return {'summary':summary,'groups':groups,'items':exact,'anomalies':anomalies,
            'source':{'file_name':path.name,'sha256':summary['source_sha256'],
                      'source_snapshot_id':summary['source_snapshot_id'],'generated_at':summary['generated_at']}}


def _validation_rows(report):
    messages={'STALE_MEMBER_STATE':'当前数据库成员原始状态已变化，需要重新生成审核集合。',
              'FAMILY_CONFLICT':'存在系列映射冲突，不能批量批准。',
              'FORMAT_WARNING':'格式异常需先单独核实，不能提前批准。'}
    rows=[]
    for group in report.get('groups',[]):
        group_id=group['group_id']
        rows.append({'check':'固定成员与原始状态','status':group.get('structure_status','UNKNOWN'),
                     'severity':'INFO' if group.get('structure_status')=='READY' else 'BLOCKING','group_id':group_id,
                     'message':'只读核对数据库冻结集合、SKU、ASIN及校验值。'})
        for code in group.get('blockers',[]):
            rows.append({'check':code,'status':'BLOCKED','severity':'BLOCKING','group_id':group_id,
                         'message':messages.get(code,'该组需人工核实并再次模拟校验。')})
        for code in group.get('warnings',[]):
            rows.append({'check':code,'status':'WARNING','severity':'WARNING','group_id':group_id,
                         'message':messages.get(code,'潜在停补或人工审核风险；不得自动批准。')})
        risk=group.get('stop_risk') or {}
        if any(type(value) is int and value>0 for value in risk.values()):
            rows.append({'check':'STOP_LINKAGE_SCOPE','status':'WARNING','severity':'WARNING','group_id':group_id,
                         'message':'停补可能联动其他明确子体；具体范围见模拟报告。'})
    rows.append({'check':'正式导入授权','status':'PENDING','severity':'BLOCKING',
                 'message':'全部决定默认 NEEDS_REVIEW；审核完成并经用户明确确认后才允许正式导入。'})
    return rows


def prepare_first_batch(source_path,database_path,output_dir,marketplace='US'):
    payload=extract_first_batch(source_path,marketplace=marketplace)
    with readonly_database(database_path) as con:
        validation=validate_first_batch(con,payload['groups'],payload['items'])
    payload['validation']=_validation_rows(validation)
    payload['summary'].update(validation_mode='READ_ONLY_DRY_RUN',
                              validation_structure_status=validation.get('structure_status'),
                              validation_decision_status=validation.get('decision_status'),
                              validation_structure=validation.get('structure_status'))
    paths=export_first_batch(payload,output_dir)
    dest=Path(output_dir);dest.mkdir(parents=True,exist_ok=True)
    with (dest/'first_batch_simulation.json').open('w',encoding='utf-8') as stream:
        json.dump(validation,stream,ensure_ascii=False,indent=2)
    with (dest/'first_batch_summary.json').open('w',encoding='utf-8') as stream:
        json.dump(payload['summary'],stream,ensure_ascii=False,indent=2)
    return dict(paths,summary=payload['summary'],simulation_path=str(dest/'first_batch_simulation.json'))


def simulate_first_batch_file(path,database_path,source_path='output/batch-review/sku_batch_review.xlsx'):
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        decisions=_rows(wb,FIRST_DECISION_SHEET)
        items=_rows(wb,FIRST_MEMBER_SHEET)
        stops=_rows(wb,FIRST_STOP_SHEET)
    finally:wb.close()
    source=load_workbook(source_path,read_only=True,data_only=False)
    try:
        source_snapshots={row.get('snapshot_id') for row in _rows(source,GROUP_SHEET,GROUP_COLUMNS)}
    finally:source.close()
    # This is validation only. No conversion to an importable APPROVED file is
    # provided, and no confirm flag can make this command mutate mappings.
    with readonly_database(database_path) as con:
        report=validate_first_batch(con,decisions,items,decisions=decisions,stops=stops)
        # A whole group can disappear from all three sheets without changing
        # any surviving group's digest. Compare the full exact-reference cohort
        # with the immutable source snapshot, not workbook-provided counts.
        if report.get('summary') is not None and 'REQUIRED_REVIEW_SCHEMA_MISSING' not in report['global_blockers']:
            snapshots={row.get('snapshot_id') for row in decisions}
            cohort_changed=len(snapshots)!=1 or len(source_snapshots)!=1 or snapshots!=source_snapshots
            if not cohort_changed:
                frozen=con.execute('''SELECT g.group_id,m.sku,m.expected_state
                    FROM review_queue_group g JOIN batch_review_member m
                    ON m.manifest_id=g.group_id WHERE g.snapshot_id=?''',(next(iter(snapshots)),))
                expected_groups=set();expected_members=set()
                try:
                    for row in frozen:
                        state=json.loads(row['expected_state'])
                        if state.get('mapping_basis')=='USER_WORKBOOK_EXACT_SKU' and state.get('mapping_status')!='APPROVED':
                            expected_groups.add(row['group_id'])
                            expected_members.add((state.get('api_marketplace'),row['sku']))
                    cohort_changed=not expected_groups or (
                        {row.get('group_id') for row in decisions}!=expected_groups
                        or {_key(row) for row in items}!=expected_members)
                except (ValueError,TypeError,AttributeError):
                    cohort_changed=True
            if cohort_changed:
                report['global_blockers'].append('FIRST_COHORT_SCOPE_CHANGED')
                report.update(status='BLOCKED',structure_status='BLOCKED',import_permitted=False)
    identities={_key(row):row for row in items}
    if (len(stops)!=len(items) or len({_key(row) for row in stops})!=len(stops)
            or any(_key(row) not in identities or any(row.get(k)!=identities[_key(row)].get(k)
                   for k in ('snapshot_id','group_id','member_digest','marketplace','api_marketplace','candidate_family')) for row in stops)):
        report['global_blockers'].append('STOP_VISIBLE_MEMBER_IDENTITY_CHANGED')
        report.update(status='BLOCKED',structure_status='BLOCKED',import_permitted=False)
    return report
