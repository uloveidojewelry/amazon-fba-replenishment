"""Read-only, SKU-level workbook using the user's monthly demand as its basis.

API sales windows remain comparison evidence.  This exporter deliberately does
not use ``sku_export._sanitize_rows``: that helper selects API 30-day demand and
would replace the monthly planner's explicitly chosen demand basis.
"""
from collections.abc import Mapping
from datetime import datetime
from pathlib import Path
import re
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from .export import sheet
from .sku_export import (
    INVENTORY_FIELDS, STOCK_FIELDS, WINDOWS, AUDIT_COLUMNS, _code,
    _inventory_details, _inbound_details, _inbound_exceptions,
    _known_number, _quantity, _safe_audit, _same_snapshot, _save_atomically,
)


MAIN = (
    ('sku', 'SKU'), ('asin', 'ASIN'), ('marketplace', '站点'),
    ('uploaded_monthly_units', '用户月销量'), ('monthly_days', '月销折算天数'),
    ('monthly_daily_sales', '月销日均件数'), ('monthly_status', '月销值状态'),
    ('sku_match_status', '美国站精确匹配状态'),
    ('fulfillable', 'FBA可售'), ('reserved', 'FBA预留'),
    ('unfulfillable', 'FBA不可售'), ('inventory_inbound', '入仓库存摘要'),
    ('eligible_inbound', '可抵扣在途'), ('unknown_eta_inbound', 'ETA未知在途'),
    ('coverage_days', '可售覆盖天数'), ('stockout_risk', '缺货风险'),
    ('recommended_qty', '建议补货件数'), ('status', '补货状态'),
    ('parameter_status', '参数确认状态'), ('recommendation_label', '建议用途'),
    ('anomaly_codes', '异常代码'), ('stop_reason', '拦截或停补说明'),
    ('sales_7d', 'API近7天销量'), ('sales_30d', 'API近30天销量'),
    ('sales_90d', 'API近90天销量'), ('source_rows', '上传表来源行'),
    ('stock_observed_at_beijing', '库存采集时间_北京时间'),
)
WINDOW_COLUMNS = tuple(
    (f'{name}_{days}d', f'API近{days}天{label}')
    for days in WINDOWS for name, label in (
        ('sales', '销量'), ('daily_sales', '日均'),
        ('observed_sales', '已观察销量'), ('observed_sales_days', '已验证天数'),
        ('sales_complete', '完整'),
    )
)
WINDOWS_TABLE = MAIN[:7] + WINDOW_COLUMNS + (
    ('monthly_30d_equivalent', '上传月销折合30天件数'),
    ('api_7d_month_equivalent', 'API7天折合30天件数'),
    ('api_90d_month_equivalent', 'API90天折合30天件数'),
    ('api_30d_ratio', '上传月销与API30天比值'),
    ('reference_review_codes', '销量差异人工审核提示'),
    ('api_30d_recommended_qty', '原API30天试算补货件数'),
    ('api_coverage_days', '原API30天可售覆盖天数'),
    ('recommended_qty', '月销量基准补货件数'),
    ('monthly_period_status', '上传月份确认状态'),
    ('demand_basis', '补货需求来源'), ('asof', 'API销量截止日期'),
    ('window_start', 'API销量开始日期'), ('sales_timezone', 'API销量时区'),
    ('source_rows', '上传表来源行'),
)
INVENTORY_TABLE = MAIN[:3] + tuple(zip(INVENTORY_FIELDS, (
    'FBA可售', 'FBA预留', 'FBA不可售', '准备入仓', '已发运入仓', '接收入仓',
))) + (
    ('inventory_inbound', '入仓库存摘要合计'), ('eligible_inbound', '可抵扣在途'),
    ('unknown_eta_inbound', 'ETA未知在途'), ('inventory_complete', '库存数据完整'),
    ('inbound_complete', '货件数量已核对'), ('inbound_rule', '在途抵扣规则'),
    ('stock_observed_at_beijing', '库存采集时间_北京时间'),
    ('stock_observed_at', '库存采集时间_原始时区'), ('inventory_source', '库存来源'),
)
INBOUND_TABLE = MAIN[:3] + (
    ('shipment_id', '货件编号'), ('fnsku', 'FNSKU'),
    ('quantity_shipped', '原始发运件数'), ('quantity_received', '原始接收件数'),
    ('remaining', '剩余在途件数'), ('eta', '人工确认ETA'),
    ('status', '货件状态'), ('code', '数量异常代码'), ('inbound_rule', '在途抵扣规则'),
    ('inventory_note', '在途说明'),
)
ANOMALY_TABLE = MAIN[:3] + (
    ('code', '异常代码'), ('reason', '异常或待确认内容'),
    ('blocking', '是否阻断补货'), ('source', '证据来源'),
)
OMITTED_TABLE = MAIN[:3] + (
    ('sales_7d', 'API近7天销量'), ('sales_30d', 'API近30天销量'),
    ('sales_90d', 'API近90天销量'), ('fulfillable', 'FBA可售'),
    ('reserved', 'FBA预留'), ('inventory_inbound', '入仓库存摘要'),
    ('stockout_risk', 'API参考缺货风险'), ('monthly_status', '月销量状态'),
    ('status', '补货状态'),
    ('recommended_qty', '月销量基准补货件数'), ('stop_reason', '处理说明'),
)
INPUT_AUDIT_KEYS = {
    'file_name', 'filename', 'sheet_name', 'sheet', 'sha256', 'file_sha256',
    'input_rows', 'row_count', 'unique_skus', 'sku_count', 'duplicate_skus',
    'duplicate_rows', 'invalid_rows', 'matched_skus', 'unmatched_skus',
    'matched_sku_count', 'unmatched_sku_count', 'monthly_units_total',
    'monthly_days', 'source_file', 'source_sha256', 'valid_rows',
    'input_data_rows', 'blank_rows', 'record_count', 'valid_record_count',
    'invalid_record_count', 'duplicate_identical_skus', 'duplicate_conflict_skus',
    'header_row', 'monthly_period_status', 'quantity_header',
}
SETTINGS_KEYS = {
    'asof', 'window_start', 'window_end', 'sales_window_start', 'sales_window_end',
    'sales_timezone', 'marketplace', 'marketplace_id', 'monthly_days',
    'monthly_period_status', 'monthly_marketplace_status', 'monthly_units_status',
    'lead_days', 'target_days', 'safety_days', 'moq', 'case_pack', 'inbound_rule',
    'max_stock_age_days', 'parameters_confirmed', 'confirmed_by', 'confirmed_at',
    'inventory_observed_at', 'inbound_observed_at', 'sales_scope',
}
SUMMARY_LABELS = {
    'cohort_count': '上传月销量SKU数', 'matched_sku_count': '匹配美国站SKU数',
    'unmatched_sku_count': '未匹配美国站SKU数', 'monthly_units_total': '上传月销量总件数',
    'recommended_sku_count': '正数试算补货SKU数', 'recommended_units': '试算补货总件数',
    'blocked_sku_count': '数据阻断SKU数', 'stopped_sku_count': '已停补SKU数',
    'inventory_positive_sku_count': '有FBA实物库存SKU数',
    'omitted_risk_sku_count': '月销量表外风险SKU数',
    'parameter_status': '参数确认状态', 'excel_path': '文件位置',
}
COMMENTS = {
    '用户月销量': '上传表提供的月销量；未说明具体月份时按30天折算，属于待确认假设，不冒充API近30天销量。',
    '月销日均件数': '用户月销量 / 月销折算天数。补货以此为基准，API 7/30/60/90天销量仅为对比证据。',
    'FBA预留': '客户订单预留、FC转运及处理库存；单独展示，不重复加入可售或在途抵扣。',
    '入仓库存摘要': '准备、发运及接收入仓摘要的合计；与货件明细重叠，不与可抵扣在途再次相加。',
    '可抵扣在途': '只使用按配置规则核对后的唯一货件余额抵扣一次；ETA未知不擅自当作已确认到货。',
    '建议补货件数': 'STOPPED强制0；数据异常为留空。SCENARIO_ONLY仅供人工试算，不是正式采购依据。',
    '剩余在途件数': '冲突行保留原始发运/接收数量，余额留空；禁止将负差额截断为0。',
}
REFERENCE_REASONS = {
    'MONTHLY_DIFFERS_FROM_API30': '上传月销折合30天与API近30天相差至少2倍，或一方为0；请核实期间、站点和单位。',
    'RECENT_7D_HIGHER_THAN_MONTHLY': '近7天销量折合30天达到上传月销的2倍或以上；请人工检查近期增长及缺货影响。',
    'API30_ZERO_RATIO_UNDEFINED': 'API近30天已确认销量为0；月销/API比值无定义，留空，不能填0或无穷大。',
}


def _report_sheet(wb, name, columns, rows):
    translated = [{label: row.get(key) for key, label in columns} for row in rows]
    ws = sheet(wb, name, [label for _, label in columns], translated)
    for cell in ws[1]:
        if cell.value in COMMENTS:
            cell.comment = Comment(COMMENTS[cell.value], 'Uloveido')
    return ws


def _beijing(value):
    try:
        instant = datetime.fromisoformat(value.replace('Z', '+00:00'))
        return instant.astimezone(ZoneInfo('Asia/Shanghai')).isoformat() if instant.tzinfo else None
    except (AttributeError, TypeError, ValueError):
        return None


def _sources(value):
    """Permit Excel row references, never stringify arbitrary nested payloads."""
    if _quantity(value):
        return str(int(value))
    if isinstance(value, str):
        return value if re.fullmatch(r'[A-Za-z0-9_\u4e00-\u9fff !:$,;\-\[\].]{1,300}', value) else None
    if isinstance(value, (tuple, list)):
        refs = [_sources(item) for item in value]
        return '; '.join(item for item in refs if item) or None
    return None


def _sort_key(row):
    risk = {'STOCKOUT_BEFORE_REPLENISHMENT': 0, 'STOCKOUT_BEFORE_ARRIVAL': 0,
            'SAFETY_STOCK_RISK': 1, 'UNKNOWN_DATA': 2, 'COVERED': 3,
            'NO_RECENT_DEMAND': 4, 'NO_SALES': 4}.get(row.get('stockout_risk'), 2)
    quantity = row.get('recommended_qty')
    return (risk, 0 if _quantity(quantity) and quantity > 0 else 1,
            0 if any(_quantity(row.get(k)) and row[k] > 0 for k in STOCK_FIELDS) else 1,
            0 if _quantity(row.get('inventory_inbound')) and row['inventory_inbound'] > 0 else 1,
            -(row.get('uploaded_monthly_units') or 0), row['sku'])


def _normalize(rows, details):
    allowed = {key for columns in (MAIN, WINDOWS_TABLE, INVENTORY_TABLE) for key, _ in columns}
    allowed.update({'daily_sales', 'parameters_confirmed', 'data_complete', 'planning_data_complete',
                    'lead_days', 'target_days', 'safety_days', 'moq', 'case_pack',
                    'matched', 'sku_matched', 'api_sku_matched', 'monthly_sku_matched'})
    values, anomalies, seen = [], [], set()
    for original in rows:
        if not isinstance(original, Mapping):
            raise ValueError('Monthly report requires mapping rows')
        sku = original.get('sku')
        if original.get('marketplace') != 'US' or not isinstance(sku, str) or not sku.strip():
            raise ValueError('Monthly report requires a US Marketplace + Seller SKU identity')
        if sku in seen:
            raise ValueError('Duplicate Marketplace + Seller SKU in monthly report')
        seen.add(sku)
        row = {key: original.get(key) for key in allowed if key in original}
        row['source_rows'] = _sources(original.get('source_rows'))
        row['demand_basis'] = 'USER_MONTHLY_WORKBOOK'
        row['reference_review_codes'] = '; '.join(
            code for code in original.get('reference_review_codes', []) if _code(code))
        row['stock_observed_at_beijing'] = _beijing(row.get('stock_observed_at'))
        extra = details.get(sku, {})
        if _same_snapshot(row.get('stock_observed_at'), extra.get('observed_at')):
            for key in INVENTORY_FIELDS[3:]:
                row.setdefault(key, extra.get(key))
        for key in INVENTORY_FIELDS + ('inventory_inbound', 'eligible_inbound', 'unknown_eta_inbound'):
            if not _quantity(row.get(key)):
                row[key] = None
        for days in WINDOWS:
            complete = (original.get(f'sales_complete_{days}d') is True
                        and original.get(f'observed_sales_days_{days}d') == days
                        and _quantity(original.get(f'sales_{days}d')))
            row[f'sales_complete_{days}d'] = complete
            row[f'sales_{days}d'] = original.get(f'sales_{days}d') if complete else None
            row[f'daily_sales_{days}d'] = row[f'sales_{days}d'] / days if complete else None
            for key in (f'observed_sales_{days}d', f'observed_sales_days_{days}d'):
                if not _quantity(row.get(key)):
                    row[key] = None
        for key in ('uploaded_monthly_units', 'monthly_daily_sales', 'monthly_days',
                    'daily_sales', 'coverage_days', 'api_coverage_days',
                    'monthly_30d_equivalent',
                    'api_7d_month_equivalent', 'api_90d_month_equivalent', 'api_30d_ratio'):
            if not _known_number(row.get(key)):
                row[key] = None
        row['inventory_complete'] = (original.get('inventory_complete') is True
                                     and all(row[k] is not None for k in INVENTORY_FIELDS))
        row['inbound_complete'] = original.get('inbound_complete') is True
        codes = {_code(code) for code in original.get('anomaly_codes', [])
                 if isinstance(code, str)} if isinstance(original.get('anomaly_codes'), (list, tuple)) else set()
        codes.discard(None)
        has_blocker = False
        for issue in original.get('anomalies') or []:
            if not isinstance(issue, Mapping) or not _code(issue.get('code')):
                continue
            code = issue['code']
            codes.add(code)
            has_blocker = has_blocker or issue.get('blocking') is True
            anomalies.append({'marketplace': 'US', 'sku': sku, 'asin': row.get('asin'),
                              'code': code, 'reason': issue.get('reason'),
                              'blocking': issue.get('blocking') is True, 'source': issue.get('source')})
        for code in original.get('reference_review_codes') or []:
            if _code(code):
                anomalies.append({'marketplace': 'US', 'sku': sku, 'asin': row.get('asin'),
                                  'code': code, 'reason': REFERENCE_REASONS.get(code, '上传月销量与API证据存在差异，请人工审核。') + ' 仅供参考，未自动加权或解除安全拦截。',
                                  'blocking': False, 'source': 'Monthly/API sales reference comparison'})
        status = original.get('status') or original.get('replenishment_status') or 'UNKNOWN'
        row['status'] = status
        qty = original.get('recommended_qty')
        row['recommended_qty'] = qty if _quantity(qty) else None
        confirmed = (original.get('parameters_confirmed') is True
                     and original.get('parameter_status') == 'CONFIRMED'
                     and original.get('monthly_period_status') == 'USER_CONFIRMED_PERIOD')
        row['parameter_status'] = 'CONFIRMED' if confirmed else 'SCENARIO_ONLY'
        matched = all(original.get(key) is not False for key in
                      ('matched', 'sku_matched', 'api_sku_matched', 'monthly_sku_matched'))
        if str(original.get('monthly_status') or '').upper() in {
                'UNMATCHED', 'UNKNOWN_SKU', 'SKU_NOT_FOUND', 'SKU_NOT_FOUND_IN_US', 'NOT_FOUND'}:
            matched = False
        if 'SKU_NOT_FOUND_US' in codes:
            matched = False
        row['_matched'] = matched
        row['sku_match_status'] = 'MATCHED_US' if matched else 'NOT_FOUND_US'
        monthly_valid = (row['uploaded_monthly_units'] is not None
                         and row['monthly_daily_sales'] is not None
                         and row['monthly_days'] is not None and row['monthly_days'] > 0
                         and row.get('monthly_status') == 'VALID')
        valid = (matched and monthly_valid and row['inventory_complete'] and row['inbound_complete']
                 and row['sales_complete_30d'] and row['sales_complete_90d']
                 and row['eligible_inbound'] is not None
                 and original.get('planning_data_complete') is True
                 and not has_blocker and status in {'ACTIVE', 'SCENARIO_ONLY'})
        if status == 'STOPPED':
            row['recommended_qty'] = 0
            row['recommendation_label'] = '已明确SKU停补：补货数量强制为0'
        elif not valid:
            row['recommended_qty'] = None
            row['recommendation_label'] = '数据异常或月销量未知：暂不能计算'
        else:
            row['recommendation_label'] = ('人工审核建议；不执行采购或Amazon业务写入' if confirmed else
                '参数或月份未确认：仅供试算，不得作为正式采购依据')
        row['anomaly_codes'] = '; '.join(sorted(codes))
        values.append(row)
    values.sort(key=_sort_key)
    return values, anomalies


def _safe_input_audit(audit):
    result = []
    if not isinstance(audit, Mapping):
        return result
    for key in sorted(INPUT_AUDIT_KEYS):
        value = audit.get(key)
        if key.endswith('sha256') or key == 'sha256':
            if not isinstance(value, str) or not re.fullmatch('[a-fA-F0-9]{64}', value):
                continue
        elif key == 'monthly_period_status':
            if not _code(value):
                continue
        elif key == 'quantity_header':
            if type(value) not in (str, int, float) or isinstance(value, str) and len(value) > 100:
                continue
        elif key in {'file_name', 'filename', 'sheet_name', 'sheet', 'source_file'}:
            if not isinstance(value, str) or len(value) > 200:
                continue
            if key != 'sheet_name' and key != 'sheet':
                value = Path(value).name
        elif not _quantity(value):
            continue
        result.append({'item': key, 'value': value})
    for code in audit.get('issues') or []:
        if _code(code):
            result.append({'item': '上传表格式待确认', 'value': code})
    return result


def _omitted_rows(rows):
    values = []
    seen = set()
    for original in rows or []:
        sku = original.get('sku') if isinstance(original, Mapping) else None
        if not isinstance(sku, str) or not sku.strip() or original.get('marketplace') != 'US':
            raise ValueError('Omitted monthly cohort requires US SKU identities')
        if sku in seen:
            raise ValueError('Duplicate omitted SKU')
        seen.add(sku)
        row = {key: original.get(key) for key, _ in OMITTED_TABLE}
        for key in ('sales_7d', 'sales_30d', 'sales_90d', 'fulfillable', 'reserved', 'inventory_inbound'):
            if not _quantity(row.get(key)):
                row[key] = None
        if original.get('status') == 'STOPPED':
            row.update(monthly_status='MONTHLY_SALES_NOT_PROVIDED', status='STOPPED',
                       recommended_qty=0,
                       stop_reason='已明确SKU停补，补货数量强制为0；未提供月销量不解除停补。' +
                                   (str(original['stop_reason']) if original.get('stop_reason') else ''))
        else:
            row.update(monthly_status='MONTHLY_SALES_NOT_PROVIDED', recommended_qty=None,
                       stop_reason='不在本次月销量表中；月销量未知，未按0处理，不生成月销基准补货建议。')
        values.append(row)
    values.sort(key=lambda row: (-(row.get('sales_30d') or 0), row['sku']))
    return values


def export_monthly_report(rows, path, con, settings, *, input_audit,
                          collection_audit=None, omitted_rows=None):
    """Publish an auditable Excel without modifying inputs or the source database."""
    if not isinstance(settings, Mapping):
        raise ValueError('Monthly report settings must be a mapping')
    if settings.get('marketplace', 'US') != 'US':
        raise ValueError('Monthly report is limited to the US marketplace')
    values, anomalies = _normalize(rows, _inventory_details(con, 'US'))
    cohort = {row['sku'] for row in values}
    inbound = _inbound_details(con, 'US', cohort)
    exceptions = _inbound_exceptions(con, 'US', cohort)
    by_sku = {row['sku']: row for row in values}
    exception_keys = {(r['sku'], r['shipment_id']) for r in exceptions}
    for item in exceptions:
        row = by_sku[item['sku']]
        row.update(inbound_complete=False, eligible_inbound=None, unknown_eta_inbound=None)
        if row['status'] != 'STOPPED':
            row.update(recommended_qty=None, status='DATA_REVIEW_REQUIRED',
                       recommendation_label='货件数量异常：剩余在途未知，暂不能计算')
        code = _code(item.get('code')) or 'INBOUND_QUANTITY_CONFLICT'
        row['anomaly_codes'] = '; '.join(sorted(set(row['anomaly_codes'].split('; ')) - {''} | {code}))
        anomalies.append({'marketplace': 'US', 'sku': row['sku'], 'asin': row.get('asin'),
                          'code': code, 'reason': '货件接收与发运数量冲突；保留原始值，剩余在途为UNKNOWN。',
                          'blocking': True, 'source': 'SP-API Inbound quantity exception'})
    anomalies = list({(row['sku'], row['code']): row for row in anomalies}.values())
    omitted = _omitted_rows(omitted_rows)
    if cohort & {row['sku'] for row in omitted}:
        raise ValueError('Monthly cohort and omitted cohort overlap')
    positive = [r for r in values if _quantity(r['recommended_qty']) and r['recommended_qty'] > 0]
    summary = {
        'cohort_count': len(values), 'matched_sku_count': sum(r['_matched'] for r in values),
        'unmatched_sku_count': sum(not r['_matched'] for r in values),
        'monthly_units_total': sum(r['uploaded_monthly_units'] for r in values
                                   if r.get('uploaded_monthly_units') is not None),
        'recommended_sku_count': len(positive), 'recommended_units': sum(r['recommended_qty'] for r in positive),
        'blocked_sku_count': sum(r['recommended_qty'] is None for r in values),
        'stopped_sku_count': sum(r['status'] == 'STOPPED' for r in values),
        'inventory_positive_sku_count': sum(any(_quantity(r.get(k)) and r[k] > 0 for k in STOCK_FIELDS) for r in values),
        'omitted_risk_sku_count': len(omitted),
        'parameter_status': 'CONFIRMED' if values and all(r['parameter_status'] == 'CONFIRMED' for r in values) else 'SCENARIO_ONLY',
        'excel_path': str(Path(path)),
    }
    wb = Workbook()
    wb.remove(wb.active)
    _report_sheet(wb, '补货清单', MAIN, values)
    _report_sheet(wb, '建议补货数量', MAIN, positive)
    _report_sheet(wb, '缺货风险', MAIN, [r for r in values if r.get('stockout_risk') in {
        'STOCKOUT_BEFORE_REPLENISHMENT', 'STOCKOUT_BEFORE_ARRIVAL',
        'SAFETY_STOCK_RISK', 'UNKNOWN_DATA'}])
    _report_sheet(wb, '月销与7_30_90对比', WINDOWS_TABLE, values)
    _report_sheet(wb, 'FBA库存明细', INVENTORY_TABLE, values)
    inbound_rows = []
    for original in inbound:
        if (original['sku'], original['shipment_id']) in exception_keys:
            continue
        row = dict(original)
        row.update(asin=by_sku[row['sku']].get('asin'),
                   remaining=row.get('units') if _quantity(row.get('units')) else None,
                   inbound_rule=by_sku[row['sku']].get('inbound_rule'),
                   inventory_note='货件余额与库存入仓摘要重叠，禁止再次相加；仅按确认规则抵扣一次。')
        row['eta'] = row.get('eta') or None
        inbound_rows.append(row)
    for original in exceptions:
        row = {key: original.get(key) for key, _ in INBOUND_TABLE}
        for key in ('quantity_shipped', 'quantity_received'):
            if not _quantity(row.get(key)):
                row[key] = None
        row.update(asin=by_sku[row['sku']].get('asin'), remaining=None, eta=None,
                   code=_code(row.get('code')) or 'INBOUND_QUANTITY_CONFLICT',
                   inbound_rule=by_sku[row['sku']].get('inbound_rule'),
                   inventory_note='数量冲突：保留原始发运与接收件数，剩余在途未知，不截断为0。')
        inbound_rows.append(row)
    inbound_rows.sort(key=lambda r: (0 if r.get('code') else 1, r['sku'], r['shipment_id']))
    _report_sheet(wb, '在途货件明细', INBOUND_TABLE, inbound_rows)
    for row in values:
        if row['recommended_qty'] is None and not any(a['sku'] == row['sku'] for a in anomalies):
            anomalies.append({'marketplace': 'US', 'sku': row['sku'], 'asin': row.get('asin'),
                              'code': 'MONTHLY_PLAN_BLOCKED', 'reason': row.get('stop_reason') or '请核对月销量、SKU匹配、销量完整性、库存及在途证据。',
                              'blocking': True, 'source': 'Monthly report validation'})
    _report_sheet(wb, '异常与待确认', ANOMALY_TABLE, anomalies)
    _report_sheet(wb, '已停补SKU', MAIN, [r for r in values if r['status'] == 'STOPPED'])
    _report_sheet(wb, '月销表外风险SKU', OMITTED_TABLE, omitted)
    lead, target, safety = (settings.get(key, default) for key, default in
                            (('lead_days', 35), ('target_days', 75), ('safety_days', 0)))
    horizon = lead + target + safety if all(_known_number(v) for v in (lead, target, safety)) else None
    parameters = [
        {'item': '重要提示', 'value': '参数、月份、站点或单位未经确认时仅供试算，不得作为正式采购依据。'},
        {'item': '参数确认状态', 'value': summary['parameter_status']},
        {'item': '补货需求来源', 'value': '用户上传月销量；API 7/30/60/90天销量作为对比，不替换用户月销量。'},
        {'item': '月销量期间假设', 'value': '上传表未注明月份时按30天折算；月份、站点及单位须由用户确认。'},
        {'item': '补货覆盖总天数', 'value': horizon},
        {'item': '计算公式', 'value': 'max(0, ceil(月销量/折算天数 × (目标覆盖天数+运输入仓天数+安全天数) - FBA可售 - 唯一可抵扣在途))；再按MOQ/装箱数取整。'},
        {'item': '美国站范围', 'value': '美国站 US / ATVPDKIKX0DER；不含加拿大。上传表所属站点仍为待确认假设。'},
        {'item': 'API销量口径', 'value': 'FBA/AFN订单件数，排除取消订单；不含MFN自配送，不扣退货或退款净额。'},
        {'item': 'API销量日期', 'value': str(settings.get('sales_window_start', settings.get('window_start', '2026-07-16'))) + ' 至 ' + str(settings.get('sales_window_end', settings.get('asof', '2026-10-13')))},
        {'item': 'API销量时区', 'value': settings.get('sales_timezone', 'America/Los_Angeles')},
        {'item': '时间显示', 'value': '库存采集时间同时保留原始时区及北京时间 Asia/Shanghai；不把销售日期切换到北京时间。'},
        {'item': '预留库存规则', 'value': 'FBA预留独立展示；不重复计为可售或在途，不再次从可售扣除。'},
        {'item': '在途抵扣规则', 'value': '库存入仓摘要与货件余额只核对不相加；仅按已配置规则对唯一货件余额抵扣一次。'},
        {'item': '未知与冲突规则', 'value': '未知为留空或UNKNOWN，绝不当作0；原API未结订单、缺失或数量冲突保持阻断。'},
        {'item': '缺货风险代码', 'value': 'STOCKOUT_BEFORE_REPLENISHMENT=预计到货前缺货；SAFETY_STOCK_RISK=安全库存风险；COVERED=当前覆盖充足；NO_RECENT_DEMAND=用户月销为0；UNKNOWN_DATA=缺少有效证据。'},
        {'item': '状态代码', 'value': 'SCENARIO_ONLY=仅供试算；DATA_REVIEW_REQUIRED=数据待核实；STOPPED=已停补；VALID=上传月销数值有效，不代表参数或系列已批准；NOT_FOUND_US=未精确匹配美国站。'},
        {'item': '对比提示规则', 'value': 'API7天/90天均折合30天比较；月销/API比值的分母为0或未知时留空。2倍差异仅用于人工筛选，不是预测权重，不自动改变建议件数。'},
        {'item': '停补规则', 'value': '明确生效的直接SKU停补强制0；本次不使用系列映射审批或系列停补。原配置保留，不能自动恢复。'},
        {'item': '操作边界', 'value': '只读报告；不创建FBA货件，不修改Listing、库存、价格，不提交采购。'},
    ]
    parameters.extend({'item': key, 'value': settings[key]} for key in sorted(SETTINGS_KEYS)
                      if key in settings and type(settings[key]) in (str, int, float, bool))
    ws = _report_sheet(wb, '参数与计算说明', (('item', '项目'), ('value', '内容')), parameters)
    ws.column_dimensions['B'].width = 110
    for row in ws.iter_rows(min_row=2, max_row=5):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='FFF2CC')
            cell.font = Font(bold=True, color='9C5700')
    safe_audit, safe_summary = _safe_audit(collection_audit)
    audit_rows = _safe_input_audit(input_audit)
    audit_rows.extend({'item': key, 'value': value} for key, value in sorted(safe_summary.items()))
    for key in ('inventory_observed_at', 'inbound_observed_at'):
        converted = _beijing(settings.get(key))
        if converted:
            audit_rows.append({'item': key + '_北京时间', 'value': converted})
    _report_sheet(wb, '数据来源与审计', (('item', '审计项目'), ('value', '审计值')), audit_rows)
    _report_sheet(wb, 'API分页与采集审计', tuple((k, k) for k in AUDIT_COLUMNS), safe_audit)
    _report_sheet(wb, '汇总统计', (('item', '统计项目'), ('value', '统计值')),
                  [{'item': SUMMARY_LABELS.get(k, k), 'value': v} for k, v in summary.items()])
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    _save_atomically(wb, output)
    return summary
