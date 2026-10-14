"""Human review workbooks for fixed SKU sets; this module never approves a SKU.

Only operational fields are exported. Raw Amazon responses, customer fields,
credentials and arbitrary caller metadata are intentionally not serialized.
"""
from collections import Counter
from pathlib import Path
import math

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


DECISION_SHEET = '批量审核决定'
MEMBERS_SHEET = 'SKU明确集合'
STOP_SHEET = '停补导入'
DECISION_COLUMNS = (
    'snapshot_id', 'group_id', 'manifest_id', 'member_digest', 'marketplace', 'api_marketplace',
    'candidate_family', 'sku_count', 'batch_eligible', 'decision', 'approved_family', 'reviewer', 'reason',
)
STOP_COLUMNS = (
    'marketplace', 'api_marketplace', 'sku', 'family', 'candidate_family',
    'stop', 'reason', 'effective_date', 'review_status', 'reviewer',
)
GROUP_COLUMNS = (
    'priority', 'priority_score', 'priority_reasons', 'snapshot_id', 'group_id', 'member_digest',
    'marketplace', 'api_marketplace', 'candidate_family', 'mapping_basis', 'naming_format',
    'sku_count', 'exact_match_count', 'stock_positive_count', 'recent_sales_positive_count',
    'unknown_status_count', 'batch_eligible',
)
MEMBER_COLUMNS = (
    'snapshot_id', 'group_id', 'member_digest', 'sku', 'asin', 'marketplace',
    'api_marketplace', 'candidate_family', 'family', 'mapping_basis',
    'mapping_status', 'replenishment_status', 'listing_status', 'recent_sales_status',
    'listing_evidence', 'fba_stock_status', 'history_status', 'naming_format',
    'priority', 'priority_reasons', 'fulfillable', 'reserved', 'unfulfillable',
    'inbound_working', 'inbound_shipped', 'inbound_receiving',
    'recent_units_observed', 'inventory_observed_at',
    'first_discovered_at', 'last_seen_at', 'review_status', 'recommended_qty',
)
ANOMALY_COLUMNS = (
    'priority', 'severity', 'group_id', 'marketplace', 'api_marketplace', 'sku',
    'candidate_family', 'code', 'message', 'source',
)
AUDIT_COLUMNS = (
    'at', 'action', 'marketplace', 'api_marketplace', 'sku', 'group_id', 'snapshot_id',
    'reviewer', 'original_family', 'approved_family', 'original_mapping_status',
    'mapping_status', 'original_replenishment_status', 'replenishment_status',
    'original_stop', 'stop', 'reason',
)
SUMMARY_KEYS = {
    'format_warning_sku_count', 'exception_sku_count',
    'snapshot_id', 'observed_at', 'asof', 'source', 'collection_status',
    'sku_count', 'total_sku_count', 'pending_sku_count', 'pending_review_count',
    'approved_sku_count', 'group_count', 'pending_group_count',
    'review_group_count', 'review_action_count', 'reviewable_group_count',
    'candidate_family_count', 'unknown_mapping_count', 'unknown_sku_count',
    'exact_match_sku_count', 'exact_sku_count', 'exact_group_count',
    'exact_match_group_count', 'conflict_sku_count', 'conflict_group_count',
    'exact_exception_sku_count', 'exact_review_action_count', 'exact_actions_saved',
    'stopped_sku_count', 'stopped_family_count', 'stop_inherited_count',
    'risk_sku_count', 'active_listing_count', 'recent_sales_count',
    'fba_stock_count', 'unknown_activity_count', 'new_sku_count',
    'listing_unknown_count', 'naming_candidate_sku_count', 'fullmatch_candidate_sku_count',
    'manual_sku_review_count', 'batch_review_count', 'review_actions_saved',
    'review_reduction', 'review_reduction_pct', 'review_reduction_percent',
    'full_history_sync', 'daily_schedule_deployed',
}
HEADER_DESCRIPTIONS = {
    'priority': '审核优先级；顺序来自审核队列，不代表已获批准。',
    'priority_reasons': '优先级依据；风险判断仅使用已获取的证据。',
    'snapshot_id': '固定审核快照标识。不得修改。',
    'group_id': '本次明确 SKU 集合标识。不得修改。',
    'member_digest': '明确 SKU 成员集合校验值。不得修改。',
    'manifest_id': '后端固定集合标识，与 group_id 一致。不得修改。',
    'candidate_family': '候选系列，仅供人工审核，不是正式映射。',
    'family': '已经人工确认的正式系列；未知时留空。',
    'decision': '仅明确填写 APPROVED 的行可申请批准。默认 PENDING 不产生映射变更。',
    'approved_family': '人工确认的系列，APPROVED 时必须填写。',
    'reviewer': '人工审核人的姓名或可追溯标识，作出决定时必须填写。',
    'reason': '人工决定或停补原因，保存在审计记录中。',
    'listing_status': '当前销售状态证据。UNKNOWN 表示未取得可信在售或下架证据。',
    'recent_sales_status': '近期销量证据。未采集完整窗口时 UNKNOWN 不表示零销量。',
    'fba_stock_status': '库存证据。UNKNOWN 不表示库存为零。',
    'mapping_status': '系列映射审核状态，与补货状态分别存储。',
    'replenishment_status': '补货安全状态。STOPPED 最终建议量必须为零。',
    'recommended_qty': '本审核工作簿不生成实际补货建议，全部为零。',
    'stop': '人工确认是否停补：TRUE / FALSE。空白表示尚未确认。FALSE 不自动恢复既有停补。',
    'review_status': '仅明确填写 APPROVED 的停补行可导入。默认 UNAPPROVED。',
    'effective_date': '停补生效日期，YYYY-MM-DD。',
    'sku_count': '只批准 SKU明确集合 工作表中属于本组的明确 SKU；未来同前缀 SKU 不受批准。',
    'batch_eligible': 'TRUE 表示候选系列明确，可以审核固定集合；FALSE 表示只按格式分类，不得批量批准未知系列。',
    'recent_units_observed': '已采集小批量窗口中的观察销量；未知时留空，不代表完整 90 天销量。',
    'listing_evidence': '可信销售状态的来源；缺少证据时 UNKNOWN，不从库存存在推测在售。',
    'naming_format': 'SKU 命名格式，仅用于分类；禁止按第一个连字符截取系列。',
    'history_status': '历史或未再发现状态证据；未知不等于已下架。',
}


def _value(value):
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValueError('Review cell exceeds the Excel text limit')
        return value
    if isinstance(value, (list, tuple, set)) and all(isinstance(v, str) for v in value):
        return _value('; '.join(sorted(value) if isinstance(value, set) else value))
    return None


def _sheet(wb, name, columns, rows, *, editable=()):
    ws = wb.create_sheet(name)
    ws.append(list(columns))
    for row_index, row in enumerate(rows, 2):
        ws.append([_value(row.get(key)) for key in columns])
        for col_index in range(1, len(columns) + 1):
            cell = ws.cell(row_index, col_index)
            # Amazon/user provided strings must never become formulas.
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if columns[col_index - 1] in editable:
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for index, key in enumerate(columns, 1):
        cell = ws.cell(1, index)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(wrap_text=True)
        if key in HEADER_DESCRIPTIONS:
            cell.comment = Comment(HEADER_DESCRIPTIONS[key], 'Uloveido')
        ws.column_dimensions[get_column_letter(index)].width = (
            46 if key in ('priority_reasons', 'message', 'basis', 'mapping_basis')
            else 36 if key in ('member_digest', 'snapshot_id', 'group_id')
            else max(16, min(30, len(key) + 4)))
    ws.row_dimensions[1].height = 30
    return ws


def _validation(ws, columns, key, options):
    col = get_column_letter(columns.index(key) + 1)
    validation = DataValidation(type='list', formula1='"' + ','.join(options) + '"')
    validation.allow_blank = True
    validation.errorTitle = '请使用允许的审核值'
    validation.error = '此字段必须使用列表中的值。'
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(f'{col}2:{col}{max(2, ws.max_row)}')


def _items(queue, groups):
    """Associate each identity with exactly one persisted group, never a prefix."""
    members = queue.get('items') or []
    indexed = {str(group['group_id']): group for group in groups}
    if len(indexed) != len(groups):
        raise ValueError('Duplicate review group identity')
    rows, identities = [], set()
    snapshot_id = queue.get('snapshot_id') or (queue.get('summary') or {}).get('snapshot_id')
    for original in members:
        row = {key: original.get(key) for key in MEMBER_COLUMNS}
        group = indexed.get(str(row['group_id']))
        if group is None:
            raise ValueError('Review SKU has no fixed review group')
        api_marketplace = original.get('api_marketplace') or group.get('api_marketplace')
        marketplace = original.get('marketplace') or group.get('marketplace')
        identity = (api_marketplace or marketplace, original.get('sku'))
        if not identity[0] or not isinstance(identity[1], str) or not identity[1]:
            raise ValueError('Review SKU is missing its marketplace identity')
        if identity in identities:
            raise ValueError('Duplicate SKU in the fixed review snapshot')
        identities.add(identity)
        row.update(snapshot_id=original.get('snapshot_id') or snapshot_id,
                   member_digest=group.get('member_digest'), marketplace=marketplace,
                   api_marketplace=api_marketplace, recommended_qty=0)
        row['family'] = original.get('family', original.get('formal_family'))
        row['mapping_basis'] = original.get('mapping_basis', original.get('basis'))
        for key in ('listing_status', 'recent_sales_status', 'fba_stock_status', 'history_status'):
            row[key] = row[key] or 'UNKNOWN'
        rows.append(row)
    counts = Counter(str(row['group_id']) for row in rows)
    for group in groups:
        if group.get('sku_count') != counts[str(group['group_id'])]:
            raise ValueError('Review group does not match its explicit SKU membership')
    order = {str(group['group_id']): position for position, group in enumerate(groups)}
    return sorted(rows, key=lambda row: (order[str(row['group_id'])], str(row['sku'])))


def _instructions(wb):
    lines = [
        ('审核顺序', '先查看 系列优先审核 和 数据异常，精确匹配上传表也只是候选，不会自动批准。'),
        ('精确匹配入口', '精确匹配优先 列出与上传表精确匹配且仍待审核的 SKU，包括格式异常。依然必须人工确认，不能把精确匹配当作批准。'),
        ('固定集合', '在 SKU明确集合 查看每组全部成员。批量批准只作用于本快照明确列出的 SKU，不能批准未来同前缀 SKU。'),
        ('批准映射', '仅修改 批量审核决定 的黄色 decision、approved_family、reviewer、reason 列。只有 APPROVED 行申请批准。'),
        ('默认安全状态', 'PENDING / UNAPPROVED 不产生批准；不修改已批准映射，不自动恢复停补。建议补货数量全部为 0。'),
        ('禁止修改标识', '不得修改 snapshot_id、group_id、member_digest 或成员集合；导入端会核验原始快照。需要变更成员时重新生成队列。'),
        ('未知证据', 'UNKNOWN 或空白数量表示数据未确认，不表示销量或库存为 0；库存接口返回 SKU 不证明其当前在售。'),
        ('销量范围', '仅使用当前已保存的小批量销量证据，不执行完整 90 天同步，不把样本缺失解释为零销量。'),
        ('停补导入', '在 停补导入 填写明确 SKU 或已批准系列、stop、reason、effective_date、reviewer，并显式将 review_status 改为 APPROVED。'),
        ('系列联动', '同站点某个已批准系列存在有效停补记录，则该系列全部已确认子体停补；不能因销量增长或库存不足解除停补。'),
        ('候选停补', 'candidate_family 仅作参考。候选映射不得当成正式系列导入；归属未知时仍拦截实际补货。'),
        ('恢复停补', '填写 stop=FALSE 不自动恢复历史停补。恢复必须通过明确人工确认流程并留存审核人及审计。'),
        ('只读边界', '导入仅改变本地审核与停补记录，不修改 Amazon Listing、库存或价格，不创建 FBA 货件。'),
        ('状态分离', 'mapping_status 表示映射审核状态，replenishment_status 表示补货安全状态，listing_status 表示当前销售证据。'),
        ('工作量口径', '审核统计中的 review_actions_saved / review_reduction_pct 仅估算决定次数的减少，不代表实际审核耗时。未知系列格式桶仍需逐 SKU 确认，不能视为一次批准全部。'),
    ]
    ws = _sheet(wb, '使用说明', ('项目', '操作说明'), [{'项目': title, '操作说明': text} for title, text in lines])
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 110
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row[0].row].height = 42


def export_batch_review(queue, output_dir):
    """Export a persisted queue and blank human decisions without mutating it.

    The caller persists ``snapshot_id`` / group membership before export. Queue
    group order is already the evidence-based review order and is preserved.
    """
    groups = list(queue.get('groups') or [])
    items = _items(queue, groups)
    snapshot_id = queue.get('snapshot_id') or (queue.get('summary') or {}).get('snapshot_id')
    group_rows = [dict(group, snapshot_id=group.get('snapshot_id') or snapshot_id) for group in groups]
    decisions = [{key: group.get(key) for key in DECISION_COLUMNS} for group in group_rows]
    for row in decisions:
        row.update(manifest_id=row['group_id'], decision='PENDING', approved_family=None, reviewer=None, reason=None)
    # A candidate is never prefilled into the authoritative family/stop fields.
    stops = [{key: item.get(key) for key in STOP_COLUMNS} for item in items]
    for row in stops:
        row.update(stop=None, reason=None, effective_date=None, review_status='UNAPPROVED', reviewer=None)

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, '系列优先审核', GROUP_COLUMNS, group_rows)
    severity_order = {'BLOCKING': 0, 'ERROR': 1, 'WARNING': 2, 'INFO': 3}
    anomalies = sorted(queue.get('anomalies') or [],
                       key=lambda row: severity_order.get(str(row.get('severity')).upper(), 4))
    _sheet(wb, '数据异常', ANOMALY_COLUMNS, anomalies)
    exact_members = {
        (str(original.get('group_id')), original.get('sku'))
        for original in queue.get('items') or []
        if original.get('exact_reference_match') is True and original.get('mapping_status') != 'APPROVED'
    }
    # Each fixed group belongs to one marketplace and _items validated every
    # marketplace+SKU identity. A workbook match never becomes an approval.
    exact_items = [row for row in items if (str(row['group_id']), row['sku']) in exact_members]
    _sheet(wb, '精确匹配优先', MEMBER_COLUMNS, exact_items)
    ws = _sheet(wb, DECISION_SHEET, DECISION_COLUMNS, decisions,
                editable=('decision', 'approved_family', 'reviewer', 'reason'))
    decision_validation = DataValidation(type='list', formula1='"PENDING,APPROVED,REJECTED"')
    decision_validation.showErrorMessage = True
    ws.add_data_validation(decision_validation)
    for row_index, group in enumerate(group_rows, 2):
        eligible = group.get('batch_eligible') is True
        for key in ('decision', 'approved_family', 'reviewer', 'reason'):
            cell = ws.cell(row_index, DECISION_COLUMNS.index(key) + 1)
            cell.protection = Protection(locked=not eligible)
            if not eligible:
                cell.fill = PatternFill('solid', fgColor='E7E6E6')
                cell.comment = Comment('只按命名格式归类，未确认候选系列；不能批量批准。', 'Uloveido')
        if eligible:
            decision_validation.add(ws.cell(row_index, DECISION_COLUMNS.index('decision') + 1))
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    _sheet(wb, MEMBERS_SHEET, MEMBER_COLUMNS, items)
    ws = _sheet(wb, STOP_SHEET, STOP_COLUMNS, stops,
                editable=('stop', 'reason', 'effective_date', 'review_status', 'reviewer'))
    _validation(ws, STOP_COLUMNS, 'review_status', ('UNAPPROVED', 'APPROVED'))
    _validation(ws, STOP_COLUMNS, 'stop', ('TRUE', 'FALSE'))
    _sheet(wb, '审核审计', AUDIT_COLUMNS, queue.get('audit') or [])
    summary = queue.get('summary') or {}
    _sheet(wb, '审核统计', ('metric', 'value'),
           [{'metric': key, 'value': summary[key]} for key in summary if key in SUMMARY_KEYS])
    _instructions(wb)

    template = Workbook()
    template.remove(template.active)
    ws = _sheet(template, STOP_SHEET, STOP_COLUMNS, stops,
                editable=('family', 'stop', 'reason', 'effective_date', 'review_status', 'reviewer'))
    _validation(ws, STOP_COLUMNS, 'review_status', ('UNAPPROVED', 'APPROVED'))
    _validation(ws, STOP_COLUMNS, 'stop', ('TRUE', 'FALSE'))
    _instructions(template)
    dest = Path(output_dir)
    dest.mkdir(parents=True, exist_ok=True)
    review_path = dest / 'sku_batch_review.xlsx'
    template_path = dest / 'confirmed_stop_list_import_template.xlsx'
    wb.save(review_path)
    template.save(template_path)
    return {'review_path': str(review_path), 'stop_template_path': str(template_path)}
